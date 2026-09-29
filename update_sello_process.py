#!/usr/bin/env python3
"""Insert/update the approved detailed and generalized Sello + AIR process views."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import shutil
import tempfile
from zipfile import ZipFile

from pptx import Presentation
from pptx.util import Inches

import build_presentation as b
from update_slide25_test_evaluation import footer, snapshot, visible

ROOT = Path(__file__).resolve().parent
DECK = ROOT / 'VITA-FL_Thesis_Presentation_TU_Berlin.pptx'
ASSET = b.SELLO_PROCESS_ASSET
RECORD = ROOT / 'assets/sello-process-integration.json'
CHAPTER = ROOT.parent / 'overleaf/chapters/chapter4.tex'
COUNT = len(b.SELLO_PROCESS_SLUGS)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def check_structure(prs):
    target = b.sello_process_target(prs)
    results_count = b.external_results_count(prs)
    assert b.sello_process_count(prs) == COUNT
    trees = [(index, slide.name) for index, slide in enumerate(prs.slides)
             if slide.name.startswith(b.ATTACK_TREE_SLIDE_PREFIX)]
    expected_trees = [(target + COUNT + offset, b.ATTACK_TREE_SLIDE_PREFIX + slug)
                      for offset, slug in enumerate(b.ATTACK_TREE_SLUGS)]
    assert trees in ([], expected_trees), 'Unexpected attack-tree position or order'
    assert len(prs.slides) == 46 + COUNT + len(trees) + results_count, 'Unexpected talk/backup page count'
    assert len(prs.slide_masters) == 2
    evaluation = target + COUNT + len(trees)
    assert prs.slides[evaluation].name == b.EVALUATION_STACK_SLIDE_NAME
    assert 'One end-to-end Phala FedAvg run' in visible(prs.slides[evaluation + 1])
    assert prs.slides[evaluation + 2].name == b.TEST_EVALUATION_SLIDE_NAME
    assert 'Conclusion and outlook' in visible(prs.slides[evaluation + 5 + results_count])
    assert 'Training and inference share one verified image' in visible(prs.slides[evaluation + 6 + results_count])
    for page, slide in enumerate(prs.slides, 1):
        assert slide.notes_slide.notes_text_frame.text.strip(), f'Empty notes on Page {page}'
        if page > 1:
            assert footer(slide).text == f'Page {page}'
        ids = [node.get('id') for node in slide.element.xpath('.//p:cNvPr')]
        assert len(ids) == len(set(ids)), f'Duplicate shape IDs on Page {page}'
        if target < page <= target + COUNT:
            assert not slide.element.xpath('.//a:blip'), 'Process diagrams must remain native and editable'
            assert 'Variant B' not in visible(slide)
            for shape in slide.shapes:
                assert shape.left >= 0 and shape.top >= 0, shape.name
                assert shape.left + shape.width <= prs.slide_width + 12700, shape.name
                assert shape.top + shape.height <= prs.slide_height + 12700, shape.name
    return prs


def source_footer(shape):
    return (shape.has_text_frame and shape.top > Inches(6.7)
            and shape.text.startswith(('Variant ', 'Generalized B', 'Page ')))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--render', action='store_true', help='Refresh all deck previews')
    args = parser.parse_args()
    assert ASSET.is_file(), f'Approved two-slide asset is not ready: {ASSET}'
    prs = Presentation(DECK)
    old_count = len(prs.slides)
    target = b.sello_process_target(prs)
    results_count = b.external_results_count(prs)
    inserting = b.sello_process_count(prs) == 0
    assert old_count - results_count in ((46, 52) if inserting else (48, 54)), \
        'Unexpected deck size; refusing to overwrite'
    if inserting:
        assert (prs.slides[target].name == b.EVALUATION_STACK_SLIDE_NAME
                or prs.slides[target].name == b.ATTACK_TREE_SLIDE_PREFIX + b.ATTACK_TREE_SLUGS[0])
    else:
        check_structure(prs)
    source = Presentation(ASSET)
    assert len(source.slides) == COUNT
    before = [snapshot(slide) for slide in prs.slides]
    before_hash = digest(DECK)
    chapter_hash = digest(CHAPTER)
    backup = Path(tempfile.mkdtemp(prefix='vita-fl-before-sello-process-')) / DECK.name
    shutil.copy2(DECK, backup)
    if inserting:
        b.insert_sello_process_slides(prs)
    else:
        for offset, (asset_slide, slug) in enumerate(zip(source.slides, b.SELLO_PROCESS_SLUGS)):
            b.apply_sello_process_asset(prs.slides[target + offset], asset_slide, target + offset + 1, slug)
    check_structure(prs)

    handle, filename = tempfile.mkstemp(prefix='.sello-process-', suffix='.pptx', dir=ROOT)
    os.close(handle)
    candidate = Path(filename)
    try:
        prs.save(candidate)
        checked = check_structure(Presentation(candidate))
        preserved = 0
        for index, old in enumerate(before):
            if not inserting and target <= index < target + COUNT:
                continue
            new_index = index + (COUNT if inserting and index >= target else 0)
            assert snapshot(checked.slides[new_index]) == old, \
                f'Existing Page {index + 1} changed beyond its page-number footer'
            preserved += 1
        for offset, asset_slide in enumerate(source.slides):
            integrated = checked.slides[target + offset]
            assert integrated.notes_slide.notes_text_frame.text == asset_slide.notes_slide.notes_text_frame.text
            assert len(integrated.shapes) == len(asset_slide.shapes)
            assert sum(source_footer(shape) for shape in asset_slide.shapes) == 1
            for approved, actual in zip(asset_slide.shapes, integrated.shapes):
                if source_footer(approved):
                    continue
                assert str(approved.element.xml) == str(actual.element.xml), (offset, approved.name)
        with ZipFile(backup) as old_zip, ZipFile(candidate) as new_zip:
            assert new_zip.testzip() is None
            for name in old_zip.namelist():
                if name.startswith(('ppt/media/', 'ppt/slideMasters/', 'ppt/slideLayouts/', 'ppt/notesMasters/')):
                    assert old_zip.read(name) == new_zip.read(name), name
        # Reapplying the approved assets to the candidate must not change its content.
        once = [snapshot(slide) for slide in checked.slides]
        for offset, (asset_slide, slug) in enumerate(zip(source.slides, b.SELLO_PROCESS_SLUGS)):
            b.apply_sello_process_asset(checked.slides[target + offset], asset_slide, target + offset + 1, slug)
        assert [snapshot(slide) for slide in checked.slides] == once, 'Asset replacement is not idempotent'
        assert digest(CHAPTER) == chapter_hash
        assert digest(DECK) == before_hash, 'Main deck changed during validation; refusing to overwrite'
        os.replace(candidate, DECK)
    finally:
        candidate.unlink(missing_ok=True)

    new_count = len(prs.slides)
    pages = list(range(target + 1, target + COUNT + 1))
    RECORD.write_text(json.dumps({
        'pages': pages, 'slides_before': old_count, 'slides_after': new_count,
        'talk_pages': new_count - 17, 'backup_pages': 17,
        'external_results_pages': results_count,
        'mode': 'insert' if inserting else 'replace',
        'preserved_other_slides_and_notes': preserved,
        'allowed_existing_change': 'page-number footers only on insertion; none outside process slides on update',
        'main_deck_sha256_before': before_hash, 'main_deck_sha256_after': digest(DECK),
        'asset_sha256': digest(ASSET), 'chapter4_sha256': chapter_hash, 'chapter4_unchanged': True,
        'backup': str(backup), 'structure_check': 'passed', 'idempotent_asset_update': 'passed',
        'approved_shape_xml_preserved_except_footer': True, 'source_notes_preserved': True,
        'media_masters_layouts_unchanged': True, 'native_editable_shapes': True,
        'visual_review': 'pending',
    }, indent=2) + '\n')
    if args.render:
        import render_preview
        render_preview.main()
    print(f'{"Inserted" if inserting else "Updated"} Sello + AIR on Pages {pages[0]}–{pages[-1]}; '
          f'{new_count} pages ({new_count - 17} talk, 17 backup).')
    print(f'Preserved {preserved} existing slides and notes; backup: {backup}')
    print(f'Checks: {RECORD}')


if __name__ == '__main__':
    main()
