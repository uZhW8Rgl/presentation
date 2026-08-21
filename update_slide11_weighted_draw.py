#!/usr/bin/env python3
"""Clarify the weighted on-chain draw on slide 11 without rebuilding the deck."""

from __future__ import annotations

import os
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

from build_presentation import (
    BLUE,
    BLUE_TINT,
    DARK,
    GREEN,
    GREEN_TINT,
    LIGHT,
    TU_RED,
    add_box,
    add_note,
    add_rich_text,
    add_text,
    rgb,
    set_run_language,
)


ROOT = Path(__file__).resolve().parent
DECK = ROOT / "VITA-FL_Thesis_Presentation_TU_Berlin.pptx"
BACKUP = Path("/tmp/VITA-FL_Thesis_Presentation_TU_Berlin.before-slide11.pptx")
TEMP = ROOT / ".VITA-FL_Thesis_Presentation_TU_Berlin.slide11.tmp.pptx"
SOURCE_URL = "https://eips.ethereum.org/EIPS/eip-4399"


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def main():
    shutil.copy2(DECK, BACKUP)
    prs = Presentation(DECK)
    assert len(prs.slides) >= 11
    slide = prs.slides[10]

    title_text = "\n".join(
        shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)
    )
    assert "Weighted aggregator selection" in title_text
    assert "SCORE CHANGES" in title_text
    assert "SCOPE" in title_text

    formula = next(
        shape
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
        and shape.text.strip().startswith("P(worker i)")
    )
    formula.text_frame.clear()
    formula.text_frame.word_wrap = True
    formula.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    formula.text_frame.margin_left = formula.text_frame.margin_right = 0
    formula.text_frame.margin_top = formula.text_frame.margin_bottom = 0
    paragraph = formula.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = "P(worker i) = (score_i + 1) / Σ(score_j + 1)"
    run.font.name = "Arial"
    run.font.bold = True
    run.font.size = Pt(17)
    run.font.color.rgb = rgb(DARK)
    set_run_language(run)

    # Remove the previous one-line draw description and both lower explanation panels.
    removable = [
        shape
        for shape in list(slide.shapes)
        if 4.25 <= shape.top / 914400 < 5.90
    ]
    assert len(removable) == 7, len(removable)
    for shape in removable:
        remove_shape(shape)

    add_box(slide, 0.82, 4.38, 11.70, 0.64, fill=LIGHT, line=TU_RED, radius=False)
    add_rich_text(
        slide,
        [
            ("1 · HASH DRAW CONTEXT   ", TU_RED, True, 10.0),
            (
                "r = uint256(keccak256(prevrandao, timestamp, block number, round, "
                "current aggregator public address))",
                DARK,
                True,
                10.2,
            ),
        ],
        1.02,
        4.57,
        11.30,
        0.22,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    add_box(slide, 0.82, 5.16, 5.78, 0.78, fill=BLUE_TINT, line=BLUE)
    add_text(slide, "2 · DRAW ONE TICKET", 1.08, 5.31, 2.18, 0.20, 10.5, BLUE, True)
    add_text(
        slide,
        "ticket = r mod W = r mod 14  ∈  {0, …, 13}",
        1.08,
        5.57,
        5.20,
        0.22,
        12.2,
        DARK,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    add_box(slide, 6.78, 5.16, 5.74, 0.78, fill=GREEN_TINT, line=GREEN)
    add_text(slide, "3 · MAP TO CUMULATIVE INTERVAL", 7.04, 5.31, 3.20, 0.20, 10.5, GREEN, True)
    add_text(
        slide,
        "A: 0   ·   B: 1–2   ·   C: 3–6   ·   D: 7–13   |   Example: 5 → C",
        7.04,
        5.57,
        5.18,
        0.22,
        10.4,
        DARK,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    source = add_rich_text(
        slide,
        [
            ("prevrandao pseudorandom value: ", DARK, False, 7.5),
            (SOURCE_URL, TU_RED, False, 7.5),
        ],
        0.90,
        6.10,
        11.54,
        0.16,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    source.text_frame.paragraphs[0].runs[-1].hyperlink.address = SOURCE_URL

    add_note(
        slide,
        "Each authorized worker receives weight score plus one. In the example the weights sum to 14, "
        "so the displayed probabilities are one, two, four, and seven divided by 14. For the normal "
        "selection, the contract hashes prevrandao, the block timestamp, block number, round, and the "
        "current aggregator public address. The resulting 256-bit value is converted to an integer and "
        "reduced modulo the total weight. This produces one ticket from zero through 13. The ticket is "
        "then mapped into the cumulative intervals in the displayed worker order: A receives zero, B "
        "one through two, C three through six, and D seven through 13. A ticket of five therefore selects "
        "Worker C. The interval widths—not the hash itself—create the probabilities. prevrandao is the "
        "EVM-accessible beacon-chain RANDAO value specified by EIP-4399 and is pseudorandom rather than "
        "an unbiased randomness oracle. Source: https://eips.ethereum.org/EIPS/eip-4399.",
    )

    prs.save(TEMP)

    check = Presentation(TEMP)
    check_slide = check.slides[10]
    visible = "\n".join(
        shape.text for shape in check_slide.shapes if getattr(shape, "has_text_frame", False)
    )
    for required in (
        "P(worker i) = (score_i + 1) / Σ(score_j + 1)",
        "current aggregator public address",
        "ticket = r mod W = r mod 14",
        "A: 0",
        SOURCE_URL,
    ):
        assert required in visible, required
    for removed in ("SCORE CHANGES", "SCOPE", "incumbent"):
        assert removed not in visible, removed

    os.replace(TEMP, DECK)
    print(DECK)
    print(f"backup: {BACKUP}")


if __name__ == "__main__":
    main()
