#!/usr/bin/env python3
"""Insert/update approved attack trees after the chain-of-custody / Sello section."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import shutil
import tempfile
from zipfile import ZipFile

from pptx import Presentation

import build_presentation as b
from update_slide25_test_evaluation import (
    attack_tree_count, attack_tree_target, chain_of_custody_index, footer,
    sello_process_count, snapshot, visible,
)

ROOT = Path(__file__).resolve().parent
DECK = ROOT / 'VITA-FL_Thesis_Presentation_TU_Berlin.pptx'
ASSET = b.ATTACK_TREE_ASSET
RECORD = ROOT / 'assets/attack-trees-integration.json'
CHAPTER = ROOT.parent / 'overleaf/chapters/chapter4.tex'
COUNT = len(b.ATTACK_TREE_SLUGS)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def check_structure(prs):
    target = attack_tree_target(prs)
    process_count = sello_process_count(prs)
    results_count = b.external_results_count(prs)
    assert len(prs.slides) == 52 + process_count + results_count, 'Unexpected talk/backup page count'
    assert attack_tree_count(prs) == COUNT
    assert len(prs.slide_masters) == 2
    names = [b.ATTACK_TREE_SLIDE_PREFIX + slug for slug in b.ATTACK_TREE_SLUGS]
    assert [prs.slides[i].name for i in range(target, target + COUNT)] == names
    assert sum(s.name.startswith(b.ATTACK_TREE_SLIDE_PREFIX) for s in prs.slides) == COUNT
    for i, text in ((chain_of_custody_index(prs), 'Cryptographic chain of custody'),
                    (target + COUNT + 1, 'One end-to-end Phala FedAvg run'),
                    (target + COUNT + 5 + results_count, 'Conclusion and outlook'),
                    (target + COUNT + 6 + results_count, 'Training and inference share one verified image')):
        assert text in visible(prs.slides[i]), (i, text)
    assert prs.slides[target + COUNT].name == b.EVALUATION_STACK_SLIDE_NAME
    assert prs.slides[target + COUNT + 2].name == b.TEST_EVALUATION_SLIDE_NAME
    for page, slide in enumerate(prs.slides, 1):
        assert slide.notes_slide.notes_text_frame.text.strip(), f'Empty notes on Page {page}'
        if page > 1:
            assert footer(slide).text == f'Page {page}'
        ids = [node.get('id') for node in slide.element.xpath('.//p:cNvPr')]
        assert len(ids) == len(set(ids)), f'Duplicate shape IDs on Page {page}'
        if target < page <= target + COUNT:
            assert 'TREE-SPECIFIC MODEL' in slide.notes_slide.notes_text_frame.text
            assert 'B light ·' not in visible(slide)
            assert sum(s.name.startswith('Attack legend swatch: ') for s in slide.shapes) == 4
            assert sum(s.name.startswith('Attack legend label: ') for s in slide.shapes) == 4
            assert not slide.element.xpath('.//a:blip'), 'Attack trees must remain native and editable'
            for s in slide.shapes:
                assert s.left >= 0 and s.top >= 0
                assert s.left + s.width <= prs.slide_width + 12700
                assert s.top + s.height <= prs.slide_height + 12700
    return prs


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--render', action='store_true', help='Refresh all deck previews')
    args = parser.parse_args()
    assert ASSET.is_file()
    prs = Presentation(DECK)
    old_count = len(prs.slides)
    process_count = sello_process_count(prs)
    results_count = b.external_results_count(prs)
    target = attack_tree_target(prs)
    assert old_count in (46 + process_count + results_count, 52 + process_count + results_count), 'Unexpected deck size; refusing to overwrite'
    inserting = attack_tree_count(prs) == 0
    assert old_count == (46 if inserting else 52) + process_count + results_count
    if inserting:
        assert 'Cryptographic chain of custody' in visible(prs.slides[chain_of_custody_index(prs)])
        assert prs.slides[target].name == b.EVALUATION_STACK_SLIDE_NAME
        assert not any(s.name.startswith(b.ATTACK_TREE_SLIDE_PREFIX) for s in prs.slides)
    else:
        check_structure(prs)
    source = Presentation(ASSET)
    assert len(source.slides) == COUNT
    before = [snapshot(s) for s in prs.slides]
    before_hash = digest(DECK)
    chapter_hash = digest(CHAPTER)
    backup = Path(tempfile.mkdtemp(prefix='vita-fl-before-attack-trees-')) / DECK.name
    shutil.copy2(DECK, backup)
    if inserting:
        b.insert_attack_tree_slides(prs, target)
    else:
        for i, (asset_slide, slug) in enumerate(zip(source.slides, b.ATTACK_TREE_SLUGS)):
            b.apply_attack_tree_asset(prs.slides[target + i], asset_slide, target + i + 1, slug)
    check_structure(prs)

    handle, filename = tempfile.mkstemp(prefix='.attack-trees-', suffix='.pptx', dir=ROOT)
    os.close(handle)
    candidate = Path(filename)
    try:
        prs.save(candidate)
        checked = check_structure(Presentation(candidate))
        preserved = 0
        for index, old in enumerate(before):
            if not inserting and target <= index < target + COUNT:
                continue
            new_index = index + (COUNT if inserting and index >= target else 0)
            assert snapshot(checked.slides[new_index]) == old, \
                f'Existing Page {index + 1} changed beyond its footer'
            preserved += 1
        for i, asset_slide in enumerate(source.slides):
            integrated = checked.slides[target + i]
            assert integrated.notes_slide.notes_text_frame.text == asset_slide.notes_slide.notes_text_frame.text
            assert len(integrated.shapes) == len(asset_slide.shapes)
            for approved, actual in zip(asset_slide.shapes, integrated.shapes):
                if approved.has_text_frame and approved.text.startswith('B light · '):
                    continue
                assert approved.element.xml == actual.element.xml, (i, approved.name)
        with ZipFile(backup) as old_zip, ZipFile(candidate) as new_zip:
            assert new_zip.testzip() is None
            for name in old_zip.namelist():
                if name.startswith(('ppt/media/', 'ppt/slideMasters/', 'ppt/slideLayouts/', 'ppt/notesMasters/')):
                    assert old_zip.read(name) == new_zip.read(name), name
        # Exercise the same generator hook without regenerating the user's existing slides.
        generated = b.prepare_template()
        for _ in range(target):
            generated.slides.add_slide(generated.slide_masters[1].slide_layouts[0])
        b.insert_attack_tree_slides(generated, target)
        for i in range(COUNT):
            assert visible(generated.slides[target + i]) == visible(checked.slides[target + i])
            assert generated.slides[target + i].notes_slide.notes_text_frame.text == source.slides[i].notes_slide.notes_text_frame.text
        assert digest(CHAPTER) == chapter_hash
        os.replace(candidate, DECK)
    finally:
        candidate.unlink(missing_ok=True)

    RECORD.write_text(json.dumps({
        'pages': list(range(target + 1, target + COUNT + 1)),
        'slides_before': old_count, 'slides_after': len(checked.slides),
        'talk_pages': len(checked.slides) - 17, 'backup_pages': 17,
        'sello_process_pages': process_count, 'external_results_pages': results_count, 'mode': 'insert' if inserting else 'replace',
        'preserved_other_slides_and_notes': preserved,
        'allowed_existing_change': 'page-number footers only on insertion; none outside trees on update',
        'main_deck_sha256_before': before_hash, 'main_deck_sha256_after': digest(DECK),
        'asset_sha256': digest(ASSET), 'chapter4_sha256': chapter_hash, 'chapter4_unchanged': True,
        'backup': str(backup), 'structure_check': 'passed', 'generator_hook': 'passed',
        'approved_shape_xml_preserved_except_footer': True, 'source_notes_preserved': True,
        'media_masters_layouts_unchanged': True, 'native_editable_shapes': True,
        'visual_review': 'pending',
    }, indent=2) + '\n')
    if args.render:
        import render_preview
        render_preview.main()
    print(f'{"Inserted" if inserting else "Updated"} attack trees on Pages {target + 1}–{target + COUNT}; '
          f'{len(checked.slides)} pages ({len(checked.slides) - 17} talk, 17 backup).')
    print(f'Preserved {preserved} existing slides and notes; backup: {backup}')
    print(f'Checks: {RECORD}')


if __name__ == '__main__':
    main()
