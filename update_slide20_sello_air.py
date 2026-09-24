#!/usr/bin/env python3
"""Apply the approved Sello/AIR asset to page 20 while preserving other slides."""
from __future__ import annotations

import hashlib
import json
import math
import os
from pathlib import Path
import shutil
import sys
import tempfile
from types import SimpleNamespace
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation

import build_presentation as b
import render_preview as renderer

ROOT = Path(__file__).resolve().parent
DECK = ROOT / 'VITA-FL_Thesis_Presentation_TU_Berlin.pptx'
ASSET = ROOT / 'assets/sello-air-a.pptx'
TARGET = 19


def visible(slide):
    return '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)


def snapshot(prs):
    return [(s.element.xml.encode("utf-8"), s.notes_slide._element.xml.encode("utf-8")) for s in prs.slides]


def main():
    prs = Presentation(DECK)
    count = len(prs.slides)
    assert count == 44
    assert any(title in visible(prs.slides[TARGET]) for title in (
        'Sello: evidence from the receiving service', 'Sello and AIR: two layers of evidence',
    ))
    before = snapshot(prs)
    original_hash = hashlib.sha256(DECK.read_bytes()).hexdigest()
    backup = Path(tempfile.mkdtemp(prefix='vita-fl-before-sello-air-')) / DECK.name
    shutil.copy2(DECK, backup)
    b.apply_sello_air_asset(prs.slides[TARGET], TARGET+1)

    # Check the changed slide only; other slides retain their existing edits.
    sys.path.insert(0, str(ROOT / 'concepts/color-navigation'))
    from build_concepts import check_layout
    check_layout(SimpleNamespace(slides=[prs.slides[TARGET]], slide_width=prs.slide_width,
                                 slide_height=prs.slide_height))
    fd, candidate_name = tempfile.mkstemp(prefix='.sello-air-', suffix='.pptx', dir=ROOT)
    os.close(fd)
    candidate = Path(candidate_name)
    try:
        prs.save(candidate)
        checked = Presentation(candidate)
        assert len(checked.slides) == count
        after = snapshot(checked)
        assert all(before[i] == after[i] for i in range(count) if i != TARGET), \
            'An unrelated slide or its notes changed'
        slide = checked.slides[TARGET]
        text = visible(slide)
        for required in ('Page 20', 'Attested Inference Receipt', 'Inference TEE', 'Transparency Log'):
            assert required in text, required
        for removed in ('Scope:', 'SCITT', 'Variant A', 'Receiver'):
            assert removed not in text, removed
        shapes = {shape.name: shape for shape in slide.shapes}
        assert str(shapes['Domain navigation: DFL badge'].fill.fore_color.rgb) == 'F5F6F7'
        assert str(shapes['Domain navigation: Agent badge'].fill.fore_color.rgb) == '1764A1'
        ids = [shape.shape_id for shape in slide.shapes]
        assert len(ids) == len(set(ids))
        links = [run.hyperlink.address for shape in slide.shapes if shape.has_text_frame
                 for p in shape.text_frame.paragraphs for run in p.runs if run.hyperlink.address]
        assert len(links) == 2
        with ZipFile(backup) as old, ZipFile(candidate) as new:
            assert new.testzip() is None
            for name in old.namelist():
                if name.startswith(('ppt/media/', 'ppt/slideMasters/', 'ppt/slideLayouts/')):
                    assert new.read(name) == old.read(name), name
        # Verify the build path independently without rebuilding the full deck.
        generated = b.prepare_template()
        b.slide_sello_protocol(generated)
        assert 'Attested Inference Receipt' in visible(generated.slides[0])
        assert 'Transparency Log' in visible(generated.slides[0])
        check_layout(generated)
        os.replace(candidate, DECK)
    finally:
        candidate.unlink(missing_ok=True)

    preview_dir = ROOT / 'preview'
    renderer.render_slide(checked, checked.slides[TARGET]).save(preview_dir/'slide-20.png')
    thumb_w, thumb_h, gutter = 720, 405, 28
    rows = math.ceil(count/2)
    sheet = Image.new('RGB', (thumb_w*2+gutter*3, thumb_h*rows+gutter*(rows+1)), (221,229,234))
    for index in range(count):
        with Image.open(preview_dir/f'slide-{index+1:02d}.png') as preview:
            thumb = preview.resize((thumb_w,thumb_h), Image.Resampling.LANCZOS)
        sheet.paste(thumb, (gutter+(index%2)*(thumb_w+gutter), gutter+(index//2)*(thumb_h+gutter)))
    sheet.save(preview_dir/'contact-sheet.png')
    record = {
        'page':20,'slides':count,'unchanged_other_slides_and_notes':count-1,
        'main_deck_sha256_before':original_hash,
        'main_deck_sha256_after':hashlib.sha256(DECK.read_bytes()).hexdigest(),
        'asset_sha256':hashlib.sha256(ASSET.read_bytes()).hexdigest(),
        'backup':str(backup),'layout':'passed','generator_import':'passed',
        'source_hyperlinks':len(links),'dfl_chip':'inactive','agent_chip':'active',
        'media_masters_and_layouts_unchanged':True,'visual_review':'pending',
    }
    destination=ROOT/'concepts/sello-air/guarantees/integration.json'
    destination.write_text(json.dumps(record,indent=2)+'\n')
    print(f'Updated page 20 in {DECK.name}. Preserved all other {count-1} slides and notes.')
    print(f'Backup: {backup}')


if __name__ == '__main__':
    main()
