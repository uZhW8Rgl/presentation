#!/usr/bin/env python3
"""Render an approximate PNG preview of the generated deck.

This is intentionally a lightweight visual QA renderer. PowerPoint remains the
authoritative renderer, but the preview catches misplaced shapes, clipping, and
unbalanced layouts in environments without LibreOffice or Microsoft Office.
"""

from __future__ import annotations

import math
import re
import xml.etree.ElementTree as ET
from io import BytesIO
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN


ROOT = Path(__file__).resolve().parent
DECK = ROOT / "VITA-FL_Thesis_Presentation_TU_Berlin.pptx"
OUTPUT = ROOT / "preview"

WIDTH = 1600
HEIGHT = 900
FONT_REGULAR = "/usr/share/fonts/dejavu/DejaVuSans.ttf"
FONT_BOLD = "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf"
FONT_MONO = "/usr/share/fonts/dejavu/DejaVuSansMono.ttf"

SVG_BLIP = "{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip"
REL_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
SVG_NUMBER = re.compile(r"[A-Za-z]|[-+]?(?:\d*\.\d+|\d+\.?)(?:[eE][-+]?\d+)?")


def rgb_tuple(value, fallback=(0, 0, 0)):
    try:
        return tuple(value.rgb)
    except Exception:
        return fallback


def alpha_for(color_format):
    try:
        color_element = color_format._color._xClr
        alpha = color_element.find("{http://schemas.openxmlformats.org/drawingml/2006/main}alpha")
        if alpha is not None:
            return round(int(alpha.get("val")) * 255 / 100000)
    except Exception:
        pass
    return 255


def font_for(run, scale):
    size_pt = float(run.font.size.pt) if run.font.size is not None else 16.0
    size_px = max(8, round(size_pt * scale))
    name = (run.font.name or "").lower()
    if "mono" in name:
        path = FONT_MONO
    elif run.font.bold:
        path = FONT_BOLD
    else:
        path = FONT_REGULAR
    return ImageFont.truetype(path, size_px)


def text_color(run):
    try:
        return tuple(run.font.color.rgb) + (255,)
    except Exception:
        return (244, 250, 253, 255)


def split_tokens(text):
    return [token for token in re.split(r"(\s+)", text) if token]


def _svg_path_polygons(path_data, curve_steps=18):
    """Approximate the common SVG path commands as filled polygons."""
    tokens = SVG_NUMBER.findall(path_data.replace(",", " "))
    index = 0
    command = None
    current = (0.0, 0.0)
    start = current
    previous_control = None
    polygon = []
    polygons = []

    def is_command(token):
        return len(token) == 1 and token.isalpha()

    def take(count):
        nonlocal index
        if index + count > len(tokens) or any(is_command(token) for token in tokens[index:index + count]):
            return None
        values = [float(token) for token in tokens[index:index + count]]
        index += count
        return values

    def absolute_pair(x, y, relative):
        return (current[0] + x, current[1] + y) if relative else (x, y)

    def finish_polygon():
        nonlocal polygon
        if len(polygon) >= 3:
            polygons.append(polygon)
        polygon = []

    while index < len(tokens):
        if is_command(tokens[index]):
            command = tokens[index]
            index += 1
        if command is None:
            break

        relative = command.islower()
        op = command.upper()
        if op == "Z":
            if polygon and polygon[-1] != start:
                polygon.append(start)
            current = start
            finish_polygon()
            previous_control = None
            command = None
            continue

        if op == "M":
            values = take(2)
            if values is None:
                command = None
                continue
            point = absolute_pair(values[0], values[1], relative)
            if polygon:
                finish_polygon()
            polygon = [point]
            current = start = point
            previous_control = None
            command = "l" if relative else "L"
            continue

        if op == "L":
            values = take(2)
            if values is None:
                command = None
                continue
            current = absolute_pair(values[0], values[1], relative)
            polygon.append(current)
            previous_control = None
            continue

        if op == "H":
            values = take(1)
            if values is None:
                command = None
                continue
            current = (current[0] + values[0] if relative else values[0], current[1])
            polygon.append(current)
            previous_control = None
            continue

        if op == "V":
            values = take(1)
            if values is None:
                command = None
                continue
            current = (current[0], current[1] + values[0] if relative else values[0])
            polygon.append(current)
            previous_control = None
            continue

        if op in {"C", "S"}:
            values = take(6 if op == "C" else 4)
            if values is None:
                command = None
                continue
            p0 = current
            if op == "C":
                p1 = absolute_pair(values[0], values[1], relative)
                p2 = absolute_pair(values[2], values[3], relative)
                p3 = absolute_pair(values[4], values[5], relative)
            else:
                p1 = (
                    2 * p0[0] - previous_control[0],
                    2 * p0[1] - previous_control[1],
                ) if previous_control else p0
                p2 = absolute_pair(values[0], values[1], relative)
                p3 = absolute_pair(values[2], values[3], relative)
            for step in range(1, curve_steps + 1):
                t = step / curve_steps
                u = 1 - t
                polygon.append((
                    u ** 3 * p0[0] + 3 * u ** 2 * t * p1[0] + 3 * u * t ** 2 * p2[0] + t ** 3 * p3[0],
                    u ** 3 * p0[1] + 3 * u ** 2 * t * p1[1] + 3 * u * t ** 2 * p2[1] + t ** 3 * p3[1],
                ))
            current = p3
            previous_control = p2
            continue

        # Unsupported commands are skipped safely instead of breaking all previews.
        command = None

    if polygon:
        finish_polygon()
    return polygons


def _svg_picture(shape, width, height):
    svg_blip = shape._element.find(f".//{SVG_BLIP}")
    if svg_blip is None:
        return None
    relation_id = svg_blip.get(REL_EMBED)
    if not relation_id:
        return None
    blob = shape.part.related_part(relation_id).blob
    root = ET.fromstring(blob)
    view_box = [float(value) for value in root.attrib.get("viewBox", "0 0 1 1").replace(",", " ").split()]
    if len(view_box) != 4 or view_box[2] == 0 or view_box[3] == 0:
        return None
    vx, vy, vw, vh = view_box
    antialias = 4
    raster_width = max(1, width * antialias)
    raster_height = max(1, height * antialias)
    raster = Image.new("RGBA", (raster_width, raster_height), (0, 0, 0, 0))
    draw = ImageDraw.Draw(raster)

    for element in root.iter():
        if not element.tag.endswith("path") or not element.get("d"):
            continue
        fill = element.get("fill", "#000000")
        if fill == "none" or not fill.startswith("#"):
            continue
        if len(fill) == 4:
            fill = "#" + "".join(character * 2 for character in fill[1:])
        color = tuple(int(fill[offset:offset + 2], 16) for offset in (1, 3, 5)) + (255,)
        for polygon in _svg_path_polygons(element.get("d")):
            points = [
                (
                    round((x - vx) / vw * raster_width),
                    round((y - vy) / vh * raster_height),
                )
                for x, y in polygon
            ]
            draw.polygon(points, fill=color)
    return raster.resize((max(1, width), max(1, height)), Image.Resampling.LANCZOS)


def paragraph_lines(draw, paragraph, max_width, scale):
    lines = [[]]
    line_width = 0
    runs = list(paragraph.runs)
    if not runs and paragraph.text:
        return [[(paragraph.text, ImageFont.truetype(FONT_REGULAR, round(16 * scale)), (244, 250, 253, 255))]]

    for run in runs:
        font = font_for(run, scale)
        color = text_color(run)
        chunks = run.text.split("\n")
        for chunk_index, chunk in enumerate(chunks):
            for token in split_tokens(chunk):
                token_width = draw.textlength(token, font=font)
                if not token.isspace() and line_width + token_width > max_width and lines[-1]:
                    lines.append([])
                    line_width = 0
                if token.isspace() and not lines[-1]:
                    continue
                lines[-1].append((token, font, color))
                line_width += token_width
            if chunk_index < len(chunks) - 1:
                lines.append([])
                line_width = 0
    return lines


def render_text(layer, shape, sx, sy):
    frame = shape.text_frame
    draw = ImageDraw.Draw(layer)
    left = round(shape.left * sx + frame.margin_left * sx)
    top = round(shape.top * sy + frame.margin_top * sy)
    right = round((shape.left + shape.width) * sx - frame.margin_right * sx)
    bottom = round((shape.top + shape.height) * sy - frame.margin_bottom * sy)
    max_width = max(1, right - left)
    scale = HEIGHT / (7.5 * 72)

    paragraph_blocks = []
    for paragraph in frame.paragraphs:
        lines = paragraph_lines(draw, paragraph, max_width, scale)
        line_metrics = []
        for line in lines:
            ascents = []
            descents = []
            for _, font, _ in line:
                ascent, descent = font.getmetrics()
                ascents.append(ascent)
                descents.append(descent)
            max_ascent = max(ascents or [round(16 * scale)])
            max_descent = max(descents or [round(4 * scale)])
            line_height = (max_ascent + max_descent) * 1.08
            line_metrics.append((line_height, max_ascent))
        paragraph_blocks.append((paragraph, lines, line_metrics))

    total_height = 0
    for index, (_, _, metrics) in enumerate(paragraph_blocks):
        total_height += sum(line_height for line_height, _ in metrics)
        if index < len(paragraph_blocks) - 1:
            total_height += round(5 * scale)

    anchor = frame.vertical_anchor
    if anchor == MSO_ANCHOR.MIDDLE:
        y = top + max(0, (bottom - top - total_height) / 2)
    elif anchor == MSO_ANCHOR.BOTTOM:
        y = bottom - total_height
    else:
        y = top

    for block_index, (paragraph, lines, metrics) in enumerate(paragraph_blocks):
        for line, (line_height, max_ascent) in zip(lines, metrics):
            widths = [draw.textlength(text, font=font) for text, font, _ in line]
            line_width = sum(widths)
            if paragraph.alignment == PP_ALIGN.CENTER:
                x = left + (max_width - line_width) / 2
            elif paragraph.alignment == PP_ALIGN.RIGHT:
                x = right - line_width
            else:
                x = left
            baseline = y + max_ascent
            for (text, font, color), token_width in zip(line, widths):
                draw.text(
                    (round(x), round(baseline)),
                    text,
                    font=font,
                    fill=color,
                    anchor="ls",
                )
                x += token_width
            y += line_height
        if block_index < len(paragraph_blocks) - 1:
            y += round(5 * scale)


def render_shape(canvas, shape, sx, sy):
    layer = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(layer)
    x0 = round(shape.left * sx)
    y0 = round(shape.top * sy)
    x1 = round((shape.left + shape.width) * sx)
    y1 = round((shape.top + shape.height) * sy)
    if x1 <= 0 or y1 <= 0 or x0 >= canvas.width or y0 >= canvas.height:
        return

    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        target_width = max(1, x1 - x0)
        target_height = max(1, y1 - y0)
        try:
            image = Image.open(BytesIO(shape.image.blob)).convert("RGBA")
            image = image.resize((target_width, target_height), Image.Resampling.LANCZOS)
        except ValueError:
            image = _svg_picture(shape, target_width, target_height)
        if image is not None:
            layer.alpha_composite(image, (x0, y0))

    elif shape.shape_type == MSO_SHAPE_TYPE.LINE:
        transform = shape._element.spPr.xfrm
        flip_h = bool(transform.flipH)
        flip_v = bool(transform.flipV)
        start = (x1 if flip_h else x0, y1 if flip_v else y0)
        end = (x0 if flip_h else x1, y0 if flip_v else y1)
        color = rgb_tuple(shape.line.color, (175, 196, 207))
        width = max(1, round((shape.line.width or 12700) / 12700 * HEIGHT / (7.5 * 72)))
        dashed = shape.line.dash_style is not None
        if dashed:
            length = math.dist(start, end)
            if length:
                ux = (end[0] - start[0]) / length
                uy = (end[1] - start[1]) / length
                step = max(8, width * 4)
                pos = 0
                while pos < length:
                    stop = min(length, pos + step)
                    draw.line((start[0] + ux * pos, start[1] + uy * pos, start[0] + ux * stop, start[1] + uy * stop), fill=color + (255,), width=width)
                    pos += step * 1.75
        else:
            draw.line((*start, *end), fill=color + (255,), width=width)

        line_properties = shape._element.spPr.ln
        tail = None if line_properties is None else line_properties.find("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
        if tail is not None and tail.get("type") != "none":
            angle = math.atan2(end[1] - start[1], end[0] - start[0])
            length = max(10, width * 5)
            spread = math.pi / 6
            points = [
                end,
                (end[0] - length * math.cos(angle - spread), end[1] - length * math.sin(angle - spread)),
                (end[0] - length * math.cos(angle + spread), end[1] - length * math.sin(angle + spread)),
            ]
            draw.polygon(points, fill=color + (255,))

    elif shape.shape_type in {MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM}:
        fill = None
        if shape.fill.type == MSO_FILL_TYPE.SOLID:
            fill = rgb_tuple(shape.fill.fore_color, (16, 44, 58)) + (alpha_for(shape.fill.fore_color),)
        outline = None
        try:
            outline = rgb_tuple(shape.line.color, (49, 82, 99)) + (alpha_for(shape.line.color),)
        except Exception:
            pass
        line_width = max(1, round((shape.line.width or 12700) / 12700 * HEIGHT / (7.5 * 72)))
        try:
            kind = shape.auto_shape_type
        except (AttributeError, ValueError):
            kind = None
        if kind == MSO_AUTO_SHAPE_TYPE.OVAL:
            draw.ellipse((x0, y0, x1, y1), fill=fill, outline=outline, width=line_width)
        elif kind in {MSO_AUTO_SHAPE_TYPE.DIAMOND, MSO_AUTO_SHAPE_TYPE.FLOWCHART_DECISION}:
            points = [
                ((x0 + x1) // 2, y0),
                (x1, (y0 + y1) // 2),
                ((x0 + x1) // 2, y1),
                (x0, (y0 + y1) // 2),
            ]
            draw.polygon(points, fill=fill)
            if outline:
                draw.line(points + [points[0]], fill=outline, width=line_width, joint="curve")
        elif kind == MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE:
            radius = max(4, round(min(x1 - x0, y1 - y0) * 0.12))
            draw.rounded_rectangle((x0, y0, x1, y1), radius=radius, fill=fill, outline=outline, width=line_width)
        else:
            draw.rectangle((x0, y0, x1, y1), fill=fill, outline=outline, width=line_width)

    if shape.has_text_frame and any(paragraph.text for paragraph in shape.text_frame.paragraphs):
        render_text(layer, shape, sx, sy)
    canvas.alpha_composite(layer)


def render_slide(prs, slide):
    sx = WIDTH / prs.slide_width
    sy = HEIGHT / prs.slide_height
    background = (255, 255, 255, 255)
    try:
        if slide.background.fill.type == MSO_FILL_TYPE.SOLID:
            background = rgb_tuple(slide.background.fill.fore_color, background[:3]) + (255,)
    except Exception:
        pass
    canvas = Image.new("RGBA", (WIDTH, HEIGHT), background)
    for shape in slide.slide_layout.slide_master.shapes:
        if not getattr(shape, "is_placeholder", False):
            render_shape(canvas, shape, sx, sy)
    for shape in slide.slide_layout.shapes:
        if not getattr(shape, "is_placeholder", False):
            render_shape(canvas, shape, sx, sy)
    for shape in slide.shapes:
        render_shape(canvas, shape, sx, sy)
    return canvas.convert("RGB")


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for stale_preview in OUTPUT.glob("slide-[0-9][0-9].png"):
        stale_preview.unlink()
    prs = Presentation(DECK)
    previews = []
    for index, slide in enumerate(prs.slides, 1):
        preview = render_slide(prs, slide)
        destination = OUTPUT / f"slide-{index:02d}.png"
        preview.save(destination, quality=95)
        previews.append(preview)

    thumb_width = 720
    thumb_height = round(thumb_width * HEIGHT / WIDTH)
    gutter = 28
    rows = math.ceil(len(previews) / 2)
    sheet = Image.new(
        "RGB",
        (thumb_width * 2 + gutter * 3, thumb_height * rows + gutter * (rows + 1)),
        (221, 229, 234),
    )
    for index, preview in enumerate(previews):
        thumb = preview.resize((thumb_width, thumb_height), Image.Resampling.LANCZOS)
        column = index % 2
        row = index // 2
        x = gutter + column * (thumb_width + gutter)
        y = gutter + row * (thumb_height + gutter)
        sheet.paste(thumb, (x, y))
    sheet.save(OUTPUT / "contact-sheet.png", quality=95)
    print(OUTPUT / "contact-sheet.png")


if __name__ == "__main__":
    main()
