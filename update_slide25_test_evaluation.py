#!/usr/bin/env python3
"""Insert/update the approved test matrix on Page 25 without rebuilding the deck.

The first run adds one slide to the 44-page deck. Later runs replace only that
slide in the 45-page deck. Existing slides and notes are compared before the
candidate file atomically replaces the original. No cloud operations occur.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import tempfile
from zipfile import ZipFile

from pptx import Presentation
from pptx.util import Inches

import build_presentation as b

ROOT = Path(__file__).resolve().parent
DECK = ROOT / "VITA-FL_Thesis_Presentation_TU_Berlin.pptx"
ASSET = ROOT / "assets/test-evaluation-a.pptx"
TARGET = 24
RECORD = ROOT / "assets/test-evaluation-integration.json"


def visible(slide):
    return "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)


def footer(slide):
    matches = [shape for shape in slide.shapes if shape.has_text_frame
               and shape.top > Inches(6.7)
               and re.fullmatch(r"Page \d+", shape.text.strip())]
    assert len(matches) == 1, "Expected exactly one existing page footer"
    return matches[0]


def set_footer(slide, number):
    texts = footer(slide).element.xpath(".//a:t")
    assert len(texts) == 1, "Refuse to reformat a multi-run footer"
    texts[0].text = f"Page {number}"


def normalized_xml(slide):
    # Only the exact numeric page text may differ in an old slide.
    element = deepcopy(slide.element)
    for text in element.xpath(".//a:t"):
        if re.fullmatch(r"Page \d+", text.text or ""):
            text.text = "Page <number>"
    return str(element.xml)


def snapshot(slide):
    return {
        "xml": normalized_xml(slide),
        "notes": str(slide.notes_slide._element.xml),
        "relationships": sorted((rel.rId, rel.reltype, rel.is_external, rel.target_ref)
                                for rel in slide.part.rels.values()),
    }


def check_structure(prs):
    assert len(prs.slides) == 45, "Expected 28 talk pages and 17 backup pages"
    assert len(prs.slide_masters) == 2
    assert "One end-to-end Phala FedAvg run" in visible(prs.slides[23])
    assert prs.slides[TARGET].name == b.TEST_EVALUATION_SLIDE_NAME
    assert "What changed over the 24 federated rounds?" in visible(prs.slides[25])
    assert "Conclusion and outlook" in visible(prs.slides[27])
    assert "Training and inference share one verified image" in visible(prs.slides[28])
    for number, slide in enumerate(prs.slides, 1):
        assert slide.notes_slide.notes_text_frame.text.strip(), f"Empty notes: Page {number}"
        if number > 1:
            assert footer(slide).text == f"Page {number}", f"Wrong footer: Page {number}"
        ids = [node.get("id") for node in slide.element.xpath(".//p:cNvPr")]
        assert len(ids) == len(set(ids)), f"Duplicate shape IDs: Page {number}"
    slide = prs.slides[TARGET]
    assert "Variant A" not in visible(slide)
    shapes = {shape.name: shape for shape in slide.shapes}
    for domain, color in (("DFL", "207548"), ("Agent", "1764A1")):
        assert str(shapes[f"Domain navigation: {domain} badge"].fill.fore_color.rgb) == color
    for shape in slide.shapes:
        assert shape.left >= 0 and shape.top >= 0, shape.name
        assert shape.left + shape.width <= prs.slide_width, shape.name
        assert shape.top + shape.height <= prs.slide_height, shape.name


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--render", action="store_true", help="Refresh all PNG previews and the contact sheet")
    args = parser.parse_args()
    if not ASSET.is_file():
        raise FileNotFoundError(f"Approved slide asset is not ready: {ASSET}")
    prs = Presentation(DECK)
    old_count = len(prs.slides)
    inserting = old_count == 44
    assert old_count in (44, 45), "Refuse to modify an unexpected deck structure"
    assert "One end-to-end Phala FedAvg run" in visible(prs.slides[23])
    if inserting:
        assert "What changed over the 24 federated rounds?" in visible(prs.slides[24])
        assert not any(slide.name == b.TEST_EVALUATION_SLIDE_NAME for slide in prs.slides)
    else:
        check_structure(prs)
    before = [snapshot(slide) for slide in prs.slides]
    before_hash = hashlib.sha256(DECK.read_bytes()).hexdigest()
    backup = Path(tempfile.mkdtemp(prefix="vita-fl-before-test-evaluation-")) / DECK.name
    shutil.copy2(DECK, backup)

    if inserting:
        b.slide_test_evaluation(prs)
        added_id = prs.slides._sldIdLst[-1]
        prs.slides._sldIdLst.remove(added_id)
        prs.slides._sldIdLst.insert(TARGET, added_id)
        for number in range(TARGET + 1, len(prs.slides) + 1):
            set_footer(prs.slides[number - 1], number)
    else:
        b.apply_test_evaluation_asset(prs.slides[TARGET], TARGET + 1)
    check_structure(prs)

    handle, filename = tempfile.mkstemp(prefix=".test-evaluation-", suffix=".pptx", dir=ROOT)
    os.close(handle)
    candidate = Path(filename)
    try:
        prs.save(candidate)
        checked = Presentation(candidate)
        check_structure(checked)
        preserved = 0
        for index, old in enumerate(before):
            if not inserting and index == TARGET:
                continue
            new_index = index + int(inserting and index >= TARGET)
            assert snapshot(checked.slides[new_index]) == old, \
                f"Existing Page {index + 1} changed beyond its page-number footer"
            preserved += 1
        source = Presentation(ASSET).slides[0]
        assert checked.slides[TARGET].notes_slide.notes_text_frame.text == source.notes_slide.notes_text_frame.text
        with ZipFile(backup) as old_zip, ZipFile(candidate) as new_zip:
            assert new_zip.testzip() is None
            for name in old_zip.namelist():
                if name.startswith(("ppt/media/", "ppt/slideMasters/", "ppt/slideLayouts/", "ppt/notesMasters/")):
                    assert new_zip.read(name) == old_zip.read(name), name

        # Independently exercise the generator hook at its real Page 25 position.
        generated = b.prepare_template()
        for _ in range(TARGET):
            generated.slides.add_slide(generated.slide_masters[1].slide_layouts[0])
        b.slide_test_evaluation(generated)
        assert visible(generated.slides[TARGET]) == visible(checked.slides[TARGET])
        assert len(generated.slides[TARGET].shapes) == len(checked.slides[TARGET].shapes)
        assert generated.slides[TARGET].notes_slide.notes_text_frame.text == source.notes_slide.notes_text_frame.text
        os.replace(candidate, DECK)
    finally:
        candidate.unlink(missing_ok=True)

    record = {
        "page": 25, "slides_before": old_count, "slides_after": 45,
        "talk_pages": 28, "backup_pages": 17,
        "mode": "insert" if inserting else "replace",
        "preserved_existing_slides_and_notes": preserved,
        "allowed_existing_change": "page-number footer only after inserted Page 25",
        "main_deck_sha256_before": before_hash,
        "main_deck_sha256_after": hashlib.sha256(DECK.read_bytes()).hexdigest(),
        "asset_sha256": hashlib.sha256(ASSET.read_bytes()).hexdigest(),
        "backup": str(backup), "structure_check": "passed", "generator_import": "passed",
        "media_masters_and_layouts_unchanged": True,
        "source_notes_preserved": True, "native_editable_shapes": True,
        "dfl_chip": "active", "agent_chip": "active", "visual_review": "pending",
    }
    RECORD.write_text(json.dumps(record, indent=2) + "\n")
    if args.render:
        import render_preview
        render_preview.main()
    print(f'{"Inserted" if inserting else "Updated"} Page 25 in {DECK.name}; 45 pages (28 talk, 17 backup).')
    print(f"Preserved {preserved} existing slides and their notes; backup: {backup}")
    print(f"Checks: {RECORD}")


if __name__ == "__main__":
    main()
