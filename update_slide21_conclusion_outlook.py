#!/usr/bin/env python3
"""Make slide 21's conclusion and outlook specific to thesis Chapter 8."""

from __future__ import annotations

import os
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from build_presentation import (
    DARK,
    GREEN,
    ORANGE,
    TU_RED,
    WHITE,
    add_bullets,
    add_note,
    add_text,
    rgb,
)


ROOT = Path(__file__).resolve().parent
DECK = ROOT / "VITA-FL_Thesis_Presentation_TU_Berlin.pptx"
BACKUP = Path("/tmp/VITA-FL_Thesis_Presentation_TU_Berlin.before-slide21.pptx")
TEMP = ROOT / ".VITA-FL_Thesis_Presentation_TU_Berlin.slide21.tmp.pptx"


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def replace_run_text(shape, text):
    paragraph = shape.text_frame.paragraphs[0]
    assert paragraph.runs
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def main():
    shutil.copy2(DECK, BACKUP)
    prs = Presentation(DECK)
    assert len(prs.slides) >= 21
    slide = prs.slides[20]

    by_text = {
        shape.text.strip(): shape
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    }
    assert "Conclusion and outlook" in by_text
    assert "TRUST-MINIMIZED—NOT TRUSTLESS" in by_text
    assert "DEMONSTRATED" in by_text
    assert "REMAINING BOUNDARIES" in by_text
    assert "CORE CONTRIBUTION" in by_text

    replace_run_text(
        by_text["TRUST-MINIMIZED—NOT TRUSTLESS"],
        "COMPOSED EVIDENCE WORKS · NEXT: STRENGTHEN SOURCE-TO-DEPLOYMENT TRUST",
    )
    replace_run_text(by_text["DEMONSTRATED"], "CONCLUSION · DEMONSTRATED")
    replace_run_text(by_text["REMAINING BOUNDARIES"], "CONCRETE NEXT STEPS")
    demonstrated_heading = by_text["DEMONSTRATED"]
    demonstrated_heading.width = Inches(4.50)
    demonstrated_heading.height = Inches(0.32)

    old_bullets = [
        shape
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
        and shape.text.strip().startswith("●")
    ]
    assert len(old_bullets) == 2, len(old_bullets)
    for shape in old_bullets:
        remove_shape(shape)

    add_bullets(
        slide,
        [
            "End-to-end chain: ledger-selected model → encrypted handoff → TDX/AIR result → three tool receipts → SCITT record",
            "Deterministic trust path: selection, decryption, validation, and inference remain outside the language model",
            "Fail-closed bindings: identities, freshness, artifacts, inputs, and outputs are checked across component boundaries",
        ],
        0.98,
        2.25,
        5.13,
        2.62,
        10.8,
        DARK,
        GREEN,
        6,
    )
    add_bullets(
        slide,
        [
            "Supply chain: signed source → build → image provenance, reproducible or attestable builds, verified SBOMs, image signatures, and vulnerability gates",
            "Independent logs: replicated SCITT/CCF operators, witnessed or gossiped checkpoints, external anchoring, and split-view tests",
            "Assurance at scale: robust aggregation and poisoning tests, native DICOM provenance, separated inference keys, and repeated runs",
        ],
        7.10,
        2.25,
        5.12,
        2.62,
        10.4,
        DARK,
        ORANGE,
        6,
    )

    remove_shape(by_text["CORE CONTRIBUTION"])
    old_core = next(
        shape
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
        and shape.text.strip().startswith("Verifiability emerges")
    )
    remove_shape(old_core)
    add_text(slide, "CORE CONCLUSION", 1.30, 5.61, 1.75, 0.24, 10.5, WHITE, True)
    add_text(
        slide,
        "VITA-FL composes scoped evidence from decentralized model production to agent consumption—"
        "a reproducible proof of concept, not a trustless or clinically validated platform.",
        3.12,
        5.48,
        8.78,
        0.42,
        10.8,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    add_note(
        slide,
        "The conclusion is specific to the evaluated systems path. VITA-FL connects the ledger-selected "
        "model, encrypted and signed handoff, TDX/AIR inference evidence, three receiver-signed tool "
        "receipts, and SCITT recording without making the conversational model a root of trust. Exact "
        "identities, freshness values, artifacts, inputs, and outputs are checked across component "
        "boundaries and mismatches fail closed. The prioritized outlook follows Chapter 8. First, the "
        "image digest needs a complete signed source-to-build-to-deployment chain through reproducible or "
        "attestable builds, signed provenance, verified software bills of materials, image signatures, "
        "vulnerability gates, and digest-pinned automation. Second, transparency should move to independent "
        "operators with durable replication, witnessed or gossiped checkpoints, external anchoring, and "
        "split-view testing. Third, assurance should expand to robust aggregation and poisoning tests, native "
        "DICOM provenance, purpose-specific inference identities with attested key delegation, larger and "
        "batched inference, and repeated failure-injection runs. These are future-work items, not demonstrated "
        "properties. The result remains a reproducible proof of concept rather than a production-ready or "
        "clinically validated platform.",
    )

    thank_you = by_text.get("Thank you · Questions?")
    if thank_you is not None:
        thank_you.left = Inches(10.03)
        thank_you.top = Inches(6.15)
        thank_you.width = Inches(2.18)
        thank_you.height = Inches(0.22)
        paragraph = thank_you.text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.RIGHT
        for run in paragraph.runs:
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = rgb(TU_RED)

    prs.save(TEMP)

    check = Presentation(TEMP)
    check_slide = check.slides[20]
    visible = "\n".join(
        shape.text for shape in check_slide.shapes if getattr(shape, "has_text_frame", False)
    )
    for required in (
        "CONCLUSION · DEMONSTRATED",
        "CONCRETE NEXT STEPS",
        "verified SBOMs",
        "signed source → build → image provenance",
        "Independent logs",
        "CORE CONCLUSION",
        "reproducible proof of concept",
    ):
        assert required in visible, required
    for removed in ("REMAINING BOUNDARIES", "Native DICOM data, larger models"):
        assert removed not in visible, removed

    os.replace(TEMP, DECK)
    print(DECK)
    print(f"backup: {BACKUP}")


if __name__ == "__main__":
    main()
