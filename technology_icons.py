"""Editable technology symbols shared by the introduction and architecture slides."""

from pathlib import Path
import xml.etree.ElementTree as ET

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from render_preview import _svg_path_polygons


ASSETS = Path(__file__).resolve().parent / "assets" / "technology-icons"
PREFIX = "Page 2 technology icon: "
PROJECTS = (
    (5.42, ("blockchain", "encryption"), "2F8F46", "EAF6ED"),
    (7.63, ("zk", "tee", "attestation", "dfl"), "1F90CC", "EAF5FA"),
    (9.83, ("agent", "mcp"), "FF6C00", "FFF1E7"),
)


def add_icon(slide, name, x, y, size, color, background, *, prefix=PREFIX, mcp_color="000000"):
    """Draw on a 32-unit square using native lines, ovals and rectangles."""
    start = len(slide.shapes)
    scale = size / 32
    stroke_scale = size / .36

    def line(points, width=1.15):
        for a, b in zip(points, points[1:]):
            s = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(x + a[0] * scale), Inches(y + a[1] * scale),
                Inches(x + b[0] * scale), Inches(y + b[1] * scale),
            )
            s.line.color.rgb = RGBColor.from_string(color)
            s.line.width = Pt(width * stroke_scale)
            s._element.spPr.get_or_add_ln().set("cap", "rnd")

    def path(data, width=1.15):
        for points in _svg_path_polygons(data, curve_steps=12):
            line(points, width)

    def shape(kind, px, py, w, h, filled=False, width=1.15):
        s = slide.shapes.add_shape(
            kind, Inches(x + px * scale), Inches(y + py * scale),
            Inches(w * scale), Inches(h * scale),
        )
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor.from_string(color if filled else background)
        s.line.color.rgb = RGBColor.from_string(color)
        s.line.width = Pt(width * stroke_scale)
        return s

    def rect(px, py, w, h, filled=False):
        return shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, px, py, w, h, filled)

    def circle(px, py, diameter, filled=False):
        return shape(MSO_AUTO_SHAPE_TYPE.OVAL, px, py, diameter, diameter, filled)

    if name == "blockchain":
        line([(9, 9), (22, 9), (22, 23), (9, 23), (9, 9)])
        for cx, cy in ((8, 8), (24, 8), (8, 24), (24, 24)):
            rect(cx - 4, cy - 4, 8, 8)
    elif name == "encryption":
        path("M9 15 L9 10 C9 0 23 0 23 10 L23 15")
        rect(5, 13, 22, 16)
        circle(14.5, 18, 3, True)
        line([(16, 21), (16, 24)])
    elif name == "zk":
        circle(2, 2, 28)
        s = slide.shapes.add_textbox(
            Inches(x + 4 * scale), Inches(y + 6 * scale),
            Inches(24 * scale), Inches(20 * scale),
        )
        f = s.text_frame
        f.margin_left = f.margin_right = f.margin_top = f.margin_bottom = 0
        f.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = f.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = "ZK"
        r.font.name = "Arial"
        r.font.size = Pt(8 * stroke_scale)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(color)
    elif name == "tee":
        for p in (10, 16, 22):
            line([(p, 2), (p, 6)])
            line([(p, 26), (p, 30)])
            line([(2, p), (6, p)])
            line([(26, p), (30, p)])
        rect(6, 6, 20, 20)
        path("M12 15 L12 12 C12 7 20 7 20 12 L20 15", .9)
        rect(11, 14, 10, 8)
    elif name == "attestation":
        path("M20 3 L6 3 L6 29 L26 29 L26 9 L20 3 L20 9 L26 9")
        line([(10, 12), (17, 12)], .9)
        line([(10, 16), (15, 16)], .9)
        line([(11, 22), (15, 26), (24, 17)], 1.5)
    elif name == "dfl":
        # Equal peers with mesh connections; no privileged central server.
        peers = ((16, 5), (28, 14), (23, 27), (9, 27), (4, 14))
        line(list(peers) + [peers[0]])
        line([peers[0], peers[2], peers[4], peers[1], peers[3], peers[0]], .8)
        for cx, cy in peers:
            circle(cx - 2.6, cy - 2.6, 5.2, True)
    elif name == "agent":
        line([(16, 4), (16, 9)])
        circle(14.5, 1.5, 3, True)
        rect(4, 9, 24, 20)
        line([(1, 15), (1, 22)])
        line([(31, 15), (31, 22)])
        circle(9, 15, 3, True)
        circle(20, 15, 3, True)
        line([(11, 24), (21, 24)])
    elif name == "mcp":
        # Preserve the official geometry; use a light mark on dark backgrounds.
        root = ET.parse(ASSETS / "mcp.svg").getroot()
        color = mcp_color
        for element in root.iter():
            if element.tag.endswith("path") and element.get("d"):
                for points in _svg_path_polygons(element.get("d"), curve_steps=16):
                    line([(px * 32 / 180, py * 32 / 180) for px, py in points], 1.55)
    else:
        raise ValueError(name)

    for index, s in enumerate(list(slide.shapes)[start:]):
        s.name = f"{prefix}{name} / {index + 1}"


def add_project_technology_icons(slide):
    for x, names, color, fill in PROJECTS:
        centers = (.71, 1.42) if len(names) == 2 else (.37, .83, 1.29, 1.75)
        for center, name in zip(centers, names):
            add_icon(slide, name, x + center - .18, 4.18, .36, color, fill)


def add_architecture_technology_icons(slide):
    """Eight familiar symbols, without cards or captions, on the dark panel."""
    rows = (
        ("blockchain", "encryption", "83D49B"),
        ("zk", "tee", "86CEEF"),
        ("attestation", "dfl", "86CEEF"),
        ("agent", "mcp", "FFB47D"),
    )
    for row, (left, right, color) in enumerate(rows):
        for center, name in zip((10.48, 11.77), (left, right)):
            add_icon(
                slide, name, center - .33, 2.22 + row * .89, .66, color, "434343",
                prefix="Page 4 technology icon: ", mcp_color="FFFFFF",
            )
    add_architecture_connections(slide)


def add_architecture_connections(slide):
    """Connect adjacent symbols through the free gaps in the two-column grid."""
    connections = []
    for row in range(4):
        cy = 2.55 + row * .89
        connections.append((10.88, cy, 11.37, cy))
        if row < 3:
            for cx in (10.48, 11.77):
                connections.append((cx, cy + .34, cx, cy + .89 - .34))
    for index, (x1, y1, x2, y2) in enumerate(connections, start=1):
        s = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2),
        )
        s.line.color.rgb = RGBColor.from_string("C1C8CC")
        s.line.width = Pt(1.1)
        end = OxmlElement("a:tailEnd")
        end.set("type", "triangle")
        end.set("w", "sm")
        end.set("len", "sm")
        s._element.spPr.get_or_add_ln().append(end)
        s.name = f"Page 4 symbol connection: {index}"


def add_research_gap_sequence(slide):
    """The requested seven-symbol sequence above slide 5's compact RQ cards."""
    sequence = (
        ("encryption", "2F8F46"),  # Security: the existing padlock symbol.
        ("tee", "1F90CC"),
        ("attestation", "1F90CC"),
        ("blockchain", "2F8F46"),
        ("dfl", "1F90CC"),
        ("mcp", "000000"),
        ("agent", "FF6C00"),
    )
    for index, (name, color) in enumerate(sequence):
        cx = 1.27 + index * 1.80
        add_icon(
            slide, name, cx - .25, 3.92, .50, color, "FFFFFF",
            prefix="Page 5 technology icon: ",
        )
        if index < len(sequence) - 1:
            s = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(cx + .35), Inches(4.17),
                Inches(cx + 1.80 - .35), Inches(4.17),
            )
            s.line.color.rgb = RGBColor.from_string("87949C")
            s.line.width = Pt(1.2)
            end = OxmlElement("a:tailEnd")
            end.set("type", "triangle")
            end.set("w", "sm")
            end.set("len", "sm")
            s._element.spPr.get_or_add_ln().append(end)
            s.name = f"Page 5 symbol connection: {index + 1}"
            if index >= 4:
                s._element.spPr.xfrm.flipH = True
