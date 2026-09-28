#!/usr/bin/env python3
"""Insert/update approved evaluation architecture after the attack trees, if present.

Locates the section after the optional Sello process and attack trees without rebuilding.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import shutil
import tempfile
from zipfile import ZipFile

from pptx import Presentation

import build_presentation as b
from update_slide25_test_evaluation import (
    attack_tree_count, chain_of_custody_index, evaluation_section_target, footer, sello_process_count,
    set_footer, snapshot, visible,
)

ROOT = Path(__file__).resolve().parent
DECK = ROOT / 'VITA-FL_Thesis_Presentation_TU_Berlin.pptx'
ASSET = ROOT / 'assets/evaluation-stack-c.pptx'
RECORD = ROOT / 'assets/evaluation-stack-integration.json'


def evaluation_stack_target(prs):
    return evaluation_section_target(prs)


def check_structure(prs):
    tree_count = attack_tree_count(prs)
    process_count = sello_process_count(prs)
    target = evaluation_stack_target(prs)
    assert len(prs.slides) == 46 + tree_count + process_count, \
        'Unexpected talk/backup page count for the evaluation and attack-tree layout'
    assert len(prs.slide_masters) == 2
    assert prs.slides[target].name == b.EVALUATION_STACK_SLIDE_NAME
    assert sum(s.name == b.EVALUATION_STACK_SLIDE_NAME for s in prs.slides) == 1
    assert 'Cryptographic chain of custody' in visible(prs.slides[chain_of_custody_index(prs)])
    assert 'One end-to-end Phala FedAvg run' in visible(prs.slides[target + 1])
    assert prs.slides[target + 2].name == b.TEST_EVALUATION_SLIDE_NAME
    assert 'What changed over the 24 federated rounds?' in visible(prs.slides[target + 3])
    assert 'Conclusion and outlook' in visible(prs.slides[target + 5])
    assert 'Training and inference share one verified image' in visible(prs.slides[target + 6])
    for number, slide in enumerate(prs.slides, 1):
        assert slide.notes_slide.notes_text_frame.text.strip(), f'Empty notes: Page {number}'
        if number > 1:
            assert footer(slide).text == f'Page {number}', f'Wrong footer: Page {number}'
        ids = [node.get('id') for node in slide.element.xpath('.//p:cNvPr')]
        assert len(ids) == len(set(ids)), f'Duplicate shape IDs: Page {number}'
    slide = prs.slides[target]
    assert 'Variant C' not in visible(slide)
    for name in ('SMEW', 'SMA', 'vita-fl-td', 'VITA-FL', 'Prometheus', 'Marimo'):
        assert name in visible(slide), name
    shapes = {shape.name: shape for shape in slide.shapes}
    for domain, color in (('DFL', '207548'), ('Agent', '1764A1')):
        assert str(shapes[f'Domain navigation: {domain} badge'].fill.fore_color.rgb) == color
    for shape in slide.shapes:
        assert shape.left >= 0 and shape.top >= 0, shape.name
        assert shape.left + shape.width <= prs.slide_width, shape.name
        assert shape.top + shape.height <= prs.slide_height, shape.name
        assert not shape.element.xpath('.//a:blip'), 'Slide must remain native and editable'


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--render', action='store_true', help='Refresh previews and contact sheet')
    args = parser.parse_args()
    assert ASSET.is_file(), f'Missing approved asset: {ASSET}'
    prs = Presentation(DECK)
    old_count = len(prs.slides)
    tree_count = attack_tree_count(prs)
    process_count = sello_process_count(prs)
    target = evaluation_stack_target(prs)
    inserting = old_count == 45 + tree_count + process_count
    assert old_count in (45 + tree_count + process_count, 46 + tree_count + process_count), \
        'Unexpected deck structure; refusing to overwrite it'
    if inserting:
        assert 'One end-to-end Phala FedAvg run' in visible(prs.slides[target])
        assert prs.slides[target + 1].name == b.TEST_EVALUATION_SLIDE_NAME
        assert not any(s.name == b.EVALUATION_STACK_SLIDE_NAME for s in prs.slides)
    else:
        check_structure(prs)
    before = [snapshot(slide) for slide in prs.slides]
    before_hash = hashlib.sha256(DECK.read_bytes()).hexdigest()
    backup = Path(tempfile.mkdtemp(prefix='vita-fl-before-evaluation-stack-')) / DECK.name
    shutil.copy2(DECK, backup)
    if inserting:
        b.slide_evaluation_stack(prs)
        slide_id = prs.slides._sldIdLst[-1]
        prs.slides._sldIdLst.remove(slide_id)
        prs.slides._sldIdLst.insert(target, slide_id)
        for number in range(target + 1, len(prs.slides) + 1):
            set_footer(prs.slides[number - 1], number)
    else:
        b.apply_evaluation_stack_asset(prs.slides[target], target + 1)
    check_structure(prs)

    handle, filename = tempfile.mkstemp(prefix='.evaluation-stack-', suffix='.pptx', dir=ROOT)
    os.close(handle)
    candidate = Path(filename)
    try:
        prs.save(candidate)
        checked = Presentation(candidate)
        check_structure(checked)
        preserved = 0
        for index, old in enumerate(before):
            if not inserting and index == target:
                continue
            new_index = index + int(inserting and index >= target)
            assert snapshot(checked.slides[new_index]) == old, \
                f'Existing Page {index+1} changed beyond its page-number footer'
            preserved += 1
        source = Presentation(ASSET).slides[0]
        assert checked.slides[target].notes_slide.notes_text_frame.text == source.notes_slide.notes_text_frame.text
        with ZipFile(backup) as old_zip, ZipFile(candidate) as new_zip:
            assert new_zip.testzip() is None
            for name in old_zip.namelist():
                if name.startswith(('ppt/media/', 'ppt/slideMasters/', 'ppt/slideLayouts/', 'ppt/notesMasters/')):
                    assert new_zip.read(name) == old_zip.read(name), name
        generated = b.prepare_template()
        for _ in range(target):
            generated.slides.add_slide(generated.slide_masters[1].slide_layouts[0])
        b.slide_evaluation_stack(generated)
        assert visible(generated.slides[target]) == visible(checked.slides[target])
        assert len(generated.slides[target].shapes) == len(checked.slides[target].shapes)
        assert generated.slides[target].notes_slide.notes_text_frame.text == source.notes_slide.notes_text_frame.text
        os.replace(candidate, DECK)
    finally:
        candidate.unlink(missing_ok=True)

    RECORD.write_text(json.dumps({
        'page': target + 1, 'slides_before': old_count, 'slides_after': len(checked.slides),
        'talk_pages': 29 + tree_count + process_count, 'backup_pages': 17,
        'attack_tree_pages': tree_count, 'sello_process_pages': process_count,
        'mode': 'insert' if inserting else 'replace',
        'preserved_existing_slides_and_notes': preserved,
        'allowed_existing_change': (f'page-number footers only, starting at inserted Page {target + 1}'
                                    if inserting else 'none outside the replaced evaluation architecture'),
        'main_deck_sha256_before': before_hash,
        'main_deck_sha256_after': hashlib.sha256(DECK.read_bytes()).hexdigest(),
        'asset_sha256': hashlib.sha256(ASSET.read_bytes()).hexdigest(),
        'backup': str(backup), 'structure_check': 'passed', 'generator_import': 'passed',
        'media_masters_and_layouts_unchanged': True,
        'source_notes_preserved': True, 'native_editable_shapes': True,
        'visual_review': 'pending',
    }, indent=2) + '\n')
    if args.render:
        import render_preview
        render_preview.main()
    print(f'{"Inserted" if inserting else "Updated"} Page {target + 1}; '
          f'{len(checked.slides)} pages ({29 + tree_count + process_count} talk, 17 backup).')
    print(f'Preserved {preserved} existing slides and notes. Backup: {backup}')
    print(f'Checks: {RECORD}')


if __name__ == '__main__':
    main()
