#!/usr/bin/env python3
"""Build separate, editable color/navigation proposals; leave the talk deck untouched."""
from __future__ import annotations

import csv
import html
import json
import sys
from collections import Counter
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx.enum.text import PP_ALIGN

HERE = Path(__file__).resolve().parent
PRESENTATION = HERE.parents[1]
WORKSPACE = PRESENTATION.parent
sys.path.insert(0, str(PRESENTATION))
import build_presentation as b
import render_preview as renderer

DFL, AGENT, INFRA = "207548", "1764A1", "7446A6"
TINT = {DFL: "EDF6F0", AGENT: "EDF4FA", INFRA: "F3EEF8"}
INK, MUTED, LINE, PALE = "263440", "5D6973", "D7DEE3", "F3F5F6"
DOMAINS = [("DFL", DFL), ("Agent / MCP / Inference / Log", AGENT), ("Blockchain / IPFS", INFRA)]
PAGES: list[dict] = []


def text(s, value, x, y, w, h, size=15, color=INK, bold=False, center=False):
    return b.add_text(s, value, x, y, w, h, size, color, bold,
                      align=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT, margin=0)


def box(s, x, y, w, h, color=None, fill=None, width=1.1):
    return b.add_box(s, x, y, w, h, fill or (TINT[color] if color else PALE),
                     color or LINE, radius=False, line_width=width)


def line(s, x1, y1, x2, y2, color=MUTED, arrow=False, dashed=False, width=1.5):
    return b.add_line_segment(s, x1, y1, x2, y2, color, width, arrow, dashed)


def nav(s, active, mode="chips", x=.65, y=1.02, scale=1):
    widths = [1.08, 4.02, 2.62]
    for (label, c), w in zip(DOMAINS, widths):
        on = c in active
        if mode == "chips":
            box(s, x, y, w*scale, .29*scale, c if on else None,
                c if on else "F5F6F7", width=.6)
            text(s, label, x+.08*scale, y+.063*scale, (w-.16)*scale, .19*scale,
                 9.5*scale, "FFFFFF" if on else MUTED, True)
        elif mode == "rail":
            text(s, label, x, y, w*scale, .21*scale, 9.4*scale, c if on else MUTED, on)
            box(s, x, y+.28*scale, (w-.12)*scale, .045*scale, c if on else None, fill=c if on else LINE, width=.2)
        x += (w+.12)*scale


def page(prs, slug, title, purpose, active, note, mode="chips"):
    s = prs.slides.add_slide(prs.slide_masters[1].slide_layouts[0])
    b.remove_slide_placeholders(s)
    text(s, title, .62, .33, 10.75, .57, 23 if len(title)<55 else 21, INK)
    nav(s, active, mode)
    b.add_divider(s, .62, 1.43, 11.95, LINE, .012)
    text(s, f"Concept {len(PAGES)+1:02d}", .62, 6.89, 1.3, .22, 9, "FFFFFF")
    text(s, "VITA-FL  |  Design proposals  |  Ramon Mehrpoya", 2.0, 6.89, 8.8, .22, 9, "FFFFFF")
    text(s, "06 September 2026", 10.92, 6.89, 1.81, .22, 8.3, "FFFFFF")
    b.add_note(s, purpose+"\n\n"+note)
    PAGES.append(dict(slug=slug, title=title, purpose=purpose, note=note))
    return s


def takeaway(s, value, color=INK, y=5.91, size=15):
    b.add_divider(s, .66, y-.12, 11.94, LINE, .01)
    text(s, value, .72, y, 11.78, .33, size, color, True, True)


def label(s, value, x, y, w, color):
    text(s, value, x, y, w, .23, 10.5, color, True)


def card(s, heading, detail, x, y, w, h, color, size=14):
    box(s, x, y, w, h, color, "FFFFFF")
    text(s, heading, x+.14, y+.15, w-.28, .37, size, color, True, True)
    text(s, detail, x+.14, y+.62, w-.28, h-.72, 12, INK, False, True)


def palette(prs):
    s = page(prs, "01-color-system", "One color per domain. One map throughout the talk.",
             "Farbsystem: A ist die empfohlene Navigation, B und C sind Alternativen.",
             {DFL, AGENT, INFRA},
             "Colors describe functional domains, not trust levels or deployment boundaries. "
             "TU red remains the existing brand color. Use explicit STOP/error labels for failures. "
             "The transparency log belongs to the blue evidence path; blockchain and IPFS share purple but retain separate roles.")
    specs = [(DFL, "DFL", "Produce the model", "Signed inputs\nWorker TEEs / FedAvg"),
             (AGENT, "AGENT / MCP / INFERENCE / LOG", "Use it and provide evidence", "Bounded tools / Receiver TEE\nAIR / receipts / transparency"),
             (INFRA, "BLOCKCHAIN / IPFS", "Coordinate and store", "Finalized model references\nEncrypted model artifacts")]
    for i, (c, name, heading, desc) in enumerate(specs):
        x=.66+i*4.05
        box(s,x,1.72,3.9,2.0,c)
        box(s,x,1.72,3.9,.075,fill=c,width=0)
        label(s,name,x+.19,1.94,3.52,c)
        text(s,heading,x+.19,2.35,3.52,.39,16,c,True)
        text(s,desc,x+.19,2.91,3.52,.6,12.5)
        text(s,"#"+c,x+.19,3.48,3.3,.17,8.2,MUTED)
    label(s,"A  PERSISTENT CHIPS  /  RECOMMENDED",.69,4.02,5.7,INK)
    nav(s,{AGENT,INFRA},x=.69,y=4.4,scale=.90)
    text(s,"Activate both domains at a handoff.",8.4,4.42,4.15,.29,12,MUTED)
    label(s,"B  UNDERLINED SECTIONS",.69,5.05,5.7,INK)
    nav(s,{AGENT},mode="rail",x=.69,y=5.43,scale=.66)
    label(s,"C  REPEATED MINI-MAP",8.4,5.05,4.1,INK)
    for x,w,name,c in [(8.4,1.12,"DFL",DFL),(9.67,2.83,"Agent / Inference / Log",AGENT)]:
        box(s,x,5.42,w,.28,c,c if c==AGENT else "FFFFFF")
        text(s,name,x+.04,5.48,w-.08,.15,8,"FFFFFF" if c==AGENT else c,True,True)
    box(s,8.4,5.77,4.1,.26,INFRA,"FFFFFF")
    text(s,"Blockchain / IPFS",8.44,5.82,4.02,.17,8,INFRA,True,True)


def architecture(prs):
    s=page(prs,"02-architecture","Two responsibility blocks, shared infrastructure",
           "Ersatz für Folie 6. Die drei Farben bezeichnen Funktionsbereiche; die zwei Verantwortungsketten bleiben erhalten.",
           {DFL,AGENT,INFRA},
           "DFL computes the model. Contracts coordinate and finalize references, not tensor arithmetic. "
           "The receiver independently resolves the authoritative bundle, fetches its artifacts, and verifies them. "
           "The log is part of the blue evidence domain, not the blockchain. The layout is a functional map, not a process/deployment isolation claim. "
           "Sources: vita-fl/dfl/neural_network/cli.py:1276; smart_contracts/src/core/GMStorage.sol:180; "
           "tee_inference/service/model_source.py:218; tee_inference/service/app.py:118.")
    for x,w,c,title,desc in [(.66,5.78,DFL,"01  DFL","Produces the global model"),(6.71,5.95,AGENT,"02  AGENT / INFERENCE","Uses the model and records evidence")]:
        box(s,x,1.72,w,2.3,c)
        label(s,title,x+.2,1.93,w-.4,c)
        text(s,desc,x+.2,2.29,w-.4,.36,16,c,True)
    for i,(title,detail) in enumerate([("Signed inputs","Medical data"),("Worker TEEs","Local training"),("FedAvg","One aggregator")]):
        x=.86+i*1.82
        card(s,title,detail,x,2.91,1.58,.87,DFL,12)
        # Compact row uses a separate single-line body.
        last=s.shapes[-1]; last.top=b.Inches(3.43); last.height=b.Inches(.22); last.text_frame.paragraphs[0].runs[0].font.size=b.Pt(10)
        if i<2: line(s,x+1.60,3.32,x+1.80,3.32,DFL,True)
    for i,(title,detail) in enumerate([("Agent / MCP","Orchestrates"),("Receiver TEE","Verifies / infers"),("Log","Records receipts")]):
        x=6.91+i*1.87
        card(s,title,detail,x,2.91,1.63,.87,AGENT,12)
        last=s.shapes[-1]; last.top=b.Inches(3.43); last.height=b.Inches(.22); last.text_frame.paragraphs[0].runs[0].font.size=b.Pt(9.8)
        if i<2: line(s,x+1.65,3.32,x+1.85,3.32,AGENT,True)
    line(s,3.3,4.03,3.3,4.63,DFL,True)
    text(s,"Publish model + references",.93,4.13,2.18,.43,10.1,DFL,True)
    line(s,9.57,4.63,9.57,4.03,INFRA,True)
    text(s,"Resolve + retrieve",9.78,4.20,2.6,.27,11,INFRA,True)
    box(s,.66,4.64,12,1.04,INFRA)
    label(s,"SHARED INFRASTRUCTURE",.9,4.86,2.76,INFRA)
    text(s,"Blockchain",3.80,4.82,3.2,.3,16,INFRA,True)
    text(s,"Finalized state / round / publisher",3.8,5.24,4.13,.22,11)
    line(s,8.03,4.85,8.03,5.47,INFRA,width=.8)
    text(s,"IPFS",8.35,4.82,3.5,.3,16,INFRA,True)
    text(s,"Encrypted model / signature / keys",8.35,5.24,4.1,.22,11)
    takeaway(s,"Finalized references and encrypted artifacts connect model production to verified use.",size=14)


def handoff(prs):
    s=page(prs,"03-model-handoff","The receiver resolves the model independently",
           "Ersatz für Folie 14. Blau und Violett sind gleichzeitig aktiv; Ledger und IPFS bleiben getrennte Rollen.",
           {AGENT,INFRA},
           "The finalized bundle contains the historical publisher-key snapshot. The request does not choose a CID, "
           "round, publisher, or local model path. Sources: agent/blockchain_source.py:357,561,656; "
           "tee_inference/service/model_source.py:218.")
    box(s,.67,1.76,3.55,.65,AGENT)
    text(s,"Agent: prepare current model",.86,1.95,3.17,.29,13.2,AGENT,True)
    line(s,4.22,2.08,5.10,2.08,AGENT,True)
    text(s,"No caller-selected CID, round, publisher or path",5.28,1.96,7.12,.30,14,AGENT,True)
    for y,title,body in [(2.82,"BLOCKCHAIN / AUTHORITY","Finalized round + CIDs\nPublisher-key snapshot"),
                         (4.25,"IPFS / ARTIFACTS","Encrypted model + signature\nRecipient key bundle")]:
        box(s,.67,y,3.55,1.15,INFRA)
        label(s,title,.88,y+.18,3.15,INFRA)
        text(s,body,.88,y+.57,3.15,.49,12.4)
        line(s,4.23,y+.59,5.10,y+.59,INFRA,True)
    box(s,5.12,2.82,7.54,2.58,AGENT)
    label(s,"MEASURED RECEIVER",5.35,3.02,6.9,AGENT)
    for i,(title,body) in enumerate([("1  RESOLVE","Finalized ledger\nbundle"),("2  RETRIEVE","Named artifacts\nUnwrap / decrypt"),("3  VERIFY","Hashes / round\nPublisher key")]):
        x=5.35+i*2.35
        text(s,title,x,3.48,2.13,.34,15,AGENT,True)
        text(s,body,x,4.01,2.10,.59,13)
    box(s,5.35,4.83,7.08,.36,AGENT,AGENT)
    text(s,"Install the exact model — or reject",5.49,4.90,6.80,.23,12.5,"FFFFFF",True,True)
    takeaway(s,"Blockchain identifies what is current. IPFS supplies bytes. The receiver checks both.",size=14)


def sequence(prs):
    s=page(prs,"04-inference-sequence","Three TEE tool calls, four log records",
           "Ersatz für Folie 16. Eine Bereichsfarbe, klar beschriftete Akteure und Pfeile. Gezeigt ist der TEE-Pfad mit Sello.",
           {AGENT},
           "Shown: receipt-enabled TEE profile. Six public MCP tools exist in total: three for TEE and three for ZK. "
           "Each receiver action publishes its receipt before returning success. The agent validates and registers the "
           "complete inference evidence bundle as a fourth record. Sello is configuration-dependent; SELLO_REQUIRED enforces it. "
           "Sources: agent/mcp_server.py:45,316; tee_inference/service/app.py:118; agent/tee_inference_client.py:691; "
           "agent_receipts/environment.py:30.", mode="rail")
    centers=[1.94,6.60,11.21]
    for x,title,sub in zip(centers,["AGENT / MCP","RECEIVER TEE","TRANSPARENCY LOG"],["Bounded orchestration","Model + image + inference","Receiver receipts + bundle"]):
        box(s,x-1.26,1.74,2.52,.67,AGENT)
        text(s,title,x-1.15,1.9,2.3,.22,12,AGENT,True,True)
        text(s,sub,x-1.7,2.47,3.4,.22,10.2,MUTED,False,True)
        line(s,x,2.84,x,5.35,LINE,dashed=True,width=1)
    for y,num,title,receipt in [(2.99,"1","Prepare model","Receipt 1"),(3.75,"2","Select image / create job","Receipt 2"),(4.51,"3","Infer / emit evidence","Receipt 3")]:
        text(s,f"{num}  {title}",2.19,y-.26,4.15,.24,12.5,AGENT,True)
        line(s,1.96,y,6.56,y,AGENT,True)
        text(s,receipt+" → register",6.9,y-.26,4.03,.24,11.6,AGENT,True)
        line(s,6.63,y,11.16,y,AGENT,True)
        text(s,"verify inclusion proof",7.0,y+.10,3.95,.20,9.6,MUTED)
        line(s,11.16,y+.34,6.63,y+.34,MUTED,True,True,.8)
        text(s,"then return success",2.19,y+.10,3.95,.20,9.6,MUTED)
        line(s,6.56,y+.34,1.96,y+.34,MUTED,True,True,.8)
    text(s,"4  Agent verifies and registers the complete inference bundle",2.19,5.18,8.62,.28,12.4,AGENT,True)
    line(s,1.96,5.55,11.16,5.55,AGENT,True,width=1.8)
    takeaway(s,"Receipt before success — for every receiver action in the Sello-enabled path.",AGENT,size=14)


def failure(prs):
    s=page(prs,"05-failure-paths","What happens when a check fails?",
           "Neue Folienidee nach Folie 14 oder 16, alternativ als Backup. Zwei Designbeispiele aus implementierten Fehlerpfaden.",
           {AGENT,INFRA},
           "These are design examples derived from rejection code, not newly performed attack experiments. "
           "In scenario 1, the bytes do not match the finalized reference/signature. In scenario 2, Sello receipts are enabled. "
           "A failed log publication causes HTTP 503 and does not roll back an already executed action. "
           "Sources: agent/blockchain_source.py:656; tee_inference/service/model_source.py:218; "
           "tee_inference/service/app.py:137-149; agent_receipts/receiver_log.py:25.")
    text(s,"Illustrative code paths · receipt publication shown with Sello enabled",.70,1.65,11.85,.24,10.5,MUTED)
    for y,num,name in [(1.98,"01","An artifact is substituted"),(3.92,"02","Receipt publication fails")]:
        label(s,num+"  "+name.upper(),.70,y,11.85,INK)
    specs=[(2.34,"IPFS / bundle","Bytes do not match\nthe finalized reference",INFRA,"Receiver verification","Hash / signature check\nfails","MODEL REJECTED","No installation"),
           (4.28,"Receiver action","Action ran; receipt\npublication is attempted",AGENT,"Transparency log","Receipt cannot\nbe committed","NO SUCCESS RESPONSE","HTTP 503; action may have run")]
    for y,t1,d1,c,t2,d2,result,detail in specs:
        card(s,t1,d1,.70,y,3.37,1.27,c)
        line(s,4.12,y+.63,4.62,y+.63,MUTED,True)
        card(s,t2,d2,4.67,y,3.65,1.27,AGENT)
        line(s,8.36,y+.63,8.86,y+.63,MUTED,True)
        box(s,8.92,y,3.67,1.27,fill=PALE)
        text(s,"×",9.1,y+.12,.38,.45,24,b.TU_RED,True)
        text(s,result,9.55,y+.22,2.87,.30,10.9,INK,True)
        text(s,detail,9.13,y+.72,3.23,.40,11.5,MUTED,False,True)
    takeaway(s,"Visible rejection explains the mechanism; it does not demonstrate every threat is covered.",size=13.5)


def evidence(prs):
    s=page(prs,"06-evidence-scope","What does each piece of evidence establish?",
           "Neue Folienidee als vereinfachter Ersatz für Folie 18 oder als Backup nach Folie 17.",
           {DFL,AGENT,INFRA},
           "Evidence has scoped guarantees. Current aggregation uses equal-weight FedAvg and does not offer semantic poisoning resistance. "
           "The inference verifier checks AIR signature, hashes, REPORTDATA, RTMR3 replay and image/contract policy, but "
           "returns dcap_collateral_verified=False. This is distinct from on-chain DFL admission. The CCF log is single-node virtual mode. "
           "Sources: dfl/neural_network/cli.py:1276; agent/tee_inference_client.py:542,636; "
           "agent_receipts/scitt.py:93; transparency_log/README.md:3,34.")
    for title,x,w in [("EVIDENCE",.85,2.64),("QUESTION IT ANSWERS",3.71,4.18),("BOUNDARY",8.45,3.72)]:
        label(s,title,x,1.78,w,MUTED)
    rows=[(DFL,"DFL statements","Which accepted inputs and output\nare bound to this round?","Signed data can still\nbe semantically harmful."),
          (INFRA,"Finalized bundle","Which model reference and publisher\nare authoritative?","IPFS availability is\na separate concern."),
          (AGENT,"AIR + quote bindings","Which model, request, response\nand workload policy are linked?","The inference verifier lacks\nfull DCAP collateral appraisal."),
          (AGENT,"Tool + log receipts","Which action was signed by the receiver,\nand which statement was recorded?","The current log is a\nsingle-node virtual CCF service.")]
    for i,(c,title,q,limit) in enumerate(rows):
        y=2.16+i*.84
        box(s,.68,y,11.95,.76,fill="FFFFFF")
        box(s,.68,y,.055,.76,fill=c,width=0)
        text(s,title,.87,y+.23,2.6,.32,13.4,c,True)
        text(s,q,3.71,y+.15,4.53,.53,12.1)
        text(s,limit,8.45,y+.15,3.89,.53,11.6,MUTED)
    takeaway(s,"Technical provenance and bindings are evidence. Clinical utility needs its own evaluation.",size=13.7)


def role_heatmap(prs):
    source=WORKSPACE/"vita-fl/data/evaluation/authoritative-phala-6w-24r-20260901/round_metrics.csv"
    rows=list(csv.DictReader(source.open()))
    assert [int(r["federated_round"]) for r in rows] == list(range(1,25))
    assert all(int(r["aggregated_model_count"])==5 for r in rows)
    counts=Counter(r["aggregator_worker"] for r in rows)
    s=page(prs,"07-worker-roles","Who aggregated in the actual 24-round run?",
           "Neue empirische Folienidee nach Folie 19 oder als Backup zu Folie 10. Direkt aus dem archivierten Einzelrun erzeugt.",
           {DFL},
           "Data: vita-fl/data/evaluation/authoritative-phala-6w-24r-20260901/round_metrics.csv and worker_activity.csv. "
           "A marks the selected aggregator, T marks a training contributor; five accepted client models per round. "
           "Counts VM0-5 are 2,3,4,4,9,2. Federated rounds 1-24 publish global model rounds 2-25. "
           "This single run shows role changes, not long-run selection fairness or recovery effectiveness; zero aborted attempts were observed.")
    text(s,"Six workers · one aggregator and five training contributors per round",.71,1.79,11.8,.40,17,DFL,True)
    x0,y0,cw,ch=1.72,2.73,.387,.408
    text(s,"Federated round",x0,2.35,8,.23,10.6,MUTED)
    text(s,"Total A",11.40,2.35,1.03,.23,10.6,MUTED,True)
    for col,r in enumerate(rows):
        text(s,str(col+1),x0+col*cw,2.58,cw-.035,.19,8.5,MUTED,False,True)
    for worker in range(6):
        y=y0+worker*ch
        text(s,"VM-"+str(worker),.73,y+.11,.88,.24,12.2,DFL,True)
        for col,r in enumerate(rows):
            a=r["aggregator_worker"]==f"VM-{worker}"
            x=x0+col*cw
            box(s,x,y,cw-.035,ch-.048,DFL,DFL if a else TINT[DFL],.5)
            text(s,"A" if a else "T",x,y+.085,cw-.035,.23,11.2,"FFFFFF" if a else DFL,a,True)
        text(s,str(counts[f"VM-{worker}"]),11.51,y+.085,.8,.27,14,DFL,True,True)
    box(s,.73,5.35,.28,.27,DFL,DFL)
    text(s,"A",.73,5.41,.28,.17,8.5,"FFFFFF",True,True)
    text(s,"Aggregation",1.10,5.39,1.75,.25,11.1,DFL)
    box(s,3.01,5.35,.28,.27,DFL,TINT[DFL])
    text(s,"T",3.01,5.41,.28,.17,8.5,DFL,False,True)
    text(s,"Local training",3.39,5.39,2.3,.25,11.1,DFL)
    text(s,"24 rounds × 5 updates = 120 contributions",6.45,5.39,6.06,.27,12.6,DFL,True)
    takeaway(s,"Role rotation is visible. One run does not establish long-run fairness or recovery behavior.",size=13.4)
    return dict(source=str(source.relative_to(WORKSPACE)),aggregator_counts=dict(sorted(counts.items())),
                rounds=len(rows),updates_per_round=5)


def hospital_privacy(prs):
    s=page(prs,"08-hospital-data-control","Hospitals want to keep control of patient data",
           "Neue Motivationsfolie nach der Arztperspektive (Folie 3), vor der technischen Antwort. Krankenhäuser wollen gemeinsam lernen und sensible Daten unter eigener Kontrolle behalten.",
           {DFL},
           "This is the target hospital scenario, not a claim of deployment in real hospitals. "
           "Hospitals want to protect patient confidentiality, decide who can access their records, "
           "and avoid pooling raw patient data in one central training repository. Federated learning "
           "addresses that motivation by training where data is held and exchanging model updates. "
           "The shared global model is an artifact, not a permanent central training server: VITA-FL "
           "selects a temporary aggregator. The code trains on local dataset shards and encrypts "
           "exported local model parameters; model updates here means full local models, not only gradients. "
           "The proof of concept was evaluated with ChestMNIST on Phala. It does not establish that "
           "real patient records remained within hospital premises. Federated learning and encryption "
           "alone do not establish a formal privacy guarantee against information leakage from models. "
           "Sources: vita-fl/dfl/neural_network/cli.py:991-1004,1069-1077; "
           "vita-fl/dfl/neural_network/README.md:9-10,22-23. "
           "Suggested narration: Hospitals have a second concern beyond the quality of a prediction: "
           "they must retain control of sensitive records. They want to benefit from collaboration "
           "without creating a central pool of raw patient data. This motivates local training and "
           "protected model exchange in VITA-FL.")
    text(s,"Patient confidentiality · control over access · no central pool of raw records",
         .72,1.76,11.9,.38,15.7,DFL,True)
    for i,name in enumerate(["Hospital A","Hospital B","Hospital C"]):
        x=.70+i*4.08
        box(s,x,2.34,3.75,2.24,DFL)
        b.add_hospital_campus_icon(s,x+.49,2.86,DFL,1.02)
        text(s,name,x+.96,2.68,2.56,.38,18,DFL,True)
        box(s,x+.20,3.29,3.35,.59,DFL,"FFFFFF")
        # Native editable padlock: an outlined shackle, a filled body and a keyhole.
        b.add_oval(s,x+.39,3.38,.24,.27,"FFFFFF",DFL,1.3)
        box(s,x+.33,3.52,.36,.23,DFL,DFL,.6)
        b.add_oval(s,x+.48,3.58,.06,.06,"FFFFFF","FFFFFF",.3)
        text(s,"Local patient records",x+.84,3.44,2.38,.27,12.2,INK,True)
        text(s,"Train locally → model update",x+.20,4.10,3.35,.28,12.3,DFL,True,True)
    for cx,endpoint in [(2.575,4.74),(10.735,8.60)]:
        line(s,cx,4.59,cx,5.36,DFL)
        line(s,cx,5.36,endpoint,5.36,DFL,True)
    line(s,6.655,4.59,6.655,5.07,DFL,True)
    text(s,"Protected model updates",1.62,5.48,3.10,.25,10.8,DFL,True,True)
    text(s,"Protected model updates",8.61,5.48,3.10,.25,10.8,DFL,True,True)
    box(s,4.79,5.10,3.75,.52,DFL,DFL)
    text(s,"Shared global model",4.98,5.22,3.37,.32,15.7,"FFFFFF",True,True)
    takeaway(s,"Train where the data is held. Collaborate through protected model exchange.",DFL,size=14.2)


def gallery():
    items=[]
    for p in PAGES:
        slug,title,purpose=(html.escape(p[k]) for k in ("slug","title","purpose"))
        items.append(f'<article><button class="preview" data-src="{slug}.png" aria-label="Vorschau vergrößern: {title}"><img src="{slug}.png" alt="{title}" loading="lazy"></button><div class="caption"><h2>{title}</h2><p>{purpose}</p><a href="{slug}.png">PNG öffnen</a></div></article>')
    document='''<!doctype html>
<html lang="de"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>VITA-FL · Folienideen & Farbcodierung</title>
<style>
*{box-sizing:border-box}body{margin:0;background:#eef1f4;color:#263440;font:16px/1.6 system-ui,sans-serif}main{max-width:1460px;margin:auto;padding:44px 28px}header{max-width:1120px;margin-bottom:30px}h1{font-size:clamp(27px,4vw,42px);line-height:1.2;letter-spacing:-.03em;margin:8px 0 16px}header p{max-width:990px}a{color:#1764a1;text-underline-offset:3px}.eyebrow{font-size:12px;font-weight:750;letter-spacing:.13em;color:#c40d1e}.tags{display:flex;gap:10px;flex-wrap:wrap;margin:22px 0}.tags span{border-radius:5px;padding:7px 13px;font-weight:650;font-size:14px;color:white}.downloads{display:flex;flex-wrap:wrap;gap:12px}.downloads a{padding:8px 14px;background:white;border:1px solid #d7dee3;border-radius:5px}section{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:25px}article{background:white;border:1px solid #d7dee3;border-radius:8px;overflow:hidden}.preview{border:0;padding:0;width:100%;background:white;display:block;cursor:zoom-in}.preview:focus-visible{outline:4px solid #1764a1;outline-offset:-4px}img{display:block;width:100%;height:auto}.caption{padding:19px 23px 22px}h2{font-size:19px;line-height:1.35;margin:0 0 8px}.caption p{font-size:14px;margin:0 0 10px;color:#52606b}.caption a{font-size:14px}dialog{border:0;background:#101821;padding:12px;max-width:98vw;max-height:98vh;color:white}dialog::backdrop{background:#111c}dialog img{max-height:85vh;width:auto;max-width:95vw;object-fit:contain}dialog button{font:inherit;padding:5px 16px;margin-bottom:10px;border:0;border-radius:4px;cursor:pointer}footer{margin-top:24px;font-size:13px;color:#5d6973}@media(max-width:800px){section{grid-template-columns:1fr}main{padding:25px 15px}}
</style><main><header><div class="eyebrow">VITA-FL / KONZEPTE / 06.09.2026</div>
<h1>Feste Farben, klarere Übergänge, vier neue Folienideen.</h1>
<p>Vorschläge auf Basis der 39 bestehenden Folien und des Vita-FL-Codes. Empfehlung: Bereichsfarben durchgehend beibehalten, an Schnittstellen mehrere Navigationsfelder aktivieren. Die Blockchain/IPFS-Leiste ist gemeinsame Infrastruktur unter den zwei Verantwortungsketten.</p>
<div class="tags"><span style="background:#207548">DFL</span><span style="background:#1764a1">Agent / MCP / Inference / Log</span><span style="background:#7446a6">Blockchain / IPFS</span></div>
<div class="downloads"><a href="VITA-FL_Color_Navigation_Concepts.pptx">8 editierbare PowerPoint-Folien</a><a href="VITA-FL_Color_Navigation_Concepts.pdf">Vorschau als PDF</a><a href="README.md">Begründung, Einordnung & Codebelege</a></div></header><section>'''+"\n".join(items)+'''</section>
<footer>Englische Folien passend zum vorhandenen Deck. Die bestehende Präsentation bleibt unverändert. PNG/PDF sind lokale Layoutvorschauen; Schriftumbrüche in PowerPoint können abweichen. Die Fehlerbeispiele wurden aus Codepfaden abgeleitet; keine neuen Angriffsversuche durchgeführt.</footer>
<dialog><button type="button" id="close">Schließen · Esc</button><img alt="Vergrößerte Folienvorschau"></dialog>
<script>const d=document.querySelector('dialog');let previous;document.querySelectorAll('.preview').forEach(b=>b.addEventListener('click',()=>{previous=b;d.querySelector('img').src=b.dataset.src;d.querySelector('img').alt=b.querySelector('img').alt;d.showModal()}));document.querySelector('#close').addEventListener('click',()=>d.close());d.addEventListener('click',e=>{if(e.target===d)d.close()});d.addEventListener('close',()=>previous?.focus());</script></main></html>'''
    (HERE/"index.html").write_text(document)


def check_layout(prs):
    # Check the same font metrics/wrapping used by the local preview renderer.
    draw=ImageDraw.Draw(Image.new("RGB",(1600,900)))
    failures=[]
    for i,s in enumerate(prs.slides,1):
        for shape in s.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left+shape.width <= prs.slide_width+12700
            assert shape.top+shape.height <= prs.slide_height+12700
            if not shape.has_text_frame or not shape.text.strip(): continue
            f=shape.text_frame
            max_w=round((shape.width-f.margin_left-f.margin_right)*1600/prs.slide_width)
            needed=0
            for para in f.paragraphs:
                lines=renderer.paragraph_lines(draw,para,max_w,900/(7.5*72))
                for row in lines:
                    metrics=[font.getmetrics() for _,font,_ in row]
                    needed+=(max((a for a,_ in metrics),default=0)+max((d for _,d in metrics),default=0))*1.08
            available=(shape.height-f.margin_top-f.margin_bottom)*900/prs.slide_height
            if needed>available+5:
                failures.append(f"slide {i}: {shape.text[:65]!r}: needs {needed:.0f}px, has {available:.0f}px")
    if failures:
        raise RuntimeError("Text overflow:\n"+"\n".join(failures))


def main():
    prs=b.prepare_template()
    prs.core_properties.title="VITA-FL — Color navigation and slide concepts"
    palette(prs); architecture(prs); handoff(prs); sequence(prs); failure(prs); evidence(prs)
    data=role_heatmap(prs)
    hospital_privacy(prs)
    check_layout(prs)
    prs.save(HERE/"VITA-FL_Color_Navigation_Concepts.pptx")
    images=[]
    for slide,meta in zip(prs.slides,PAGES):
        img=renderer.render_slide(prs,slide)
        img.save(HERE/(meta["slug"]+".png"))
        images.append(img)
    images[0].save(HERE/"VITA-FL_Color_Navigation_Concepts.pdf",save_all=True,append_images=images[1:],resolution=144)
    # Content proposals without the palette page, in a two-column overview.
    overview_rows=(len(images)-1+1)//2
    sheet=Image.new("RGB",(1648,20+490*overview_rows),(232,237,241))
    draw=ImageDraw.Draw(sheet)
    font=ImageFont.truetype(renderer.FONT_BOLD,21)
    for i,(im,meta) in enumerate(zip(images[1:],PAGES[1:])):
        x=16+(i%2)*816; y=16+(i//2)*490
        sheet.paste(im.resize((800,450),Image.Resampling.LANCZOS),(x,y))
        draw.text((x+3,y+456),meta["title"],fill="#263440",font=font)
    sheet.save(HERE/"overview.png")
    (HERE/"concepts.json").write_text(json.dumps(dict(palette=dict(dfl=DFL,agent=AGENT,infrastructure=INFRA),pages=PAGES,heatmap=data),indent=2,ensure_ascii=False)+"\n")
    gallery()
    print(f"Created {len(images)} checked previews, PPTX, PDF and HTML gallery in {HERE}")


if __name__=="__main__":
    main()
