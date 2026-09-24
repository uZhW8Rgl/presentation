#!/usr/bin/env python3
"""Build three slide-28 proposals without changing the main presentation."""
import hashlib
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
import common as c
import design_a, design_b, design_c

DESIGNS = [
    ('a-shared-image', 'A · Gemeinsames Image und interne Schlüssel', design_a),
    ('b-inference-path', 'B · Vom verschlüsselten Modell zur Antwort', design_b),
    ('c-worker-roles', 'C · Dasselbe Image, unterschiedliche Rollen', design_c),
]


def main():
    main_deck = c.ROOT / 'VITA-FL_Thesis_Presentation_TU_Berlin.pptx'
    before = hashlib.sha256(main_deck.read_bytes()).digest()
    prs = c.b.prepare_template()
    prs.core_properties.title = 'VITA-FL — Worker and inference alternatives'
    for _, _, design in DESIGNS:
        design.build(prs)
    c.check_layout(prs)
    prs.save(c.HERE / 'VITA-FL_Worker_Inference_Alternatives.pptx')
    reopened = Presentation(c.HERE / 'VITA-FL_Worker_Inference_Alternatives.pptx')
    assert len(reopened.slides) == 3
    images = []
    for slide, (slug, _, _) in zip(reopened.slides, DESIGNS):
        ids = [shape.shape_id for shape in slide.shapes]
        assert len(ids) == len(set(ids))
        assert 'SECURITY SCOPE:' in slide.notes_slide.notes_text_frame.text
        preview = c.renderer.render_slide(reopened, slide)
        preview.save(c.HERE / (slug+'.png'))
        images.append(preview)
    images[0].save(c.HERE / 'VITA-FL_Worker_Inference_Alternatives.pdf',
                   save_all=True, append_images=images[1:], resolution=120)
    if not (c.HERE/'original-slide-28.png').exists():
        original = Presentation(main_deck)
        c.renderer.render_slide(original, original.slides[27]).save(c.HERE/'original-slide-28.png')
    overview = Image.new('RGB', (1660,1050), '#E8ECEF')
    draw = ImageDraw.Draw(overview)
    font = ImageFont.truetype(c.renderer.FONT_BOLD, 21)
    for index, (preview, (_, label, _)) in enumerate(zip(images, DESIGNS)):
        x = 20+index*820 if index < 2 else 430
        y = 20 if index < 2 else 535
        overview.paste(preview.resize((800,450), Image.Resampling.LANCZOS), (x,y))
        draw.text((x+5,y+464), label, font=font, fill='#263440')
    overview.save(c.HERE/'overview.png')
    assert before == hashlib.sha256(main_deck.read_bytes()).digest()
    print('Built 3 editable alternatives, PNGs, PDF and overview. Layout and notes checked; main deck unchanged.')


if __name__ == '__main__':
    main()
