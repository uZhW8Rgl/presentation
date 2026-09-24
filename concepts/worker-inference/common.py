"""Native, editable alternatives for the combined worker/inference slide."""
from pathlib import Path
import sys
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'concepts' / 'color-navigation'))
import build_presentation as b
import render_preview as renderer
from build_concepts import check_layout
from technology_icons import add_icon

GREEN, BLUE, PURPLE = '207548', '1764A1', '7446A6'
INK, MUTED, WHITE = '263440', '67747C', 'FFFFFF'
TINT = {GREEN:'EDF6F0', BLUE:'EDF4FA', PURPLE:'F3EEF8', INK:'F3F5F6'}

NOTES = '''The active Phala deployment uses exactly one dfl-worker service and one immutable image reference per worker CVM. That image packages the Node DFL service, Python training code, native tee_inference code and agent_receipts. Worker 0 starts training services and the inference receiver as processes in the same container. Other workers contain the same code but do not start inference. This is not a pair of independent container images.

Inference can load encrypted model publications using the existing participant RSA private key inside the worker CVM. Node owns RSA initialization: it obtains a dstack-derived wrapping key, creates or unseals the random RSA key, and materializes temporary PEM files in /run/vita-fl. Inference uses the same runtime RSA file. RSA unwraps model-transport symmetric keys; it is not the receipt-signing key. Sello and AIR use separate Ed25519 signing keys. Paths/context separate cryptographic purposes but are not an isolation mechanism between processes with access to the same dstack socket. No private-key export to the external agent is part of the intended inference interface. The external agent receives results and signed evidence, and can request/use public verification material.

The joint OCI image Digest covers both training and inference code. The on-chain verifier derives the image Digest and role-policy hash from the measured app_compose, binds the Compose hash into the event log, replays RTMR3, and compares it with the certificate-verified TDX quote. Image and role-policy checks authorize participation. The parser requires one worker service and one image. Training-only and combined-inference policies are separately allowed; each CVM has its own Compose/event log/quote and potentially different final RTMR3.

SECURITY SCOPE: these diagrams show intended key custody under the approved application. On-chain admission is distinct from Phala KMS key-release authorization. Startup loads keys before entering the registration state machine. The repository does not independently bind managed-KMS release to DeviceRegistry image/role allowlists. If a malicious upgrade of the same app is authorized by the KMS under the same key root, refusing that version's on-chain registration alone cannot prevent it from retrieving keys. Actual deployed KMS/upgrade authorization was not inspected. Trusted image execution and KMS/upgrade policy are necessary assumptions; the TEE does not isolate approved processes from one another. Do not present these diagrams as unconditional protection against malicious same-app upgrades.

PER-INFERENCE VERIFICATION: the agent checks AIR signatures, context, report-data binding, replayed RTMR3, measured Compose and its expected image Digest. Terraform derives this expected inference Digest from worker_image. The per-call Python verifier reports dcap_collateral_verified=False; do not describe it as an independent complete DCAP validation for every request. Receiver admission and registry-bound identity provide additional trust context.

TOOL TRANSPORT: MCP Tool request labels the conceptual tool call. In the deployed agent, LangChain wrappers invoke imported functions from mcp_server.py; the remote inference request uses HTTP REST. This diagram does not claim a public MCP transport server on the worker.

Source map, checked against local prototype on 20 September 2026:
vita-fl/phala/dstack-compose.worker.phala.tftpl:1–18,34–37,87–104
vita-fl/phala/dynamic-workers/main.tf:67
vita-fl/phala/main.tf:180
vita-fl/dfl/Dockerfile:90–99
vita-fl/dfl/start_node_neural_network.sh:332–365
vita-fl/dfl/node_server/src/server.ts:1823–1868,2890–2923
vita-fl/dfl/node_server/src/action_key.ts:65–98
vita-fl/dfl/node_server/src/participant_key.ts:247–286
vita-fl/tee_inference/service/model_source.py:146–220
vita-fl/tee_inference/service/attestation.py:49–60,109–121
vita-fl/agent/blockchain_source.py:467–480
vita-fl/smart_contracts/src/core/DeviceRegistry.sol:487–490
vita-fl/smart_contracts/src/attestation/AppComposeImage.sol:588–592
vita-fl/agent/tee_inference_client.py:445–472,605–624
'''

def page(prs, code, title, description):
    s = b.new_content_slide(prs, len(prs.slides)+1, title, '')
    b.add_domain_navigation(s, 28)
    for shape in s.shapes:
        if shape.has_text_frame and shape.text.startswith('Page '):
            shape.text_frame.paragraphs[0].runs[0].text = 'Variant '+code
    b.add_note(s, description+'\n\n'+NOTES)
    return s

def text(s, value, x, y, w, h=.35, size=16, color=INK, bold=False, center=True):
    shape = b.add_text(s, value, x, y, w, h, size, color, bold,
                      align=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT,
                      valign=MSO_ANCHOR.MIDDLE, margin=0)
    shape.name = 'Worker inference: '+value.replace('\n', ' / ')
    return shape

def panel(s, x, y, w, h, color=INK, fill=None, width=1.2):
    return b.add_box(s, x, y, w, h, fill=fill or TINT[color], line=color, line_width=width)

def route(s, points, color=INK, width=1.8, arrow=True):
    for i, (a, z) in enumerate(zip(points, points[1:])):
        b.add_line_segment(s, *a, *z, color, width, arrow=arrow and i==len(points)-2)

def icon(s, name, x, y, size, color=INK, fill=WHITE):
    if name == 'key':
        b.add_oval(s, x, y+.10*size, .42*size, .42*size, fill, color, 1.6)
        route(s, [(x+.41*size,y+.31*size),(x+.94*size,y+.31*size)], color, 2, False)
        for pos in [.73,.91]:
            route(s, [(x+pos*size,y+.31*size),(x+pos*size,y+.55*size)], color, 2, False)
    else:
        add_icon(s, name, x, y, size, color, fill, prefix='Worker inference icon: ')

def scope(s):
    text(s, 'On-chain admission checks the workload. KMS policy controls key release.',
         .65, 5.92, 12.03, .36, 15, MUTED)
