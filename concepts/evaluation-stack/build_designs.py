#!/usr/bin/env python3
"""Create native, editable architecture slide alternatives; preserve the main deck."""
import hashlib
import html
import json
from pathlib import Path
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
WORKSPACE = ROOT.parent
sys.path.insert(0, str(ROOT))
import build_presentation as b
import render_preview as renderer
sys.path.insert(0, str(ROOT / 'concepts' / 'color-navigation'))
from build_concepts import check_layout

INK, MUTED, WHITE, LINE = '263440', '65737D', 'FFFFFF', 'DCE3E8'
RED, BLUE, PURPLE, GREEN = 'C40D1E', '1764A1', '7446A6', '207548'
PALE, BLUE_TINT, PURPLE_TINT, GREEN_TINT = 'F2F5F7', 'EDF4FA', 'F3EEF8', 'EDF6F0'
SOURCES = [
    'SMEW/README.md', 'SMEW/smew/daemon.py', 'SMEW/smew/runner.py',
    'SMEW/smew/analysis/vita_fl.py', 'SMEW/reporting.py',
    'sustainability-measurement-agent/sma/sma.py',
    'vita-fl-td/README.md', 'vita-fl-td/src/vita_fl_td/cli.py',
    'vita-fl-td/src/vita_fl_td/runner.py',
    'vita-fl-td/src/vita_fl_td/transparency_audit.py',
]
NOTES = '''PURPOSE AND SCOPE
Meta-level architecture of the current evaluation workflow, checked against local
implementation on 27 September 2026. This slide describes responsibilities and
execution boundaries, not the outcome of a new run or a measured success rate.
English matches the existing presentation. These are alternatives for selection;
the main deck and its historical evaluation record are unchanged.

COMPONENT ROLES
SMEW (Sustainability Measurement Experiment Workflow) defines campaign matrices
and repetitions, distributes runs to daemons, supervises the measurement process,
collects reports and validates completeness and overall status. The orchestrator
can be hosted separately from the daemon. Diagrams group both under SMEW.
SMA (Sustainability Measurement Agent) starts vita-fl-td as a local subprocess,
records a measurement window with configured pre/post time, collects process
metadata and result files, and fetches Prometheus range-query data for that window
after execution. Prometheus performs the continuous scrape; SMA is not a sensor.
vita-fl-td drives the deployed system through its real control API, chain RPC and
cloud-agent endpoints; selected fault scenarios also use the Phala API. It checks
scenario-specific assertions, on-chain events and published evidence. When audit
is enabled, it re-verifies SCITT, AIR/TDX and Sello evidence from the transparency
log. It reuses production verifier primitives and obtains fresh Phala quote
appraisals; it does not implement an independent Intel DCAP verifier.
VITA-FL is the deployed system under test: distributed training and aggregation,
agent tools, attested inference, contracts, storage, transparency log and metrics.
The test driver controls a predeployed system; it does not deploy the cloud stack.

EXECUTION AND DATA FLOWS
In the evaluated Phala setup, SMEW daemon, SMA and the test driver execute on the
experiment host; VITA-FL and its Prometheus monitoring execute in Phala. Other
Docker profiles and a separately hosted SMEW orchestrator are possible.
SMEW daemon -> SMA: launch and supervise a parameterized run.
SMA -> TD: start a subprocess scenario and await its completion.
TD -> VITA-FL: API/RPC control and configured fault injection.
VITA-FL -> TD: state, events and published evidence for assertions.
VITA-FL -> Prometheus: exported progress, operating and gas metrics.
Prometheus -> SMA: time series queried for the recorded run window.
TD -> SMA -> SMEW reports: verdict, checks, JUnit and archived run artifacts.
Saved reports -> SMEW Marimo notebook: compare and visualize completed runs.
Marimo does not start the campaign or re-run the cryptographic verification.

INTERPRETATION LIMITS
A successful subprocess exit alone is not a passed system evaluation. TD supplies
semantic verdicts and evidence; SMEW also requires a successful measurement run
and a valid report. Configured assertions determine the claim supported by a run.
SMA's name does not imply energy measurement: this integration uses exported
Prometheus metrics (including training progress and gas use), not measured Joules.
The diagrams simplify internal VITA services and do not represent every API call.
Nested evaluation scopes in variant B are a logical hierarchy: SMEW is not one
OS process containing the other programs, and VITA is a separate deployed system.
Solid neutral arrows mean launch/control, blue arrows mean evidence/report return,
and dashed purple arrows mean retrieval of time series. In A, reverse arrows
summarize the result/evidence return at each boundary, not one identical payload.
'''


def txt(s, value, x, y, w, h=.34, size=16, color=INK, bold=False, center=False):
    q = b.add_text(s, value, x, y, w, h, size, color, bold,
                   align=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT,
                   valign=MSO_ANCHOR.MIDDLE, margin=0)
    q.name = 'Evaluation stack: ' + value.replace('\n', ' / ')
    return q


def box(s, x, y, w, h, fill=PALE, border=LINE, radius=True, width=1):
    return b.add_box(s, x, y, w, h, fill=fill, line=border, radius=radius, line_width=width)


def route(s, points, color=INK, dashed=False, width=1.8):
    for i, (start, end) in enumerate(zip(points, points[1:])):
        b.add_line_segment(s, *start, *end, color=color, width=width,
                           arrow=(i == len(points)-2), dashed=dashed)


def page(prs, letter, title, subtitle):
    s = b.new_content_slide(prs, len(prs.slides)+1, title, 'EVALUATION · HOW THE TOOLS WORK TOGETHER')
    b.add_domain_navigation(s, 25, active_domains={'DFL', 'Agent'})
    for q in s.shapes:
        if q.has_text_frame and q.text.startswith('Page '):
            q.text_frame.paragraphs[0].runs[0].text = 'Variant ' + letter
            q.width = b.Inches(1.1)
    txt(s, subtitle, .69, 1.52, 11.95, .57, 20)
    b.add_note(s, 'VARIANT '+letter+'\n\n'+NOTES+'\nSOURCE MAP (workspace-relative):\n'+'\n'.join(SOURCES))
    return s


def legend(s, y=5.99, start=.72):
    x = start
    for label, color, dash, advance in [
        ('Run control', INK, False, 2.1),
        ('Results / evidence', BLUE, False, 2.95),
        ('Metrics from Prometheus', PURPLE, True, 3.9),
    ]:
        route(s, [(x, y+.12), (x+.36, y+.12)], color, dash, 1.6)
        txt(s, label, x+.45, y-.015, advance-.48, .28, 11.5, MUTED)
        x += advance


def design_a(prs):
    s = page(prs, 'A', 'From campaign to verified results',
             'One workflow connects experiment planning, measurement and system checks.')
    xs = [.70, 3.79, 6.88, 9.97]
    names = ['SMEW', 'SMA', 'vita-fl-td', 'VITA-FL']
    roles = ['Campaigns &\nrepetitions', 'Run process &\ncollect metrics', 'Drive scenarios &\nverify evidence', 'Train, aggregate &\nrun inference']
    steps = ['01  PLAN', '02  MEASURE', '03  TEST', '04  EXECUTE']
    for x, name, role, step in zip(xs, names, roles, steps):
        txt(s, step, x+.05, 2.21, 2.56, .30, 12, MUTED, True)
        box(s, x, 2.62, 2.66, 1.47, fill=BLUE_TINT if name=='vita-fl-td' else PALE,
            border=BLUE if name=='vita-fl-td' else LINE)
        txt(s, name, x+.17, 2.77, 2.32, .43, 22, BLUE if name=='vita-fl-td' else INK, True)
        txt(s, role, x+.17, 3.32, 2.32, .59, 16)
    for x in xs[:-1]:
        route(s, [(x+2.68, 3.10), (x+3.06, 3.10)])
        route(s, [(x+3.06, 3.81), (x+2.68, 3.81)], BLUE)
    route(s, [(11.30, 4.11), (11.30, 4.64), (5.12, 4.64), (5.12, 4.11)], PURPLE, True)
    txt(s, 'Time series for the run window', 6.04, 4.22, 4.58, .31, 13, PURPLE, center=True)
    route(s, [(2.03, 4.11), (2.03, 4.98)], BLUE)
    box(s, .70, 5.02, 8.79, .77, fill=BLUE_TINT, border=LINE)
    txt(s, 'Collected report', .91, 5.13, 2.30, .31, 17, BLUE, True)
    txt(s, 'Verdicts + evidence + measurements', 3.28, 5.13, 5.95, .31, 17)
    route(s, [(9.53, 5.40), (9.91, 5.40)], BLUE)
    box(s, 9.97, 5.02, 2.66, .77, fill=WHITE, border=BLUE)
    txt(s, 'Marimo analysis', 10.12, 5.18, 2.36, .36, 17, BLUE, True, True)
    legend(s)


def design_b(prs):
    s = page(prs, 'B', 'Three evaluation layers, one deployed system',
             'Campaign → measured run → test scenario; VITA-FL is the external system under test.')
    box(s, .70, 2.30, 8.15, 3.42, fill=PALE, border=LINE)
    txt(s, 'SMEW', .96, 2.57, 2.0, .42, 23, INK, True)
    txt(s, 'Campaign', .96, 3.10, 2.02, .35, 16, INK, True)
    txt(s, 'Select scenarios\nand repeat runs', .96, 3.60, 1.96, .88, 16)
    box(s, 3.12, 2.96, 5.46, 2.48, fill=PURPLE_TINT, border='C8B9DB')
    txt(s, 'SMA', 3.39, 3.17, 1.70, .42, 22, PURPLE, True)
    txt(s, 'Run scope', 3.39, 3.70, 1.70, .34, 16, PURPLE, True)
    txt(s, 'Process, time\nwindow &\nmeasurements', 3.39, 4.12, 1.73, 1.07, 15.5)
    box(s, 5.35, 3.67, 2.97, 1.48, fill=WHITE, border=BLUE, width=1.5)
    txt(s, 'vita-fl-td', 5.57, 3.83, 2.52, .40, 21, BLUE, True)
    txt(s, 'Scenario logic\n& assertions', 5.57, 4.36, 2.52, .64, 16)
    box(s, 10.10, 2.30, 2.53, 3.42, fill=WHITE, border=INK, width=1.5)
    txt(s, 'VITA-FL', 10.31, 2.57, 2.10, .43, 23, INK, True)
    txt(s, 'Deployed in Phala', 10.31, 3.16, 2.10, .63, 16, MUTED)
    txt(s, 'Training\nAggregation\nInference\nPublished evidence', 10.31, 4.08, 2.10, 1.38, 15.5)
    route(s, [(8.33, 4.39), (10.07, 4.39)])
    txt(s, 'API calls', 8.92, 4.01, 1.08, .28, 12, INK, center=True)
    route(s, [(10.07, 4.95), (8.33, 4.95)], BLUE)
    txt(s, 'Evidence', 8.92, 5.01, 1.08, .30, 12, BLUE, center=True)
    txt(s, 'Each run combines test verdicts, artifacts and measurements for later analysis in Marimo.',
        .72, 5.91, 11.93, .34, 14, MUTED)


def design_c(prs):
    s = page(prs, 'C', 'Local evaluation, deployed system',
             'The test driver stays outside VITA-FL and checks its observable behavior.')
    box(s, .70, 2.28, 7.29, 3.56, fill=PALE, border=LINE)
    box(s, 8.53, 2.28, 4.10, 3.56, fill=WHITE, border=INK, width=1.5)
    txt(s, 'EXPERIMENT HOST', .95, 2.42, 6.75, .29, 13, MUTED, True)
    txt(s, 'PHALA DEPLOYMENT', 8.79, 2.42, 3.56, .29, 13, MUTED, True)
    for name, x, w, desc, color in [
        ('SMEW', .95, 1.83, 'Campaigns', INK),
        ('SMA', 3.15, 1.80, 'Measured run', PURPLE),
        ('vita-fl-td', 5.32, 2.42, 'Assertions', BLUE),
    ]:
        box(s, x, 2.96, w, 1.00, fill=WHITE, border=color)
        txt(s, name, x+.12, 3.08, w-.24, .36, 19, color, True, True)
        txt(s, desc, x+.09, 3.54, w-.18, .27, 12.5, MUTED, center=True)
    route(s, [(2.82, 3.40), (3.11, 3.40)])
    route(s, [(4.99, 3.40), (5.28, 3.40)])
    box(s, 9.06, 2.96, 3.15, 1.33, fill=PALE, border=LINE)
    txt(s, 'VITA-FL', 9.25, 3.10, 2.77, .40, 22, INK, True)
    txt(s, 'Training · inference\ncontracts · evidence', 9.25, 3.61, 2.77, .56, 14)
    route(s, [(7.78, 3.23), (9.02, 3.23)])
    txt(s, 'API / RPC', 7.90, 2.91, 1.07, .26, 11.5, center=True)
    route(s, [(9.02, 3.78), (7.78, 3.78)], BLUE)
    txt(s, 'Evidence', 7.93, 3.83, 1.04, .27, 11.5, BLUE, center=True)
    box(s, 9.06, 4.72, 3.15, .76, fill=PURPLE_TINT, border='C8B9DB')
    txt(s, 'Prometheus', 9.24, 4.83, 2.79, .34, 18, PURPLE, True)
    txt(s, 'Recorded system metrics', 9.24, 5.19, 2.79, .22, 11, PURPLE)
    route(s, [(10.63, 4.31), (10.63, 4.68)], PURPLE, True)
    route(s, [(9.02, 5.08), (4.10, 5.08), (4.10, 4.00)], PURPLE, True)
    txt(s, 'Time series for run window', 4.57, 4.66, 3.72, .28, 13, PURPLE, center=True)
    box(s, .95, 4.77, 2.80, .70, fill=BLUE_TINT, border=LINE)
    txt(s, 'Reports + Marimo', 1.10, 4.86, 2.5, .31, 16, BLUE, True)
    txt(s, 'Compare saved runs', 1.10, 5.21, 2.5, .20, 10.5, MUTED)
    route(s, [(1.87, 4.00), (1.87, 4.73)], BLUE)
    route(s, [(3.55, 4.00), (3.55, 4.36), (2.73, 4.36), (2.73, 4.73)], BLUE)
    route(s, [(6.54, 4.00), (6.54, 4.32), (4.65, 4.32), (4.65, 4.00)], BLUE)
    txt(s, 'SMA retrieves the time series for its measurement window after the scenario ends.',
        .72, 5.98, 11.91, .30, 12.5, MUTED)


def design_d(prs):
    s = page(prs, 'D', 'Four tools, four responsibilities',
             'Each component answers a different question in the evaluation.')
    x1, x2, x3 = .91, 3.53, 8.38
    for label, x, w in [('COMPONENT', x1, 2.35), ('EVALUATION QUESTION', x2, 4.60), ('CONTRIBUTION', x3, 4.0)]:
        txt(s, label, x, 2.16, w, .29, 12, MUTED, True)
    rows = [
        ('SMEW', 'What should run, and how often?', 'Campaigns & run status'),
        ('SMA', 'What happened during the run?', 'Process & measurement report'),
        ('vita-fl-td', 'Did the expected behavior occur?', 'Assertions & verified evidence'),
        ('VITA-FL', 'What does the deployed system do?', 'Models, events & receipts'),
    ]
    for i, (name, question, contribution) in enumerate(rows):
        y = 2.60 + i*.71
        box(s, .70, y, 11.93, .66, fill=BLUE_TINT if i==2 else (PALE if i%2==0 else WHITE), border=WHITE, radius=False, width=0)
        txt(s, name, x1, y+.08, 2.35, .49, 21, BLUE if i==2 else INK, True)
        txt(s, question, x2, y+.05, 4.57, .55, 16.5)
        txt(s, contribution, x3, y+.05, 3.97, .55, 16)
    box(s, .70, 5.63, 11.93, .50, fill=BLUE_TINT, border=LINE)
    txt(s, 'Together in Marimo: scenario outcomes + measurements + evidence',
        .92, 5.69, 11.49, .35, 18, BLUE, True, True)


DESIGNS = [
    ('a-workflow', 'A · Ablaufkette', 'Wer startet wen, und wie werden Ergebnisse gesammelt?', design_a),
    ('b-execution-layers', 'B · Verschachtelte Ebenen', 'Kampagne, Messlauf und Szenario als ineinanderliegende Verantwortungen.', design_b),
    ('c-deployment-boundary', 'C · Testrechner und Phala', 'Wo die Komponenten laufen und wie Steuerung, Belege und Messdaten fließen.', design_c),
    ('d-responsibilities', 'D · Vier Rollen', 'Eine kompakte Übersicht: Welche Frage beantwortet jedes Werkzeug?', design_d),
]


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def gallery():
    cards = []
    for slug, label, description, _ in DESIGNS:
        cards.append(f'''<article><h2>{html.escape(label)}</h2><p>{html.escape(description)}</p>
<button class="preview" data-image="{slug}.png" aria-label="{html.escape(label)} vergrößern"><img src="{slug}.png" alt="{html.escape(label)} – Folienvorschlag" /></button>
<p class="links"><a href="{slug}.pptx">PowerPoint</a><a href="{slug}.png">Bild öffnen</a></p></article>''')
    (HERE/'index.html').write_text('''<!doctype html><html lang="de"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>VITA-FL · Evaluationswerkzeuge</title>
<style>body{margin:0;background:#edf1f4;color:#263440;font:16px/1.5 system-ui,sans-serif}header,main{max-width:1440px;margin:auto;padding:24px}header{padding-bottom:0}h1{margin:0}header p{max-width:950px}a{color:#1764a1}nav,.links{display:flex;gap:24px;flex-wrap:wrap}main{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:24px}article{background:white;border-radius:12px;padding:20px;box-shadow:0 3px 18px #2634400c}h2{font-size:21px;margin:0}article p{margin:8px 0 16px}.preview{border:0;padding:0;background:none;width:100%;cursor:zoom-in}.preview img{width:100%;height:auto;display:block;border:1px solid #dce3e8;box-sizing:border-box}dialog{width:min(96vw,1600px);max-width:96vw;padding:8px;border:0;background:white}dialog::backdrop{background:#14212be6}dialog img{width:100%;display:block}#close{display:block;margin:6px 0 8px auto;padding:8px 18px;cursor:pointer}@media(max-width:900px){main{grid-template-columns:1fr}}
</style><header><h1>SMEW / SMA / vita-fl-td / VITA-FL</h1><p>Vier editierbare Folienentwürfe im Stil der TU-Berlin-Präsentation. Empfehlung: <strong>C</strong> erklärt das Zusammenspiel und die Grenze zwischen Testrechner und Phala am deutlichsten; <strong>A</strong> eignet sich für eine kurze Einführung.</p><nav><a href="VITA-FL_Evaluation_Stack_Alternatives.pptx">Alle Varianten als PowerPoint</a><a href="VITA-FL_Evaluation_Stack_Alternatives.pdf">PDF-Vorschau</a><a href="overview.png">Gesamtübersicht</a></nav></header><main>'''+''.join(cards)+'''</main><dialog id="zoom"><button id="close">Schließen · Esc</button><img alt="Vergrößerte Folie"></dialog><script>const modal=document.querySelector('#zoom');document.querySelectorAll('[data-image]').forEach(b=>b.addEventListener('click',()=>{modal.querySelector('img').src=b.dataset.image;modal.showModal()}));document.querySelector('#close').addEventListener('click',()=>modal.close());modal.addEventListener('click',e=>{if(e.target===modal)modal.close()});</script></html>''', encoding='utf-8')


def main():
    main_deck = ROOT/'VITA-FL_Thesis_Presentation_TU_Berlin.pptx'
    before = digest(main_deck)
    source_map = [{'path': p, 'sha256': digest(WORKSPACE/p)} for p in SOURCES]
    prs = b.prepare_template()
    prs.core_properties.title = 'VITA-FL — How the evaluation tools work together'
    for _, _, _, build in DESIGNS:
        build(prs)
    check_layout(prs)
    output = HERE/'VITA-FL_Evaluation_Stack_Alternatives.pptx'
    prs.save(output)
    reopened = Presentation(output)
    check_layout(reopened)
    previews=[]
    for s, (slug, _, _, build) in zip(reopened.slides, DESIGNS):
        visible = '\n'.join(q.text for q in s.shapes if q.has_text_frame)
        assert all(name in visible for name in ['SMEW','SMA','vita-fl-td','VITA-FL','Marimo'])
        assert 'SOURCE MAP' in s.notes_slide.notes_text_frame.text
        preview = renderer.render_slide(reopened, s)
        preview.save(HERE/(slug+'.png'))
        previews.append(preview)
        individual=b.prepare_template()
        build(individual)
        check_layout(individual)
        individual.save(HERE/(slug+'.pptx'))
        assert len(Presentation(HERE/(slug+'.pptx')).slides)==1
    previews[0].save(HERE/'VITA-FL_Evaluation_Stack_Alternatives.pdf', save_all=True,
                     append_images=previews[1:], resolution=120)
    overview=Image.new('RGB',(1660,1040),'#E8ECEF')
    draw=ImageDraw.Draw(overview)
    font=ImageFont.truetype(renderer.FONT_BOLD,21)
    for i,(preview,(_,label,_,_)) in enumerate(zip(previews,DESIGNS)):
        x,y=20+(i%2)*820,16+(i//2)*515
        draw.text((x+4,y),label,font=font,fill='#263440')
        overview.paste(preview.resize((800,450),Image.Resampling.LANCZOS),(x,y+37))
    overview.save(HERE/'overview.png')
    gallery()
    assert digest(main_deck)==before, 'Main presentation was changed'
    (HERE/'validation.json').write_text(json.dumps({
        'variant_count':len(prs.slides),'editable':'Native PowerPoint shapes and text',
        'layout':'passed; shape bounds and preview font metrics',
        'pptx_round_trip':'passed; four-slide deck and four one-slide files',
        'speaker_notes':'present; responsibilities, boundaries and source map',
        'main_deck_unchanged':True,'main_deck_sha256':before,
        'source_map':source_map,
        'preview_renderer':'Local Pillow renderer; final PowerPoint font layout can differ',
        'cloud_runs':'None; architecture presentation only',
    },indent=2)+'\n')
    print('Created and checked 4 editable slide alternatives, individual PPTX files, PNGs, PDF and gallery.')
    print('Main presentation unchanged: '+before)


if __name__=='__main__':
    main()
