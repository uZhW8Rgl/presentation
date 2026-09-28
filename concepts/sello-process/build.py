#!/usr/bin/env python3
"""Three native slide proposals separating Sello and AIR for the inference MCP call."""
from pathlib import Path
from copy import deepcopy
import hashlib,html,json,sys
from PIL import Image,ImageDraw,ImageFont
from pptx import Presentation
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
BASE=ROOT.parent
sys.path.insert(0,str(ROOT))
import build_presentation as b
import render_preview as renderer
sys.path.insert(0,str(ROOT/'concepts/color-navigation'))
from build_concepts import check_layout
INK='343434';MUTED='686868';LINE='B7B7B7';PALE='F4F4F4';WHITE='FFFFFF'
SELLO='A84B13';AIR='245AA5'
COLORS={'sello':SELLO,'air':AIR,'shared':MUTED}
FILLS={'sello':'FFF3E8','air':'EDF3FD','shared':'FFFFFF'}
STEM='VITA-FL_Sello_Process_Alternatives'
MAIN=ROOT/'VITA-FL_Thesis_Presentation_TU_Berlin.pptx'
CHAPTER=BASE/'overleaf/chapters/chapter4.tex'
STEPS=[
 {'id':1,'actor':'Agent','label':'Create authorization token','meaning':'The agent runtime holds the owner/issuer role and creates the short-lived token with audience, scopes, subject, PoP binding, owner HPKE public key and allowed log.'},
 {'id':2,'actor':'Agent / Inference TEE','label':'Verify TEE session','meaning':'Verify challenged RA-TLS session evidence and receiver/TLS/policy binding before transmitting the protected token-bearing request.'},
 {'id':3,'actor':'Agent → Inference TEE','label':'Request + token + PoP','meaning':'Receiver checks session, request PoP and authorization claims. The third MCP operation is then executed in AIR step A1.'},
 {'id':4,'actor':'Inference TEE','label':'Hash · encrypt · sign','meaning':'Hash exact logical tool input and raw response; HPKE-encrypt receipt body to the owner; then COSE-sign the encrypted envelope with the dedicated receiver Sello key.'},
 {'id':5,'actor':'Inference TEE → Transparency Log','label':'Publish Sello receipt','meaning':'Wrap the Sello envelope in a separately signed SCITT statement and submit to the log allowed by the authorization token.'},
 {'id':6,'actor':'Transparency Log → Inference TEE','label':'Return and verify inclusion evidence','meaning':'Log returns inclusion evidence; the receiver verifies it and stores the publication bundle before releasing a successful response.'},
 {'id':7,'actor':'Inference TEE → Agent','label':'Result + receipt + reference','meaning':'Shared response carries the AIR evidence bundle in its body, the Sello receipt in headers, and the Sello publication-bundle URL and transaction reference.'},
 {'id':8,'actor':'Agent / Inference TEE / Transparency Log','label':'Verify receipt and inclusion','meaning':'Agent verifies signature, decrypts receipt and checks token, action, exact input/output, session and status. It fetches the publication bundle from the TEE, gets log keys from the log, and verifies inclusion of the exact receipt bytes.'},
 {'id':9,'actor':'Agent → Transparency Log','label':'Publish verification context','meaning':'Current RA-TLS path additionally publishes original session evidence and owner-encrypted token/input/output audit context; both registrations must succeed before Sello-verified tool completion.'},
]
AIR_STEPS=[
 {'id':'A1','actor':'Inference TEE','label':'Run inference + create AIR','meaning':'Execute inference, request an inference-specific quote from dstack, sign AIR and construct the complete evidence bundle. Sello steps 4–9 then bind, publish and verify the returned bundle as tool output.'},
 {'id':'A2','actor':'Agent Runtime','label':'Check context + quote policy','meaning':'Validate session, request/model provenance and hashes, quote structure, AIR key and nonce. Check expected REPORTDATA and configured MRTD/RTMR0–2 allowlist before contacting Phala.'},
 {'id':'A3','actor':'Agent Runtime → Phala API → Agent Runtime','label':'POST quote / verify TDX / return JSON','meaning':'POST the complete inference quote as {"hex": quote.hex()} over approved HTTPS to /api/v1/attestations/verify. Phala returns JSON to the calling agent. Eventlog and Compose are not sent in this request.'},
 {'id':'A4','actor':'Agent Runtime','label':'Check verifier JSON','meaning':'Require success=true and quote.verified=true, validate the lookup checksum and compare the expected measurements and REPORTDATA with the submitted quote.'},
 {'id':'A5','actor':'Agent Runtime → Phala API → Agent Runtime','label':'GET raw quote / return original bytes','meaning':'A separate GET to /api/v1/attestations/raw/{checksum} on the same approved HTTPS origin returns the complete raw quote bytes directly to the agent.'},
 {'id':'A6','actor':'Agent Runtime','label':'Match quote bytes + verify AIR','meaning':'Compare complete returned quote bytes with the original. Check REPORTDATA binding, AIR signature, nonce, model/request/response/quote hashes, policy version, identity and signed measurements.'},
 {'id':'A7','actor':'Agent Runtime','label':'Verify RTMR3','meaning':'Replay the event log locally with SHA-384 and compare the computed RTMR3 with the measurement in the authenticated quote.'},
 {'id':'A8','actor':'Agent Runtime','label':'Check deployment policy','meaning':'Check measured Compose hash, expected image digest, contract/chain/RPC trust-root pins; then check output finiteness, probability range and threshold consistency.'},
 {'id':'A9','actor':'Agent Runtime → Transparency Log → Agent Runtime','label':'Publish AIR bundle + return inclusion evidence','meaning':'Sign and submit the complete verified AIR evidence bundle, wait for publication evidence and obtain public SCITT verification keys. This represents a multi-request publication flow, not exactly one HTTP roundtrip.'},
 {'id':'A10','actor':'Agent Runtime','label':'Verify inclusion + return success','meaning':'Verify inclusion of the exact signed statement with trusted log keys and check that publication binds the verified bundle. Store evidence and only then return verified-and-transparency-logged success.'},
]
FLOW=['1–3','A1',4,5,6,7,8,9]+[f'A{i}' for i in range(2,11)]
SOURCES=[
 ('vita-fl/transport_security/attestation.py',[302,338,372,381],'Delegate the inference quote to Phala, check the API result and bind it to the exact quote bytes.'),
 ('vita-fl/agent/run_agent.py',[329,365,369,384,587,593,993,1004],'The LLM can select prepared tools; explicit skill commands also have a direct route. Tool implementations enforce their verification sequence.'),
 ('vita-fl/agent/sello_client.py',[104,142,145,224,245,348],'Authorization, receipt verification, receiver bundle retrieval, and supplementary RA-TLS session/context publication.'),
 ('vita-fl/transport_security/ratls_client.py',[112,170],'Verify the attested TLS session before the protected request.'),
 ('vita-fl/tee_inference/service/app.py',[146,239,253,263,330,382],'Receiver authorization, action execution, fail-closed publication before response, and publication-bundle endpoint.'),
 ('vita-fl/agent_receipts/sello_v1.py',[117,160,263,306,363,406],'Token, receipt commitments, HPKE encryption then COSE signing, owner verification.'),
 ('vita-fl/agent_receipts/receiver_log.py',[25,50],'Receiver publishes and stores the inclusion bundle.'),
 ('vita-fl/agent_receipts/scitt.py',[93,106,175,216],'SCITT publication and exact-byte inclusion verification; public log-key retrieval.'),
 ('vita-fl/agent_receipts/environment.py',[30,75],'Receiver key and owner/issuer roles in current deployment.'),
 ('vita-fl/phala/dstack-compose.worker.phala.tftpl',[92],'Worker RA-TLS profile.'),
 ('vita-fl/phala/dstack-compose.contracts.phala.tftpl',[280,299],'Agent RA-TLS profile and owner key environment configuration.'),
 ('vita-fl/agent/tee_inference_client.py',[245,269,521,568,578,609,618,699,720,755],'Sello completes first; AIR appraisal and separate domain publication follow.'),
 ('vita-fl/tee_inference/service/attestation.py',[69,138],'Inference-specific AIR receipt, quote and complete evidence bundle.'),
]
LIMITS='''SCOPE AND MODELING
The diagrams depict run_and_verify_tee_inference, the third MCP function, in the current RA-TLS profile. Model loading and job preparation have already happened. Shared steps 1–3 group token creation, challenged session verification, request PoP and receiver authorization. The protected request follows successful session appraisal. Sello labels 4–9 retain their meaning. AIR numbering is now chronological and consistent across all variants: A1 creation → Sello 4–9 → A2–A8 appraisal → A9 publication → A10 inclusion and verified success. B displays every AIR number individually; A/C summarize A2–A8 and A9–A10 in grouped labels.
PLANNING AND RUNTIME
Agent Runtime is programmed client code inside the selected tool. The LLM chooses prepared tools and can reuse session artifacts; explicit skill commands also have a direct route. Within the tool, authorization, evidence checks and required publications follow fixed code. The LLM cannot skip an individual check in that invocation. This is not a claim that all natural-language output is cryptographically verified.
ACTORS AND EVIDENCE
A/C use three areas. B has four lanes, from top to bottom: Agent Runtime, Inference TEE, Transparency Log, Phala API. The runtime also holds the owner/issuer role; its HPKE private key is not sent to the receiver or log. The session quote and the later inference quote are distinct. AIR, Sello and outer SCITT statements use separate signing keys.
Sello hashes the exact logical tool input (canonical job_id) and complete AIR output bytes. Encrypt first, then COSE-sign the Sello receipt envelope. HPKE protects its hashes/metadata, not the whole inference response or separately published AIR evidence. The shared reply carries the AIR bundle in its body, Sello receipt in headers and a receiver publication-bundle reference. The agent retrieves that publication bundle from the TEE and public log keys from the log. Sello step 8 groups these retrievals and receipt/inclusion checks. Step 9 groups original session-evidence publication and owner-encrypted token/input/output audit export, both with inclusion checks.
EXPANDED AIR VERIFICATION
A2 performs local context/quote policy checks before network delegation. A3 is the POST of the inference-specific quote and direct JSON response from Phala to the agent. A4 validates flags, lookup ID and returned quote fields. A5 is a second HTTP transaction: GET the raw quote and return bytes to the agent. A6 compares all quote bytes and verifies AIR signature/bindings. A7 replays RTMR3; A8 checks measured deployment policy and basic output consistency. A9 is separate AIR-domain publication, with receipt waiting and log-key retrieval; A10 verifies inclusion and matching bundle hash before success. Both Phala responses are explicitly drawn back to the agent. They are ordinary responses, not callbacks, and neither is sent to the Inference TEE or transparency log. The earlier Phala session-quote check remains condensed in shared steps 1–3. A3 delegates cryptographic TDX appraisal; the agent does not independently validate the Intel collateral chain. The Phala API JSON is not a signed EAT.
LIMITS
Local checks, including RTMR3 replay, are mandatory but do not independently authenticate the hardware signature without the delegated quote check. RTMR3 replay verifies event-log consistency; Compose/image/policy appraisal determines whether the measured workload is allowed. Log inclusion does not prove inference arithmetic, clinical accuracy, availability or global log consistency. Publication failure prevents verified success but does not roll back prior computation. No new deployment or system test was run.
COLOR LEGEND
Dark orange identifies Sello handling and its RA-TLS audit extension. Dark blue identifies inference/AIR creation, appraisal, publication and inclusion. Gray identifies shared authorization/session setup and transport of both layers. Actor areas are neutral. Colors identify protocol layers, not secure/unsafe status. Every slide includes a color legend; all body elements are editable PowerPoint shapes. B additionally distinguishes solid execution-order arrows from dashed AIR data-reference arrows. The dashed paths A1 → 7 and 7 → A2 trace the original AIR bundle into the shared response and its later use by the agent. They are not additional HTTP requests and do not start A2 in parallel; solid 9 → A2 retains the implemented Sello-completion dependency.
'''


def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def text(s,t,x,y,w,h=.3,size=14,color=INK,bold=False,center=False,name=None):
 q=b.add_text(s,t,x,y,w,h,size,color,bold,align=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT,valign=MSO_ANCHOR.MIDDLE,margin=0)
 q.name=name or 'Sello process: '+t.replace('\n',' / ')
 return q

def rect(s,x,y,w,h,fill=WHITE,border=LINE,radius=True):
 q=b.add_box(s,x,y,w,h,fill=fill,line=border,radius=radius,line_width=.8)
 q.name='Sello AIR process panel'
 return q

def route(s,points,arrow=True,layer='shared',width=1.5,dashed=False,color=None):
 color=color or COLORS[layer]
 for i,(a,z) in enumerate(zip(points,points[1:])):
  q=b.add_line_segment(s,*a,*z,color,width,arrow=arrow and i==len(points)-2,dashed=dashed)
  q.name='Layer connector: '+layer

def mark(s,ids,title,x,y,w,h=.40,layer='sello',size=13.3,detail=None):
 rect(s,x,y,w,h,FILLS[layer],COLORS[layer])
 prefix='AIR steps: ' if layer=='air' else 'Sello steps: '
 if detail:
  text(s,title,x+.10,y+.035,w-.20,.27,size,COLORS[layer],True,False,prefix+','.join(map(str,ids)))
  text(s,detail,x+.10,y+.32,w-.20,h-.355,11.0,COLORS[layer])
 else:text(s,title,x+.08,y+.035,w-.16,h-.07,size,COLORS[layer],True,True,prefix+','.join(map(str,ids)))

def message(s,ids,title,x1,x2,y,layer='sello',label_x=None,label_w=None,size=12.8,mutual=False):
 route(s,[(x1,y),(x2,y)],layer=layer)
 if mutual:route(s,[(x2,y+.05),(x1,y+.05)],layer=layer,width=1.0)
 xx=label_x if label_x is not None else min(x1,x2)+.05
 ww=label_w if label_w is not None else abs(x2-x1)-.10
 prefix='AIR steps: ' if layer=='air' else 'Sello steps: '
 text(s,title,xx,y-.255,ww,.225,size,COLORS[layer],True,True,prefix+','.join(map(str,ids)))

def legend(s):
 for x,w,label,layer in [(0.70,3.0,'Sello · tool-call receipt','sello'),(4.36,3.9,'AIR · inference evidence','air'),(9.08,3.44,'Shared · auth / transport','shared')]:
  q=rect(s,x,1.42,.15,.15,FILLS[layer],COLORS[layer],False);q.name='Protocol legend: '+layer
  text(s,label,x+.23,1.385,w,.225,10.7,COLORS[layer],True)

def page(prs,letter,title,subtitle):
 s=b.new_content_slide(prs,len(prs.slides)+1,title,subtitle)
 for q in list(s.shapes):
  if q.name.startswith('Domain navigation: '):b.remove_shape(q)
  elif q.width>b.Inches(11) and b.Inches(1.28)<=q.top<=b.Inches(1.33):
   q.fill.solid();q.fill.fore_color.rgb=b.rgb(LINE);q.line.color.rgb=b.rgb(LINE)
  elif q.has_text_frame and q.text.startswith('Page '):
   q.text_frame.paragraphs[0].runs[0].text='Variant '+letter;q.width=b.Inches(1.15)
 s.name='Sello AIR process proposal '+letter
 refs=[{'path':p,'lines':lines,'finding':f,'sha256':digest(BASE/p)} for p,lines,f in SOURCES]
 b.add_note(s,'SELLO + AIR PROCESS · VARIANT '+letter+'\nCURRENT IMPLEMENTATION: 28 September 2026\nMCP: run_and_verify_tee_inference\n\nSELLO AND SHARED STEPS\n'+json.dumps(STEPS,ensure_ascii=False,indent=2)+'\n\nAIR STEPS\n'+json.dumps(AIR_STEPS,ensure_ascii=False,indent=2)+'\n\nORDER\n'+json.dumps(FLOW,ensure_ascii=False)+'\n\n'+LIMITS+'\nSOURCE REFERENCES\n'+json.dumps(refs,ensure_ascii=False,indent=2))
 legend(s)
 return s

def actors_columns(s,areas,top=1.73,bottom=6.23):
 for actor,x,w in areas:
  rect(s,x,top,w,bottom-top,PALE,PALE,False)
  rect(s,x,top,w,.33,INK,INK,False)
  text(s,actor,x+.08,top+.025,w-.16,.28,16.4,WHITE,True,True,'Sello actor: '+actor)

def sequence(prs):
 s=page(prs,'A','Sello + AIR: one inference call','VERTICAL SEQUENCE · THIRD MCP FUNCTION · DISTINCT EVIDENCE LAYERS')
 centers=[2.55,6.66,10.77]
 actors_columns(s,[('Agent',.65,3.80),('Inference TEE',4.76,3.80),('Transparency Log',8.87,3.80)])
 for x in centers:route(s,[(x,2.07),(x,6.20)],arrow=False,color=LINE,width=.8,dashed=True)
 message(s,[1,2,3],'1–3  Token · trusted session · call',centers[0],centers[1],2.38,layer='shared',mutual=True,size=12.4)
 mark(s,['A1'],'A1  Run inference + create AIR',4.96,2.56,3.40,.43,layer='air',size=13.3)
 mark(s,[4],'4  Hash · encrypt · sign Sello',4.96,3.12,3.40,.43,size=13.1)
 message(s,[5],'5  Publish Sello receipt',centers[1],centers[2],3.87)
 message(s,[6],'6  Inclusion proof · TEE checks',centers[2],centers[1],4.27,size=12.5)
 message(s,[7],'7  AIR bundle + Sello receipt + URL',centers[1],centers[0],4.68,layer='shared',size=12.2)
 mark(s,[8],'8  Verify Sello + inclusion',.83,4.91,3.44,.43,size=13.1)
 message(s,[9],'9  Log session + encrypted audit context',centers[0],centers[2],5.60,label_x=4.35,label_w=6.1,size=12.6)
 mark(s,[f'A{i}' for i in range(2,9)],'A2–A8  Verify AIR + TDX',.83,5.74,3.44,.40,layer='air',size=12.9)
 message(s,['A9','A10'],'A9–A10  Publish AIR + verify inclusion',centers[0],centers[2],6.18,layer='air',label_x=4.45,label_w=5.9,size=13.0)
 return s

def lanes(prs, generalized=False):
 if generalized:
  s=page(prs,'B','Sello + AIR: verification overview','THIRD MCP FUNCTION · GROUPED VERIFICATION RESPONSIBILITIES')
  s.name='Sello AIR process generalized B'
  for q in s.shapes:
   if q.has_text_frame and q.text=='Variant B':
    q.element.xpath('.//a:t')[0].text='Generalized B'
  note=s.notes_slide.notes_text_frame.text
  b.add_note(s,note+'\nGENERALIZED COPY\nA2/A4 groups context and verifier-result checks before and between API calls. A3/A5 groups Phala quote verification and raw-quote retrieval. The bidirectional API path represents two HTTP transactions with agent-side JSON validation between them, not one atomic verifier call. A6–A8 groups byte equality, AIR signature/bindings, RTMR3 replay and deployment/output policy. Solid flow into A6–A8 is taken only after the API checks complete. A10 has been omitted as a visible diagram step at the user request; runtime inclusion verification remains mandatory in the implementation and is implicit in successful A9 publication. Original step IDs are retained for mapping to the detailed slide. Dashed A1 → 7 → A2/A4 traces the same AIR bundle; all Sello steps are unchanged.\n')
 else:
  s=page(prs,'B','Sello + AIR: complete verification flow','THIRD MCP FUNCTION · AIR STEPS A1–A10 · FOLLOW STEP NUMBERS')
 areas=[('Agent Runtime',1.79,1.68),('Inference TEE',3.57,.64),('Transparency Log',4.31,.64),('Phala API',5.05,.64)]
 for actor,y,h in areas:
  rect(s,.65,y,12.02,h,PALE,PALE,False)
  rect(s,.65,y,1.68,h,INK,INK,False)
  label=actor.replace('Agent Runtime','Agent\nRuntime').replace('Transparency Log','Transparency\nLog').replace('Inference TEE','Inference\nTEE').replace('Phala API','Phala\nAPI')
  text(s,label,.76,y+.05,1.46,h-.10,14.0 if actor=='Transparency Log' else 15.2,WHITE,True,True,'Sello actor: '+actor)
 x0,x1,x2,x3,x4,x5=2.51,3.91,5.31,6.86,8.36,9.96
 w=1.30;wide=2.51;h=.54;yt=3.62;yl=4.36;yp=5.10
 # Common setup and TEE creation of the AIR bundle.
 route(s,[(x0+w/2,2.60),(x0+w/2,yt)])
 route(s,[(x0+w,yt+h/2),(x1,yt+h/2)])
 # Sello seal, log, proof check, shared reply and agent receipt verification.
 route(s,[(x1+w/2,yt+h),(x1+w/2,yl)],layer='sello')
 route(s,[(x1+w-.20,yl),(x1+w-.20,4.26),(x2+w/2,4.26),(x2+w/2,yt+h)],layer='sello')
 route(s,[(x2,yt+h/2),(5.26,yt+h/2),(5.26,2.14),(x2,2.14)])
 route(s,[(x2+w/2,2.41),(x2+w/2,2.64)],layer='sello')
 route(s,[(x2+w,2.91),(6.66,2.91),(6.66,4.23),(6.43,4.23),(6.43,yl)],layer='sello')
 route(s,[(x2+w,yl+h/2),(6.76,yl+h/2),(6.76,2.415),(x3,2.415)],layer='sello')
 # Dashed arrows trace the same AIR bundle, separately from execution order.
 route(s,[(3.55,yt),(3.55,3.30),(4.93,3.30),(4.93,2.00),(x2,2.00)],layer='air',dashed=True,width=1.7)
 route(s,[(x2+w,2.12),(x3,2.12)],layer='air',dashed=True,width=1.7)
 text(s,'AIR bundle',3.68,3.06,1.36,.21,10.5,AIR,center=True)
 if generalized:
  # Overview loop: pre-/post-API checks and the two quote API transactions.
  route(s,[(7.43,3.03),(7.43,yp)],layer='air')
  route(s,[(8.55,yp),(8.55,3.03)],layer='air')
  route(s,[(9.14,2.48),(x5,2.48)],layer='air')
  route(s,[(x5+wide,2.48),(12.59,2.48),(12.59,yl+h/2),(x5+wide,yl+h/2)],layer='air')
 else:
  # A3 POST/JSON roundtrip, then A5 GET/raw-bytes roundtrip.
  route(s,[(x3+w/2,2.90),(x3+w/2,yp)],layer='air')
  route(s,[(x3+w,yp+h/2),(8.26,yp+h/2),(8.26,2.415),(x4,2.415)],layer='air')
  route(s,[(x4+w/2,2.90),(x4+w/2,yp)],layer='air')
  route(s,[(x4+w,yp+h/2),(9.77,yp+h/2),(9.77,2.10),(x5,2.10)],layer='air')
  # Ordered local checks; publication must precede final inclusion/success.
  route(s,[(x5+wide/2,2.35),(x5+wide/2,2.43)],layer='air')
  route(s,[(x5+wide/2,2.73),(x5+wide/2,2.81)],layer='air')
  route(s,[(x5+wide,2.96),(12.59,2.96),(12.59,yl+h/2),(x5+wide,yl+h/2)],layer='air')
  route(s,[(x5,yl+h/2),(9.87,yl+h/2),(9.87,3.29),(x5,3.29)],layer='air')
 mark(s,[1,2,3],'1–3  Auth\n+ session',x0,1.95,w,.65,layer='shared',size=11.8)
 mark(s,['A1'],'A1  Infer\n+ AIR',x0,yt,w,h,layer='air',size=11.8)
 mark(s,[4],'4  Seal\nSello',x1,yt,w,h,size=11.8)
 mark(s,[5],'5  Log Sello\nreceipt',x1,yl,w,h,size=11.5)
 mark(s,[6],'6  Check\ninclusion',x2,yt,w,h,size=11.8)
 mark(s,[7],'7  AIR +\nSello reply',x2,1.87,w,h,layer='shared',size=11.5)
 mark(s,[8],'8  Verify\nSello',x2,2.64,w,h,size=11.8)
 mark(s,[9],'9  Session\n+ audit',x2,yl,w,h,size=11.5)
 if generalized:
  mark(s,['A2','A4'],'A2 / A4  Agent checks',x3,1.93,2.28,1.10,layer='air',size=12.0,detail='Context + API results\nBefore / after API calls')
  mark(s,['A3','A5'],'A3 / A5  Verify quote\n+ retrieve raw bytes',x3,yp,2.28,h,layer='air',size=11.6)
  mark(s,['A6','A7','A8'],'A6–A8  Verify AIR',x5,1.93,wide,1.10,layer='air',size=13.0,detail='Signature · bindings\nRTMR3 · deployment policy')
  mark(s,['A9'],'A9  Log AIR bundle',x5,yl,wide,h,layer='air',size=12.4)
  text(s,'passed',9.22,2.245,.64,.20,9.9,AIR,center=True)
  rect(s,6.93,3.72,1.00,.30,PALE,PALE,False)
  text(s,'POST / GET',6.96,3.755,.94,.20,10.0,AIR,center=True)
  rect(s,8.05,3.72,1.00,.30,PALE,PALE,False)
  text(s,'JSON / bytes',8.06,3.755,.98,.20,9.7,AIR,center=True)
 else:
  mark(s,['A2'],'A2  Context\nREPORTDATA\nplatform policy',x3,1.93,w,.97,layer='air',size=10.5)
  mark(s,['A3'],'A3  Verify\nTDX quote',x3,yp,w,h,layer='air',size=11.6)
  mark(s,['A4'],'A4  Check\nJSON flags\n+ quote fields',x4,1.93,w,.97,layer='air',size=11.0)
  mark(s,['A5'],'A5  Return\nraw quote',x4,yp,w,h,layer='air',size=11.6)
  mark(s,['A6'],'A6  Match quote bytes\nAIR signature + bindings',x5,1.85,wide,.50,layer='air',size=11.1)
  mark(s,['A7'],'A7  Verify RTMR3',x5,2.43,wide,.30,layer='air',size=11.4)
  mark(s,['A8'],'A8  Deployment policy',x5,2.81,wide,.30,layer='air',size=11.4)
  mark(s,['A9'],'A9  Log AIR bundle\n+ inclusion evidence',x5,yl,wide,h,layer='air',size=11.5)
  mark(s,['A10'],'A10  Inclusion → success',x5,3.14,wide,.30,layer='air',size=11.4)
  # Labels are placed in free lane space; filled masks interrupt only their own line.
  rect(s,7.01,3.72,1.00,.30,PALE,PALE,False)
  text(s,'POST quote',7.04,3.755,.94,.20,10.0,AIR,center=True)
  rect(s,8.55,3.72,.92,.30,PALE,PALE,False)
  text(s,'GET raw',8.58,3.755,.86,.20,10.0,AIR,center=True)
  rect(s,7.97,4.51,.58,.24,PALE,PALE,False)
  text(s,'JSON',8.00,4.535,.52,.19,9.8,AIR,center=True)
  text(s,'bytes',9.21,4.54,.48,.18,9.7,AIR,center=True)
 # A separate line-style legend keeps data references distinct from control flow.
 route(s,[(.73,5.87),(1.32,5.87)],layer='shared',width=1.5)
 text(s,'Solid: execution order',1.48,5.765,2.82,.21,10.5,MUTED)
 route(s,[(5.00,5.87),(5.59,5.87)],layer='air',dashed=True,width=1.7)
 text(s,'Dashed: AIR data reference',5.75,5.765,3.92,.21,10.5,AIR)
 text(s,'LLM selects the tool · Runtime enforces verification.',.68,6.00,11.98,.25,13.2,INK,True)
 return s

def handoff(prs):
 s=page(prs,'C','Sello + AIR: two distinct receipts','THIRD MCP FUNCTION · SELLO BINDS THE RETURNED AIR BUNDLE')
 actors_columns(s,[('Agent',.65,3.36),('Inference TEE',4.57,3.98),('Transparency Log',9.11,3.56)],bottom=6.04)
 route(s,[(3.80,2.63),(4.76,2.63)])
 route(s,[(6.56,2.89),(6.56,3.07)])
 route(s,[(8.34,3.51),(9.29,3.51)],layer='sello')
 route(s,[(9.29,3.94),(8.81,3.94),(8.81,4.43),(8.34,4.43)],layer='sello')
 route(s,[(4.76,4.43),(4.28,4.43),(4.28,3.935),(3.80,3.935)])
 route(s,[(2.33,4.12),(2.33,4.20)],layer='sello')
 route(s,[(2.33,4.57),(2.33,4.97),(9.29,4.97)],layer='sello')
 route(s,[(10.89,5.20),(10.89,5.32),(4.28,5.32),(4.28,5.70),(3.80,5.70)],layer='sello',width=1.25)
 route(s,[(2.33,5.94),(2.33,6.18),(10.89,6.18),(10.89,5.94)],layer='air')
 mark(s,[1,2,3],'1–3  Prepare trusted call',.86,2.28,2.94,.69,layer='shared',detail='Token · session · request',size=13.6)
 mark(s,['A1'],'A1  Inference + AIR receipt',4.76,2.34,3.58,.55,layer='air',size=14.0)
 mark(s,[4],'4  Create Sello receipt',4.76,3.07,3.58,.87,detail='Hash of complete AIR output bundle\nHPKE encryption · COSE signature',size=15.0)
 mark(s,[5,6],'5  Register Sello\n6  Return inclusion proof',9.29,3.12,3.20,1.01,size=14.2)
 mark(s,[],'Check Sello inclusion',4.76,4.18,3.58,.50,size=14.4)
 mark(s,[7],'7  AIR + Sello reply',.86,3.75,2.94,.37,layer='shared',size=13.0)
 mark(s,[8],'8  Verify Sello',.86,4.20,2.94,.37,size=13.1)
 mark(s,[9],'9  Log audit context',9.29,4.74,3.20,.46,size=13.2)
 text(s,'Session + encrypted owner context',4.63,4.72,4.20,.20,11.0,SELLO,center=True)
 mark(s,[f'A{i}' for i in range(2,9)],'A2–A8  Verify AIR / TDX',.86,5.46,2.94,.48,layer='air',size=13.2)
 mark(s,['A9','A10'],'A9–A10  Log + inclusion',9.29,5.46,3.20,.48,layer='air',size=13.2)
 return s

DESIGNS=[
 ('a-sequence','A · Vertikales Sequenzdiagramm','Die genaue Reihenfolge und jede Übergabe zwischen den drei Akteuren.',sequence),
 ('b-swimlanes','B · Horizontale Rollenbahnen','Vollständiger Ablauf A1–A10: zwei Phala-API-Rundwege, lokale Prüfungen und abschließende Log-Inclusion.',lanes),
 ('c-receipt-journey','C · Zwei Belege im Zusammenspiel','Drei Bereiche mit getrennten AIR- und Sello-Belegen sowie zwei Publikationswegen.',handoff),
]

def validate(prs):
 check_layout(prs)
 for s in prs.slides:
  generalized=s.name=='Sello AIR process generalized B'
  expanded=s.name=='Sello AIR process proposal B' or generalized
  expected_actors={'Agent','Inference TEE','Transparency Log'} | ({'Phala API'} if expanded else set())
  assert {q.name.removeprefix('Sello actor: ').replace('Agent Runtime','Agent') for q in s.shapes if q.name.startswith('Sello actor: ')}==expected_actors
  ids=[q.shape_id for q in s.shapes];assert len(ids)==len(set(ids))
  covered=set();air_covered=set()
  for q in s.shapes:
   assert not q.element.xpath('.//a:blip'),'Only editable native body shapes'
   if q.name.startswith('Sello steps: '):covered.update(int(v) for v in q.name.removeprefix('Sello steps: ').split(',') if v)
   if q.name.startswith('AIR steps: '):air_covered.update(v for v in q.name.removeprefix('AIR steps: ').split(',') if v)
  assert covered==set(range(1,10)),covered
  assert air_covered=={f'A{i}' for i in range(1,10 if generalized else 11)},air_covered
  if generalized:
   assert not any('A10' in q.text for q in s.shapes if q.has_text_frame)
  assert {q.name.removeprefix('Protocol legend: ') for q in s.shapes if q.name.startswith('Protocol legend: ')}=={'sello','air','shared'}
  assert 'SOURCE REFERENCES' in s.notes_slide.notes_text_frame.text
  assert 'Encrypt first' in s.notes_slide.notes_text_frame.text


def gallery():
 cards=''.join(f'<article><h2>{html.escape(title)}</h2><p>{html.escape(detail)}</p><a href="{slug}.png"><img src="{slug}.png" alt="{html.escape(title)}"></a><p><a href="{slug}.pptx">Editierbare Einzelfolie</a> · <a href="{slug}.png">Große Vorschau</a></p></article>' for slug,title,detail,_ in DESIGNS)
 (HERE/'index.html').write_text('''<!doctype html><html lang="de"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Sello + AIR · Drei Prozessentwürfe</title><style>body{margin:0;background:#eeeeee;color:#333;font:17px/1.5 system-ui}main{max-width:1650px;margin:auto;padding:26px}h1{margin:0}h2{margin:0;font-size:23px}article{background:white;padding:20px;margin:24px 0;border:1px solid #ccc}img{width:100%;display:block}a{color:#333;text-underline-offset:3px}p{margin:10px 0}footer{font-size:14px;color:#666}</style><main><h1>Sello + AIR: zwei Nachweisschichten</h1><p>Variante B erweitert Agent Runtime, Inference TEE und Transparency Log um die Phala API. B zeigt die AIR-Schritte A1–A10 einzeln; A/C fassen sie mit denselben Nummernbereichen zusammen. Alle Entwürfe zeigen den dritten MCP-Aufruf: Inferenz mit Sello und AIR. Dunkelorange = Sello, Dunkelblau = AIR, Grau = gemeinsame Autorisierung und Transport. Eine Legende steht auf jeder Folie. Die Diagramme sind editierbare PowerPoint-Elemente.</p><p><a href="VITA-FL_Sello_Process_Alternatives.pptx">Alle Varianten · PowerPoint</a> · <a href="b-generalized.png">Vereinfachte Kopie</a> · <a href="../../assets/sello-process-b.pptx">Detail + vereinfachte Folie</a> · <a href="VITA-FL_Sello_Process_Alternatives.pdf">PDF</a> · <a href="overview.png">Übersicht</a></p>'''+cards+'''<footer>Empfehlung A für präzise Reihenfolge; B für den Gesamtprozess; C für die transportierten Belege. Body und Publication-Bundle sind nicht dasselbe: das Receipt kommt in der Antwort, das Bundle vom TEE, die Logkeys vom Log. Die TEE erzeugt AIR vor dem Sello-Receipt. Nach Sello-Prüfung und RA-TLS-Auditexport prüft und publiziert der Agent das AIR-Domainbundle separat. Die ausgewählte Detailfolie und ihre vereinfachte Kopie sind als eigenständiges Asset für die Hauptpräsentation verfügbar. Kapitel 4 bleibt unverändert. Codebelege und Grenzen stehen in den Sprechernotizen und im README.</footer></main></html>''')

def main():
 before={str(p):digest(p) for p in (MAIN,CHAPTER)}
 for path,lines,_ in SOURCES:
  assert all(0<line<=len((BASE/path).read_text().splitlines()) for line in lines),(path,lines)
 prs=b.prepare_template();prs.core_properties.title='VITA-FL — Sello and AIR: separate evidence layers'
 for slug,title,detail,fn in DESIGNS:
  fn(prs)
  single=b.prepare_template();fn(single);validate(single);single.save(HERE/(slug+'.pptx'))
 validate(prs);prs.save(HERE/(STEM+'.pptx'))
 checked=Presentation(HERE/(STEM+'.pptx'));validate(checked)
 images=[]
 for s,(slug,*_) in zip(checked.slides,DESIGNS):
  im=renderer.render_slide(checked,s);im.save(HERE/(slug+'.png'));images.append(im)
 images[0].save(HERE/(STEM+'.pdf'),save_all=True,append_images=images[1:],resolution=120)
 # Preserve the selected detailed slide and export its generalized companion.
 pair=b.prepare_template();pair.core_properties.title='VITA-FL — detailed and generalized Sello process'
 lanes(pair);lanes(pair,generalized=True);validate(pair)
 asset=ROOT/'assets/sello-process-b.pptx';pair.save(asset)
 checked_pair=Presentation(asset);validate(checked_pair)
 general=b.prepare_template();lanes(general,generalized=True);validate(general)
 general.save(HERE/'b-generalized.pptx')
 general_image=renderer.render_slide(checked_pair,checked_pair.slides[1])
 general_image.save(HERE/'b-generalized.png')
 images[1].save(HERE/'b-presentation-pair.pdf',save_all=True,append_images=[general_image],resolution=120)
 assert [str(q.element.xml) for q in checked_pair.slides[0].shapes]==[str(q.element.xml) for q in checked.slides[1].shapes]

 overview=Image.new('RGB',(1632,1530),'#ECECEC');draw=ImageDraw.Draw(overview);font=ImageFont.truetype(renderer.FONT_BOLD,24)
 for i,(im,(_,title,*_)) in enumerate(zip(images,DESIGNS)):
  y=12+i*506;draw.text((18,y),title,font=font,fill='#343434');overview.paste(im.resize((832,468),Image.Resampling.LANCZOS),(16,y+33))
  summaries=[['Dunkelorange: Sello-Aufrufbeleg','Dunkelblau: AIR-Inferenznachweis','Grau: gemeinsamer Aufruf und Transport','Mit Legende auf jeder Folie'],['A1–A10 lückenlos nummeriert','Phala: POST → JSON, GET → Bytes','Agent: Quote, AIR, RTMR3, Policy','Log-Publikation + Inclusion-Prüfung'],['3 große Verantwortungsbereiche','AIR-Bundle entsteht in der TEE','Sello bindet dessen Rückgabe','Agent prüft und publiziert AIR separat']]
  for j,label in enumerate(summaries[i]):draw.text((890,y+100+j*58),label,font=ImageFont.truetype(renderer.FONT_REGULAR,23),fill='#444444')
 overview.save(HERE/'overview.png');gallery()
 assert before=={str(p):digest(p) for p in (MAIN,CHAPTER)}
 (HERE/'validation.json').write_text(json.dumps({'slides':3,'single_slide_pptx':4,'approved_presentation_pair':2,'generalized_groups':['A2/A4','A3/A5','A6–A8'],'generalized_A10_visible':False,'layout':'passed','roundtrip':'passed','sello_and_shared_steps_1_to_9':'passed','air_steps_A1_to_A10':'passed','air_numbering':'B individual; A/C grouped ranges of the same steps','three_color_legend_on_each':'passed','scope':'run_and_verify_tee_inference (third MCP call)','color_map':COLORS,'actor_areas':{'A':3,'B':4,'C':3},'phala_inference_quote_before_agent_rtmr3':'passed','native_shapes':True,'main_and_chapter_preserved':before,'code_sources':[{'path':path,'sha256':digest(BASE/path)} for path,_,_ in SOURCES],'visual_review':'pending'},indent=2)+'\n')
 print('Built three proposals and the detailed/generalized presentation pair; main deck and Chapter 4 unchanged by this builder.')
if __name__=='__main__':main()
