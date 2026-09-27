#!/usr/bin/env python3
"""Revised horizontal attack trees: native PowerPoint, Chapter 4 with explicit implementation context."""
from pathlib import Path
import hashlib, html, json, re, sys
from copy import deepcopy
from collections import Counter
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[2]
CHAPTER=ROOT.parent/'overleaf/chapters/chapter4.tex'
MAIN=ROOT/'VITA-FL_Thesis_Presentation_TU_Berlin.pptx'
sys.path.insert(0,str(ROOT))
import build_presentation as b
import render_preview as renderer
sys.path.insert(0,str(ROOT/'concepts/color-navigation'))
from build_concepts import check_layout
INK,MUTED,WHITE,LINE='2D2D2D','666666','FFFFFF','B8B8B8'
RED,GREEN,AMBER='C40D1E','267347','94600B'
RED_TINT,GREEN_TINT,PALE,AMBER_TINT='FBECEF','EDF7EF','F7F7F7','FFF4DF'
ASSESSMENT_STYLE={
    'neutral':(WHITE,LINE,INK),
    'possible':(RED_TINT,RED,RED),
    'boundary':(AMBER_TINT,AMBER,AMBER),
    'blocked':(GREEN_TINT,GREEN,GREEN),
}
X=[.65,3.12,5.90,8.95]
W=[1.88,2.16,2.38,3.72]
DATA=json.loads((HERE/'trees.json').read_text())

def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def nodes(n):
    yield n
    for c in n.get('children',[]): yield from nodes(c)
def leaves(n): return [v for v in nodes(n) if not v.get('children')]
def txt(s,t,x,y,w,h,size=13,color=INK,bold=False,center=False):
    shape=b.add_text(s,t,x,y,w,h,size,color,bold,align=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT,valign=MSO_ANCHOR.MIDDLE,margin=0)
    shape.name='Attack tree: '+t.replace('\n',' / ')
    return shape

def box(s,x,y,w,h,fill,border=LINE):
    return b.add_box(s,x,y,w,h,fill=fill,line=border,radius=True,line_width=.8)

def line(s,a,z,color=LINE,width=1.25):
    return b.add_line_segment(s,*a,*z,color=color,width=width)

def gate(s,op,x,cy):
    w=.39 if op=='OR' else .46
    color=INK
    box(s,x-w/2,cy-.135,w,.27,WHITE,LINE)
    txt(s,op,x-w/2+.015,cy-.12,w-.03,.24,10.3,color,True,True)

def place(root):
    terminal=leaves(root)
    assert 4<=len(terminal)<=8,len(terminal)
    step=4.12/len(terminal)
    for i,node in enumerate(terminal):
        node['_cy']=1.66+step*(i+.5)
        node['_x'],node['_w'],node['_h']=X[3],W[3],min(.56,step-.045)
    def inner(n,depth):
        kids=n.get('children',[])
        if not kids: return n['_cy']
        assert depth<3
        centers=[inner(c,depth+1) for c in kids]
        n['_cy']=(centers[0]+centers[-1])/2
        n['_x'],n['_w']=X[depth],W[depth]
        n['_h']=2.0 if depth==0 else (.86 if depth==1 else .76)
        return n['_cy']
    inner(root,0)

def edges(s,n):
    kids=n.get('children',[])
    if not kids: return
    end=n['_x']+n['_w']; cy=n['_cy']; gx=end+.28; rail=end+.53
    color=LINE
    line(s,(end,cy),(rail,cy),color)
    line(s,(rail,kids[0]['_cy']),(rail,kids[-1]['_cy']),color)
    for c in kids:
        line(s,(rail,c['_cy']),(c['_x'],c['_cy']),color)
        edges(s,c)
    gate(s,n['op'],gx,cy)

def draw_node(s,n,depth=0):
    x,w,h,cy=n['_x'],n['_w'],n['_h'],n['_cy']
    kids=n.get('children',[])
    fill,border,fg=ASSESSMENT_STYLE[n['assessment']]
    q=box(s,x,cy-h/2,w,h,fill,border)
    q.name='Assessed node: '+n['assessment']+' / '+n['label']
    if depth==0:
        txt(s,'ATTACKER GOAL',x+.11,cy-h/2+.11,w-.22,.20,9.1,fg,True)
        txt(s,n['label'],x+.11,cy-h/2+.38,w-.22,h-.48,15.5,fg,True)
    elif kids:
        txt(s,n['label'],x+.12,cy-h/2+.05,w-.24,h-.10,13.1 if depth==1 else 12.4,fg,True)
    else:
        txt(s,n.get('short_label',n['label']),x+.12,cy-h/2+.025,w-.24,h-.05,14.0,INK)
    for c in kids: draw_node(s,c,depth+1)


def add_legend(s):
    """One assessment palette applies to goals, subgoals and leaf conditions."""
    txt(s,'LEGEND',.69,5.84,.64,.22,8.5,MUTED,True)
    entries=[
        (1.40,2.05,'Action / condition','neutral'),
        (4.00,2.32,'Possible in threat model','possible'),
        (7.03,2.12,'Assumption boundary','boundary'),
        (9.76,2.62,'Blocked under assumptions','blocked'),
    ]
    for x,width,label,status in entries:
        fill,border,_=ASSESSMENT_STYLE[status]
        q=box(s,x,5.86,.17,.17,fill,border)
        q.name='Attack legend swatch: '+label
        q=txt(s,label,x+.25,5.84,width,.22,10.0,INK)
        q.name='Attack legend label: '+label


def slide(prs,t,i):
    s=b.new_content_slide(prs,i,t['title'],f"CHAPTER 4 · ATTACK TREE {i} / {len(DATA['trees'])} · "+' / '.join(t['threats']))
    b.add_domain_navigation(s,31,active_domains=set(t['domains']))
    # Keep navigation outside the security color scale; remove the blue domain badge.
    for q in s.shapes:
        if q.name.startswith('Domain navigation: ') and q.name.endswith(' badge'):
            domain=q.name.removeprefix('Domain navigation: ').removesuffix(' badge')
            q.fill.solid();q.fill.fore_color.rgb=b.rgb('555555' if domain in t['domains'] else 'F5F5F5')
            q.line.color.rgb=b.rgb('555555' if domain in t['domains'] else LINE)
        if q.name.startswith('Domain navigation: ') and q.name.endswith(' label'):
            domain=q.name.removeprefix('Domain navigation: ').removesuffix(' label')
            for para in q.text_frame.paragraphs:
                for run in para.runs:run.font.color.rgb=b.rgb(WHITE if domain in t['domains'] else MUTED)
        if q.width>b.Inches(11) and b.Inches(1.28)<=q.top<=b.Inches(1.33):
            q.fill.solid();q.fill.fore_color.rgb=b.rgb(LINE);q.line.color.rgb=b.rgb(LINE)
    for q in s.shapes:
        if q.has_text_frame and q.text.startswith('CHAPTER 4'):
            q.width=b.Inches(6.8)
    txt(s,'OR: any route · AND: all conditions',7.60,1.02,3.79,.22,9.5,MUTED)
    for q in s.shapes:
        if q.has_text_frame and q.text.startswith('Page '):
            q.text_frame.paragraphs[0].runs[0].text=f"B light · {i}"
            q.width=b.Inches(1.1)
    # Pale field and column captions retain B's left-to-right hierarchy.
    b.add_box(s,.55,1.52,12.20,4.33,PALE,PALE,radius=False,line_width=0)
    tree=deepcopy(t['root']);place(tree);edges(s,tree);draw_node(s,tree)
    add_legend(s)
    txt(s,t['takeaway'],.69,6.06,11.96,.21,10.5,INK,True)
    note=DATA['modeling_notes']+'\n\nTREE-SPECIFIC MODEL\n'+json.dumps(t,ensure_ascii=False,indent=2)
    b.add_note(s,note)
    return s

def combined_assessment(op, states):
    if op == 'OR':
        return ('possible' if 'possible' in states else 'boundary' if 'boundary' in states
                else 'blocked' if all(v == 'blocked' for v in states) else 'neutral')
    return ('blocked' if 'blocked' in states else 'boundary' if 'boundary' in states
            else 'possible' if 'possible' in states else 'neutral')


def validate():
    assert digest(CHAPTER)==DATA['source_sha256']
    chapter=CHAPTER.read_text(); coverage=set()
    for t in DATA['trees']:
        for ref in t['threats']:
            assert '{threat:'+ref+'}' in chapter,ref
            coverage.add(ref)
        for n in nodes(t['root']):
            assert n.get('note'),n['label']
            assert n['assessment'] in ASSESSMENT_STYLE and n['assessment_reason']
            for ref in n.get('refs',[]):
                prefix='threat' if re.fullmatch(r'T\d+',ref) else ('assumption' if ref.startswith('AS') else 'req')
                assert '{'+prefix+':'+ref+'}' in chapter,ref
            if n.get('children'):
                assert n['op'] in ('OR','AND') and len(n['children'])>=2
                assert n['assessment'] == combined_assessment(n['op'], [c['assessment'] for c in n['children']]), n['label']
            else:
                assert n['attacker_capability'] and n['required_credentials'], n['label']
        assert t['source_lines']
        assert t['implementation_context']['sources']
        for source in t['implementation_context']['sources']:
            assert digest(ROOT.parent/source['path']) == source['sha256'], source['path']
            assert all(0 < line <= len((ROOT.parent/source['path']).read_text().splitlines()) for line in source['lines']), source['path']
    return sorted(coverage,key=lambda v:int(v[1:]))

def gallery():
    ts=DATA['trees']
    buttons=''.join(f'<button data-i="{i}">{html.escape(t["name_de"])}</button>' for i,t in enumerate(ts))
    meta=[{'slug':t['slug'],'name':t['name_de'],'takeaway':t['takeaway'],'root':t['root']} for t in ts]
    page='''<!doctype html><html lang="de"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>VITA-FL Attack Trees · B hell</title><style>body{margin:0;background:#f3f6f8;color:#263440;font:16px/1.5 system-ui}main{max-width:1450px;margin:auto;padding:24px}nav{display:flex;gap:10px;flex-wrap:wrap;margin:20px 0}button{font:inherit;background:white;border:1px solid #bdceda;border-radius:7px;padding:9px 14px;cursor:pointer}button[aria-pressed=true]{background:#555555;color:white}a{color:#333333}img{display:block;width:100%;background:white;border:1px solid #dce3e8}details{background:white;margin-top:20px;padding:18px}li{margin:8px 0}.note{color:#65737d;font-size:14px}h1{margin:0}</style><main><h1>Attack Trees · B auf hellem Grund</h1><p>Sechs konkrete Angreiferziele aus Kapitel 4. OR: alternative Wege. AND: gemeinsam notwendige Bedingungen. Die Farben bewerten jeden Knoten unabhängig von seiner Ebene: Rot = Verfügbarkeit mit genannter externer Fähigkeit angreifbar; Gelb = konkrete externe Annahme-/Nachweisgrenze; Grün = durch spezifizierte Checks unter den Annahmen blockiert; Weiß = neutrale Handlung oder Voraussetzung. Alle sechs Bäume wurden gegen den Code geprüft. Jeder Pfad benennt Akteur und benötigte Signierbefugnis. Der Betrieb einer TEE bedeutet keinen Zugriff auf ihre Schlüssel. Die Bewertung ist kein experimentelles Ergebnis und keine pauschale Sicherheitsgarantie. Die kurzen Stichworte auf den Folien sind unten ausführlich erklärt. Die Bäume sind Modelle, keine bestätigten Exploits. Grün gilt nur für die gezeigten Pfade. <a href="implementation-review.md">Prüfergebnis und Grenzen</a></p><p><a href="VITA-FL_Attack_Trees_B_Light.pptx">Editierbare PowerPoint</a> · <a href="VITA-FL_Attack_Trees_B_Light.pdf">PDF</a> · <a href="overview.png">Alle sechs Bäume</a></p><nav>BUTTONS</nav><a id="full"><img id="preview"></a><p id="takeaway"></p><details open><summary>Inhalt und Voraussetzungen im Detail</summary><div id="detail"></div></details><p class="note">Kapitel 4 unverändert. Die sechs freigegebenen Folien sind auf Seiten 24–29 der Hauptpräsentation integriert. Quellen und Modellierungsgrenzen stehen in den Sprechernotizen.</p></main><script>const data=META;function node(n){const li=document.createElement('li');const strong=document.createElement('strong');strong.textContent=n.label+(n.op?' ['+n.op+']':'');li.append(strong);const p=document.createElement('div');p.className='note';p.textContent=n.note+' Bewertung: '+n.assessment_reason+(n.attacker_capability?' Akteur/Fähigkeit: '+n.attacker_capability+' Benötigte Autorität: '+n.required_credentials:'')+' Quellen: '+(n.refs||[]).join(', ');li.append(p);if(n.children){const ul=document.createElement('ul');n.children.forEach(c=>ul.append(node(c)));li.append(ul)}return li}function show(i){const t=data[i];document.querySelector('#preview').src=t.slug+'.png';document.querySelector('#preview').alt=t.name;document.querySelector('#full').href=t.slug+'.png';document.querySelector('#takeaway').textContent=t.takeaway;const d=document.querySelector('#detail');d.replaceChildren();const ul=document.createElement('ul');ul.append(node(t.root));d.append(ul);document.querySelectorAll('button').forEach(b=>b.setAttribute('aria-pressed',Number(b.dataset.i)===i))}document.querySelectorAll('button').forEach(b=>b.addEventListener('click',()=>show(Number(b.dataset.i))));show(0);</script></html>'''
    (HERE/'index.html').write_text(page.replace('BUTTONS',buttons).replace('META',json.dumps(meta,ensure_ascii=False)))

def main():
    before={str(p):digest(p) for p in (CHAPTER,MAIN)}
    coverage=validate(); prs=b.prepare_template()
    prs.core_properties.title='VITA-FL — Detailed attack trees, horizontal light'
    for i,t in enumerate(DATA['trees'],1): slide(prs,t,i)
    check_layout(prs)
    target=HERE/'VITA-FL_Attack_Trees_B_Light.pptx';prs.save(target)
    reopened=Presentation(target);check_layout(reopened)
    images=[]
    for s,t in zip(reopened.slides,DATA['trees']):
        assert 'TREE-SPECIFIC MODEL' in s.notes_slide.notes_text_frame.text
        assert len([q for q in s.shapes if q.name.startswith('Attack legend swatch: ')])==4
        assert len([q for q in s.shapes if q.name.startswith('Attack legend label: ')])==4
        im=renderer.render_slide(reopened,s);im.save(HERE/(t['slug']+'.png'));images.append(im)
    images[0].save(HERE/'VITA-FL_Attack_Trees_B_Light.pdf',save_all=True,append_images=images[1:],resolution=120)
    out=Image.new('RGB',(1640,1510),'#E9EEF2');draw=ImageDraw.Draw(out);font=ImageFont.truetype(renderer.FONT_BOLD,22)
    for i,(im,t) in enumerate(zip(images,DATA['trees'])):
        x=15+815*(i%2);y=10+500*(i//2)
        draw.text((x,y),t['name_de'],font=font,fill='#263440');out.paste(im.resize((800,450)),(x,y+36))
    out.save(HERE/'overview.png');gallery()
    assert before=={str(p):digest(p) for p in (CHAPTER,MAIN)}
    stats=[{'slug':t['slug'],'nodes':len(list(nodes(t['root']))),'leaves':len(leaves(t['root']))} for t in DATA['trees']]
    (HERE/'validation.json').write_text(json.dumps({'slides':len(images),'coverage':coverage,'structure':stats,'layout':'passed','pptx_roundtrip':'passed','four_status_legend_on_each_slide':'passed','root_assessments':{t['slug']:t['root']['assessment'] for t in DATA['trees']},'neutral_OR_AND_gates':True,'assessment_composition':'passed','all_leaf_actor_and_credentials':'passed','implementation_source_hashes':'passed','assessment_counts':dict(Counter(n['assessment'] for t in DATA['trees'] for n in nodes(t['root']))),'source_main_preserved':before,'visual_review':'pending'},indent=2)+'\n')
    print(json.dumps({'slides':len(images),'structure':stats,'source_main':'unchanged'}))
if __name__=='__main__':main()
