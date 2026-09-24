"""Shared native-shape helpers and source notes for key-use slide alternatives."""

from pathlib import Path
import sys
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'concepts' / 'color-navigation'))
import build_presentation as b
import attestation_slides
import render_preview as renderer
from build_concepts import check_layout
from technology_icons import add_icon

BLUE, PURPLE, GREEN = '1764A1', '7446A6', '207548'
INK, MUTED, WHITE = '263440', '67747C', 'FFFFFF'
TINT = {BLUE:'EDF4FA', PURPLE:'F3EEF8', GREEN:'EDF6F0', INK:'F3F5F6'}

_notes_deck = b.prepare_template()
COMMON_NOTES = attestation_slides.keys(_notes_deck, b).notes_slide.notes_text_frame.text
COMMON_NOTES += '''

Concrete uses reviewed for the alternatives: the secp256k1 action key signs protocol transactions, local-update commitments and aggregation statements. The random RSA-3072 key signs model packages and unwraps the symmetric encryption keys used for model transport. The derived AES-256-GCM sealing key protects RSA private-key state on disk; it is distinct from the fresh symmetric keys used to encrypt model bytes. Sello uses Ed25519 to sign receiver-created tool-action receipts submitted to the transparency log; AIR uses a separate Ed25519 key to sign inference evidence verified by the agent. The log proves recording; it does not independently prove computation. Public action, RSA and Sello identities are enrollment-bound; AIR has a separate inference-quote binding.
'''


def page(prs, code, title, description):
    s = b.new_content_slide(prs, len(prs.slides)+1, title, '')
    b.add_domain_navigation(s, 15)
    for shape in s.shapes:
        if shape.has_text_frame and shape.text.startswith('Page '):
            shape.text_frame.paragraphs[0].runs[0].text = 'Variant '+code
    b.add_note(s, description+'\n\n'+COMMON_NOTES)
    return s


def text(s, value, x, y, w, h=.3, size=14, color=INK, bold=False, center=True):
    shape = b.add_text(s,value,x,y,w,h,size,color,bold,
                       align=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT,
                       valign=MSO_ANCHOR.MIDDLE,margin=0)
    shape.name = 'Key use: '+value.replace('\n',' / ')
    return shape


def panel(s,x,y,w,h,color,fill=None,width=1.1):
    return b.add_box(s,x,y,w,h,fill=fill or TINT[color],line=color,line_width=width)


def route(s,points,color=INK,width=1.8,arrow=True):
    for index,(start,end) in enumerate(zip(points,points[1:])):
        b.add_line_segment(s,*start,*end,color,width,
                           arrow=arrow and index==len(points)-2)


def icon(s,name,x,y,size,color,fill=WHITE):
    if name=='key':
        b.add_oval(s,x,y+.13*size,.42*size,.42*size,fill,color,1.5)
        route(s,[(x+.41*size,y+.34*size),(x+.94*size,y+.34*size)],color,2,False)
        for pos in [.73,.91]:
            route(s,[(x+pos*size,y+.34*size),(x+pos*size,y+.57*size)],color,2,False)
    else:
        add_icon(s,name,x,y,size,color,fill,prefix='Key use icon: ')


def pill(s,value,x,y,w,color,size=14,h=.38):
    panel(s,x,y,w,h,color)
    text(s,value,x+.08,y+.035,w-.16,h-.07,size,color,True)
