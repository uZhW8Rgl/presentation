"""Native slide helpers and audited regular-round aggregation flow."""
from pathlib import Path
import sys
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
sys.path.insert(0,str(ROOT))
sys.path.insert(0,str(ROOT/'concepts'/'color-navigation'))
import build_presentation as b
import render_preview as renderer
from build_concepts import check_layout
from technology_icons import add_icon

GREEN, PURPLE, BLUE = '207548', '7446A6', '1764A1'
INK, MUTED, WHITE = '263440', '67747C', 'FFFFFF'
TINT={GREEN:'EDF6F0',PURPLE:'F3EEF8',BLUE:'EDF4FA',INK:'F3F5F6'}

NOTES='''These diagrams show a regular training round, not round-0 bootstrap. The aggregation rule is equal-weight FedAvg, with no validation-loss gate in the active policy.

Workers send encrypted, signed local-model packages directly to the current aggregator over authenticated HTTPS. These local updates are not stored in IPFS. The receiver verifies package signatures and expected context, decrypts the package, compares its plaintext model hash with the worker commitment, and records the worker's action-key-signed commitment on chain through GMStorage. One logical worker has at most one consistent accepted submission.

The selected aggregator closes the input window only when the policy-required minimum is reached. AggregationPolicy freezes an ordered input fingerprint and explicit accepted count. inputRoot is a rolling Keccak accumulator over accepted worker/commitment pairs in admission order, not a Merkle root. Closing disallows further submissions. Deadline expiry does not automatically close the round or bypass the minimum. The visible phrase frozen input set means this fingerprint plus count and per-worker records; it is not an on-chain copy of the model tensors.

Within the TEE, server.ts stages authorized model files whose SHA-256 hashes match the corresponding on-chain records and rechecks the hashes after copying. The local file count must exactly equal the immutable accepted count. Python then sorts unique worker-address filenames canonically, loads all staged models, checks the count, and computes the arithmetic mean of every corresponding state-dict tensor with equal weight per accepted model. This ordering is distinct from admission order in the rolling input commitment. Invalid or incomplete inputs block aggregation. Deterministic execution here refers to the pinned implementation and ordered inputs, not universal bitwise reproducibility across arbitrary hardware or software environments.

The resulting model is RSA-signed, encrypted in an AES-GCM model bundle, and published to IPFS together with its encrypted-bundle signature and per-recipient key bundle. The aggregator obtains all three CIDs before creating the aggregation statement. The action key signs an EIP-712 statement binding round, aggregator, input fingerprint/count, algorithm/policy, output model and bundle hashes, publication hash over the three CIDs, and nonce. The signature domain binds chain ID and GMStorage. This action signature is separate from the RSA model-package signature.

GMStorage.finalizeRoundWithAggregation validates current aggregator authorization, closed inputs, nonempty references and hashes, and unpublished round state. AggregationPolicy reconstructs the statement using stored frozen inputs and policy, then verifies the registered action-key signature. Solidity does not rerun FedAvg; correct computation is trusted to the previously attested application holding that key. The single finalization transaction records model references/evidence, marks the round complete, updates counters and advances the round. IPFS publication happened earlier and is not atomic with the blockchain transaction. Next-aggregator selection is a subsequent call, not part of this finalization transaction.

Source map, checked against local prototype on 20 September 2026:
vita-fl/dfl/node_server/src/server.ts:625–702,1310–1365,1930–2004,2229–2330,2786–2828
vita-fl/dfl/neural_network/cli.py:1248–1320
vita-fl/dfl/node_server/src/ipfs.ts:300–367
vita-fl/dfl/node_server/src/bc_client.ts:542–650
vita-fl/smart_contracts/src/core/GMStorage.sol:180–248,294–373
vita-fl/smart_contracts/src/core/AggregationPolicy.sol:16,203–215,260–305
vita-fl/smart_contracts/src/core/AggregatorSelection.sol:169–195
'''


def page(prs,code,title,description):
    s=b.new_content_slide(prs,len(prs.slides)+1,title,'')
    b.add_domain_navigation(s,18)
    for shape in s.shapes:
        if shape.has_text_frame and shape.text.startswith('Page '):
            shape.text_frame.paragraphs[0].runs[0].text='Variant '+code
    b.add_note(s,description+'\n\n'+NOTES)
    return s


def text(s,value,x,y,w,h=.3,size=14,color=INK,bold=False,center=True):
    shape=b.add_text(s,value,x,y,w,h,size,color,bold,
        align=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT,
        valign=MSO_ANCHOR.MIDDLE,margin=0)
    shape.name='Aggregation: '+value.replace('\n',' / ')
    return shape


def panel(s,x,y,w,h,color,fill=None,width=1.1):
    return b.add_box(s,x,y,w,h,fill=fill or TINT[color],line=color,line_width=width)


def route(s,points,color=INK,width=1.8,arrow=True):
    for index,(start,end) in enumerate(zip(points,points[1:])):
        b.add_line_segment(s,*start,*end,color,width,
                          arrow=arrow and index==len(points)-2)


def icon(s,name,x,y,size,color,fill=WHITE):
    add_icon(s,name,x,y,size,color,fill,prefix='Aggregation icon: ')


def pill(s,value,x,y,w,color,size=14,h=.38):
    panel(s,x,y,w,h,color)
    text(s,value,x+.08,y+.035,w-.16,h-.07,size,color,True)


def model(s,x,y,w=.58,h=.63,color=GREEN):
    panel(s,x,y,w,h,color,WHITE)
    for row in range(2):
        for col in range(3):
            b.add_box(s,x+.09+col*(w-.18)/3,y+.12+row*(h-.24)/2,
                      (w-.23)/3,(h-.31)/2,
                      fill=color if (row+col)%2 else TINT[color],line=WHITE,
                      radius=False,line_width=0)
