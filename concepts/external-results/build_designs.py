#!/usr/bin/env python3
"""Build three editable results-slide proposals from archived external runs."""
from pathlib import Path
import argparse
import ast
import base64
from datetime import datetime
import hashlib
import html
import json
import sys
from zipfile import ZipFile
from collections import Counter
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
WORKSPACE = ROOT.parent
sys.path.insert(0, str(ROOT))
import build_presentation as b
import render_preview as renderer
sys.path.insert(0, str(ROOT / 'concepts/color-navigation'))
from build_concepts import check_layout

INK, MUTED, LINE, WHITE, PALE = '263440', '65737D', 'DCE3E8', 'FFFFFF', 'F3F5F6'
GREEN, GREEN_TINT = '247044', 'EDF6F0'
AMBER, AMBER_TINT = '926500', 'FFF7E2'
GRAY, GRAY_TINT = '626C74', 'F0F2F4'
RECOVERY = WORKSPACE / 'SMEW/reports/vita-fl-recovery-5r-61b96e51/2026_09_25_17_22_19_749db20e/subprocess_output'
PARTIAL = WORKSPACE / 'SMEW/reports/vita-fl-partial-5r-61b96e51/2026_09_25_17_35_46_15585dd8/subprocess_output/result.json'
OLDER = WORKSPACE / 'SMEW/reports/vita-fl-recovery-5r-98664d28/2026_09_25_16_34_03_e48f9a7c/subprocess_output/result.json'


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def completed_security_evidence(report_directory, recovery):
    """Reject scratch/incomplete runs before any proposal artifacts are written."""
    report = Path(report_directory).expanduser().resolve(strict=True)
    assert report.is_relative_to((WORKSPACE / 'SMEW/reports').resolve()), 'Use the canonical SMEW reports directory, not scratch.'
    smew_path, sma_path = report / 'smew.json', report / 'run.json'
    result_path, archive_path = report / 'subprocess_output/result.json', report / 'subprocess_output/run-artifacts.zip'
    smew, sma = json.loads(smew_path.read_text()), json.loads(sma_path.read_text())
    assert smew['state'] == 'succeeded' and smew['exit_code'] == 0, 'SMEW run has not succeeded.'
    assert sma['status'] == 'ok', 'SMA run has not completed successfully.'
    assert sma['user_data']['subprocess_exit_code'] == 0 and not sma['user_data']['subprocess_cancelled']
    result = json.loads(result_path.read_text())
    assert result['scenario'] == 'security' and result['status'] == 'PASS', 'TD security scenario is not PASS.'
    assert result['checks'] and all(c['passed'] is True for c in result['checks'])
    assert not result['errors'] and not result['measurement_errors']
    with ZipFile(archive_path) as archive:
        assert json.loads(archive.read('result.json')) == result, 'Archived TD result differs from the exported result.'
        security = json.loads(archive.read('security.json'))
        audit = json.loads(archive.read('transparency-audit.json'))
    inventory = WORKSPACE / 'vita-fl-td/src/vita_fl_td/security_bridge.py'
    assignments = [node for node in ast.parse(inventory.read_text()).body if isinstance(node, ast.Assign)]
    expected = next(ast.literal_eval(node.value) for node in assignments
                    if any(isinstance(target, ast.Name) and target.id == 'REQUIRED_CHECKS' for target in node.targets))
    assert len(expected) == 16 and len(set(expected)) == 16
    checks = security['checks']
    assert tuple(c['name'] for c in checks) == expected, 'Security probe set is incomplete or unexpected.'
    assert security['ok'] is True and security['status'] == 'PASS'
    assert security['stage'] == 'security-probes-completed' and security['execution_mode'] == 'cloud-agent-security'
    assert all(c['passed'] is True and c['status'] == 'PASS' for c in checks)
    assert result['security_checks'] == checks, 'TD and security bridge checks disagree.'
    assert audit['status'] == 'PASS' and len(audit['checks']) == 17
    assert all(c['passed'] is True for c in audit['checks'])
    audit_groups = Counter(c['name'].split(':')[0] for c in audit['checks'])
    assert audit_groups == {'scitt': 10, 'air-tdx': 1, 'session': 3, 'sello': 3}
    condition = result['condition']
    assert condition['training']['worker_count'] == 6 and condition['training']['rounds'] == 5
    assert condition['training']['client_limit'] == 5 and condition['training']['model_submission_deadline_ms'] == 45000
    assert len(result['participation']['rounds']) == 5 and result['participation']['all_workers_active']
    assert len(result['participation']['workers']) == 6
    context, previous = result['context'], recovery['context']
    assert context['environment'] == 'phala'
    assert context.get('test_driver_source_sha256') and previous.get('test_driver_source_sha256')
    assert context['test_driver_source_sha256'] != previous['test_driver_source_sha256']
    assert context['deployment_images'] and context['deployment_images'] == previous['deployment_images'], 'Do not claim identical images for a changed deployment.'
    assert context['deployment_app_id'] and context['deployment_app_id'] != previous['deployment_app_id'], 'Expected a separate new deployment.'
    run_date = datetime.fromisoformat(result['started_at']).date().isoformat()
    assert run_date == '2026-09-29', 'These proposals compare recovery on 25 Sep with security on 29 Sep.'
    groups = {
        'access_and_configuration': [name for name in expected if name.startswith(('ui_', 'control_'))],
        'contract_guards': [name for name in expected if name.startswith('coordination_')],
        'positive_inference': [name for name in expected if name == 'agent_positive_control'],
        'negative_inference_authorization': [name for name in expected if name.startswith('inference_')],
    }
    assert [len(names) for names in groups.values()] == [5, 4, 1, 6]
    http_checks = [c for c in checks if c['actual'].get('http_status') is not None]
    contract_checks = [c for c in checks if c['name'].startswith('coordination_')]
    assert len(http_checks) == 10 and all(c['actual']['http_status'] == 401 for c in http_checks)
    assert all(c['actual'].get('reason_matches', c['actual'].get('challenge_matches')) is True
               for c in http_checks)
    assert len(contract_checks) == 4 and all(
        c['actual']['reverted'] and c['actual']['reason_matches'] and
        c['boundary'] == 'deployed-contract-eth_call-simulation' for c in contract_checks)
    configuration = next(c for c in checks if c['name'] == 'control_configuration_unchanged')
    assert configuration['actual'] == configuration['expected']
    command = condition['inference']['command']
    policy_path = Path(command[command.index('--audit-policy') + 1]).resolve(strict=True)
    assert policy_path.is_relative_to(WORKSPACE), 'Security policy source must be a local workspace file.'
    policy = json.loads(policy_path.read_text())
    assert hashlib.sha256(base64.b64decode(policy['scitt_keys_cbor_base64'])).hexdigest() == policy['scitt_keys_sha256']
    policy_evidence = {'path': str(policy_path.relative_to(WORKSPACE)), 'sha256': digest(policy_path),
                       'scitt_keys_sha256': policy['scitt_keys_sha256'], 'trust_bootstrap': policy['trust_bootstrap']}
    sources = [smew_path, sma_path, report / 'manifest.json', result_path, archive_path, inventory, policy_path]
    return {
        'run_date': run_date, 'run_id': result['run_id'], 'status': result['status'],
        'report_directory': str(report.relative_to(WORKSPACE)),
        'smew_state': smew['state'], 'smew_exit_code': smew['exit_code'], 'sma_status': sma['status'],
        'checks_passed': len(checks), 'check_groups': groups, 'coverage_gaps': security['coverage_gaps'],
        'observed_outcomes': {'unauthorized_http_401': 10, 'expected_eth_call_reverts': 4,
                              'configuration_unchanged': 1, 'authorized_inference_passed': 1},
        'audit_passed': len(audit['checks']), 'audit_groups': dict(audit_groups),
        'audit_duration_seconds': audit['duration_seconds'], 'audit_limitations': audit['limitations'],
        'successful_training_rounds': 5, 'active_workers': 6,
        'deployment_app_id': context['deployment_app_id'], 'recovery_deployment_app_id': previous['deployment_app_id'],
        'same_deployment_images': True, 'deployment_images': context['deployment_images'],
        'test_driver_source_sha256': context.get('test_driver_source_sha256'),
        'recovery_test_driver_source_sha256': previous.get('test_driver_source_sha256'),
        'provisioning_terraform_parallelism': context.get('provisioning_terraform_parallelism'),
        'runtime_compose_sha256': context.get('runtime_compose_sha256'),
        'audit_policy': policy_evidence,
        'sources': {str(path.relative_to(WORKSPACE)): digest(path) for path in sources},
    }


def completed_collection_evidence(report_directories, recovery, security=None):
    """Require archived, controlled early/partial/zero runs before showing PASS."""
    assert len(report_directories) == 3, 'Collection needs exactly three canonical reports.'
    runs, sources = {}, {}
    open_topic = '0x8bcf39476f12f2198587ea4025f7f3b26decc7e5b991c682cdab31748a79fea9'
    close_topic = '0x16a8faf1a158dfdc2ee3ddfd2e8bfe4df79f18858372c90b7d467167f2fac57f'
    for directory in report_directories:
        report = Path(directory).expanduser().resolve(strict=True)
        assert report.is_relative_to((WORKSPACE / 'SMEW/reports').resolve()), 'Use canonical reports, not scratch.'
        smew_path, sma_path = report / 'smew.json', report / 'run.json'
        result_path = report / 'subprocess_output/result.json'
        archive_path = report / 'subprocess_output/run-artifacts.zip'
        smew, sma = json.loads(smew_path.read_text()), json.loads(sma_path.read_text())
        assert smew['state'] == 'succeeded' and smew['exit_code'] == 0, 'Collection SMEW run has not succeeded.'
        assert sma['status'] == 'ok' and sma['user_data']['subprocess_exit_code'] == 0
        assert not sma['user_data']['subprocess_cancelled']
        result = json.loads(result_path.read_text())
        scenario = result['scenario']
        assert scenario in {'early', 'partial', 'zero'} and scenario not in runs, 'Expected one distinct run per scenario.'
        assert result['status'] == 'PASS' and result['checks']
        assert all(c['passed'] is True for c in result['checks'])
        assert not result['errors'] and not result['measurement_errors']
        with ZipFile(archive_path) as archive:
            assert json.loads(archive.read('result.json')) == result
            events = json.loads(archive.read('chain-events.json'))
            identity = json.loads(archive.read('chain.json'))
            observations = [json.loads(line) for line in archive.read('observations.jsonl').decode().splitlines() if line]
        condition = result['condition']
        assert condition['scenario'] == scenario and condition['upload_gate']['kind'] == 'control'
        assert condition['verify_worker_participation'] is True
        assert condition['training'] == {'worker_count': 6, 'client_limit': 5, 'rounds': 5,
                                         'epoch': 1, 'model_submission_deadline_ms': 45000}
        allowed_count = {'early': 5, 'partial': 2, 'zero': 0}[scenario]
        assert condition['allowed_clients'] == allowed_count

        def checked(name, actual=True):
            matching = [c for c in result['checks'] if c['name'] == name]
            assert matching and all(c['actual'] == actual and c['expected'] == actual for c in matching), name

        checked('round_client_limit', 5)
        checked('exact_successful_round_count', 6)  # Includes bootstrap; displayed count excludes it.
        checked('all_workers_actually_worked')
        gate = result['upload_gate']
        allowed, blocked, roster = map(set, (gate['allowed_participants'], gate['blocked_participants'], gate['roster']))
        aggregator = gate.get('observed_aggregator') or gate['selected_aggregator']
        assert len(roster) == 6 and aggregator in roster
        assert len(allowed) == allowed_count and len(blocked) == 5 - allowed_count
        assert allowed.isdisjoint(blocked) and allowed | blocked == roster - {aggregator}
        assert all(type(gate['blocked_counts'].get(account)) is int and gate['blocked_counts'][account] > 0
                   for account in blocked), 'Receiver must confirm each blocked client.'
        for name, value in [('gate_round', 1), ('gate_roster_size', 6), ('gate_allowed_count', allowed_count),
                            ('gate_blocked_count', 5 - allowed_count), ('gate_clients_disjoint', True),
                            ('gate_all_clients_accounted', sorted(roster - {aggregator})),
                            ('gate_blocking_confirmed_by_receiver', True)]:
            checked(name, value)
        participation = result['participation']
        assert participation['all_workers_active'] is True and participation['bootstrap_included'] is False
        assert len(participation['workers']) == 6 and len(participation['rounds']) == 5
        attempts = result['attempts']
        assert attempts['aborted'] == ([1] if scenario == 'zero' else [])
        assert attempts['completed'] == ([0, 2, 3, 4, 5, 6] if scenario == 'zero' else [0, 1, 2, 3, 4, 5])
        assert [r['round'] for r in participation['rounds']] == attempts['completed'][1:]
        context = result['context']
        assert context['environment'] == 'phala' and context['test_driver_source_sha256']
        assert context['deployment_images'] == recovery['context']['deployment_images'], 'Image pins changed.'
        run_date = datetime.fromisoformat(result['started_at']).date().isoformat()
        assert run_date == '2026-09-29', 'Update the slide dates when consuming a different-day collection run.'
        policy_events = [e for e in events if not e.get('removed')
                         and e['address'].lower() == identity['aggregation_policy'].lower()
                         and len(e.get('topics', [])) > 1 and int(e['topics'][1], 16) == 1]
        opened = [e for e in policy_events if e['topics'][0] == open_topic]
        closed = [e for e in policy_events if e['topics'][0] == close_topic]
        assert len(opened) == 1, 'Exactly one policy opening for the controlled attempt is required.'
        words = [int(opened[0]['data'][i:i+64], 16) for i in range(2, len(opened[0]['data']), 64)]
        target, opening, deadline = words[:3]
        assert target == 5 and deadline - opening == 45
        timing = {'opened_timestamp': opening, 'deadline_timestamp': deadline,
                  'basis': 'PolicyOpened data and InputsClosed block timestamp; not local-timer instrumentation'}
        details = {}
        if scenario in {'early', 'partial'}:
            checked('closed_input_count', allowed_count)
            checked('round_1_expected_count', allowed_count)
            checked('gate_exact_allowed_contributors', sorted(allowed))
            first = participation['rounds'][0]
            assert first['input_count'] == allowed_count and set(first['contributors']) == allowed
            assert len(closed) == 1
            assert int(closed[0]['data'][2:66], 16) == allowed_count
            if closed[0].get('blockTimestamp') is not None:
                stamp = int(closed[0]['blockTimestamp'], 16)
                timing.update(closed_timestamp=stamp, elapsed_seconds=stamp-opening)
                assert (stamp < deadline if scenario == 'early' else stamp >= deadline)
            checked('observed_close_before_deadline' if scenario == 'early' else 'partial_closure_not_before_deadline')
            if scenario == 'partial':
                checked('closure_event_present')
            timing['relation_to_deadline'] = 'before 45 s' if scenario == 'early' else 'at or after 45 s'
        else:
            assert not closed, 'The empty target attempt must never close for aggregation.'
            for name, value in [('empty_round_not_published', False), ('empty_round_not_completed', False),
                                ('empty_round_has_no_inputs', 0), ('empty_round_not_closed', False),
                                ('empty_round_no_success_increment', 1), ('zero_round_never_completed', False),
                                ('zero_replacement_selected', True), ('recovery_quorum_event_proven', True)]:
                checked(name, value)
            empty = [o for o in observations if o['kind'] == 'empty_collection']
            observed = [o for o in observations if o['kind'] == 'empty_deadline_observed']
            assert empty and len(observed) == 1
            assert all(not o['value']['evidence']['published'] and not o['value']['target_completed']
                       and o['value']['policy']['accepted'] == 0 for o in empty)
            duration = condition['zero_observation_seconds']
            assert duration >= 10 and observed[0]['value'] == {'deadline': deadline, 'observation_seconds': duration}
            assert datetime.fromisoformat(observed[0]['observed_at']).timestamp() >= deadline + duration
            quorum = result['recovery_quorum']
            assert quorum['required_reports'] == 3 and quorum['eligible_reporters'] == 5
            assert quorum['below_quorum_preserved'] is True
            assert [v['report_count'] for v in quorum['votes']] == [1, 2, 3]
            assert [v['quorum_reached'] for v in quorum['votes']] == [False, False, True]
            timing['relation_to_deadline'] = f'no publication through deadline + {duration:g} s'
            details = {'quorum_required': 3, 'eligible_reporters': 5, 'below_quorum_preserved': True,
                       'empty_observation_seconds_after_deadline': duration,
                       'recovery_timings_seconds': {key: value for key, value in result['timings_seconds'].items()
                                                   if key.startswith('zero_')}}
        run_sources = [smew_path, sma_path, report / 'manifest.json', result_path, archive_path]
        provisioning = {}
        if 'provisioning_terraform_parallelism' in context or 'runtime_compose_sha256' in context:
            assert context['provisioning_terraform_parallelism'] == 2
            compose_hash = context['runtime_compose_sha256']
            assert isinstance(compose_hash, str) and len(compose_hash) == 64
            assert all(char in '0123456789abcdef' for char in compose_hash)
            plan_path = WORKSPACE / 'vita-fl-td/generated/collection-check-20260929/parallelism-plan.json'
            plan = json.loads(plan_path.read_text())
            assert plan['applied'] is True and plan['actions'] == ['update'] and not plan['new_cvm_creation']
            assert plan['container_images_unchanged'] and plan['encrypted_environment_unchanged']
            assert plan['app_id'].removeprefix('app_') == context['deployment_app_id']
            assert plan['new_compose_sha256'] == compose_hash and plan['old_compose_sha256'] != compose_hash
            assert plan['change'] == {'service': 'control-api',
                                      'environment': {'TF_CLI_ARGS_apply': '-parallelism=2'}}
            provisioning = {'provisioning_terraform_parallelism': 2, 'runtime_compose_sha256': compose_hash,
                            'previous_runtime_compose_sha256': plan['old_compose_sha256'],
                            'provisioning_config_changed_after_security_run': True,
                            'runtime_config_change': plan['change'],
                            'runtime_config_change_source': str(plan_path.relative_to(WORKSPACE))}
            run_sources.append(plan_path)
            restart_path = WORKSPACE / 'vita-fl-td/generated/collection-check-20260929/restart-verification.json'
            restart = json.loads(restart_path.read_text())
            policy_path = Path(restart['policy_path']).resolve(strict=True)
            assert policy_path.is_relative_to(WORKSPACE)
            assert restart['runtime_parallelism'] == 2 and restart['old_security_policy_preserved'] is True
            assert restart['platform_allowlist_and_sello_identity_unchanged'] is True
            assert restart['trust_bootstrap'] == 'authenticated-deployment-ui'
            assert digest(policy_path) == restart['policy_sha256']
            policy = json.loads(policy_path.read_text())
            assert hashlib.sha256(base64.b64decode(policy['scitt_keys_cbor_base64'])).hexdigest() == policy['scitt_keys_sha256']
            assert policy['scitt_keys_sha256'] == restart['new_scitt_keys_sha256'] == context['audit_log_keyset_sha256']
            assert restart['old_scitt_keys_sha256'] != restart['new_scitt_keys_sha256']
            command = condition['inference']['command']
            assert Path(command[command.index('--audit-policy') + 1]).resolve() == policy_path
            if security is not None:
                prior = security['audit_policy']
                prior_path = WORKSPACE / prior['path']
                assert digest(prior_path) == prior['sha256']
                assert prior['scitt_keys_sha256'] == restart['old_scitt_keys_sha256']
                prior_policy = json.loads(prior_path.read_text())
                for key in ('allowed_platform_measurements', 'token_issuer_public_key', 'owner_subject', 'owner_key_env'):
                    assert prior_policy[key] == policy[key], f'Unexpected trust-policy change: {key}'
            provisioning['audit_policy'] = {'path': str(policy_path.relative_to(WORKSPACE)),
                                             'sha256': digest(policy_path),
                                             'scitt_keys_sha256': policy['scitt_keys_sha256'],
                                             'previous_scitt_keys_sha256': restart['old_scitt_keys_sha256'],
                                             'trust_bootstrap': restart['trust_bootstrap'],
                                             'old_security_policy_preserved': True,
                                             'platform_allowlist_and_sello_identity_unchanged': True,
                                             'restart_verification_source': str(restart_path.relative_to(WORKSPACE))}
            run_sources.extend([restart_path, policy_path])
        hashes = {str(path.relative_to(WORKSPACE)): digest(path) for path in run_sources}
        sources.update(hashes)
        runs[scenario] = {'run_id': result['run_id'], 'run_date': run_date, 'status': 'PASS',
                          'report_directory': str(report.relative_to(WORKSPACE)),
                          'assertions_passed': len(result['checks']), 'allowed_clients': allowed_count,
                          'blocked_clients': len(blocked), 'successful_training_rounds': 5,
                          'active_workers': 6, 'attempts': attempts,
                          'contribution_counts': [r['input_count'] for r in participation['rounds']],
                          'controlled_attempt_timing': timing, 'deployment_app_id': context['deployment_app_id'],
                          'test_driver_source_sha256': context['test_driver_source_sha256'],
                          'sources': hashes, **details, **provisioning}
    assert set(runs) == {'early', 'partial', 'zero'} and len({r['run_id'] for r in runs.values()}) == 3
    failed_attempts = []
    for parent in sorted({Path(directory).expanduser().resolve().parent for directory in report_directories}):
        for report in sorted(parent.iterdir()):
            result_path = report / 'subprocess_output/result.json'
            if not result_path.is_file():
                continue
            record = json.loads(result_path.read_text())
            if record.get('scenario') not in runs or record.get('status') != 'ERROR':
                continue
            failure_sources = [path for path in [report / 'smew.json', report / 'run.json', report / 'manifest.json',
                                                  result_path, report / 'subprocess_output/run-artifacts.zip']
                               if path.is_file()]
            hashes = {str(path.relative_to(WORKSPACE)): digest(path) for path in failure_sources}
            sources.update(hashes)
            failed_attempts.append({'report_directory': str(report.relative_to(WORKSPACE)),
                                    'scenario': record['scenario'], 'run_id': record['run_id'],
                                    'status': record['status'], 'errors': record.get('errors', []),
                                    'sources': hashes})
    return {'scenarios_passed': 3, 'scenarios': {name: runs[name] for name in ('early', 'partial', 'zero')},
            'earlier_failed_attempts': failed_attempts,
            'same_deployment_images': True, 'sources': sources,
            'runtime_provisioning_config_changed': any(r.get('provisioning_config_changed_after_security_run')
                                                      for r in runs.values()),
            'coverage_limit': 'One successful execution per controlled scenario; selected collection boundaries, no reliability estimate.'}


def evidence(security_report=None, collection_reports=None):
    r = json.loads((RECOVERY / 'result.json').read_text())
    with ZipFile(RECOVERY / 'run-artifacts.zip') as z:
        audit = json.loads(z.read('transparency-audit.json'))
        events = json.loads(z.read('chain-events.json'))
    partial = json.loads(PARTIAL.read_text())
    old = json.loads(OLDER.read_text())
    assert r['status'] == 'PASS' and all(c['passed'] for c in r['checks'])
    assert audit['status'] == 'PASS' and len(audit['checks']) == 17
    assert all(c['passed'] for c in audit['checks'])
    q = r['recovery_quorum']
    assert [v['report_count'] for v in q['votes']] == [1, 2, 3]
    assert [v['quorum_reached'] for v in q['votes']] == [False, False, True]
    assert q['below_quorum_preserved'] and q['eligible_reporters'] == 5
    rounds = r['participation']['rounds']
    assert [v['input_count'] for v in rounds] == [4, 4, 5, 5, 5]
    assert r['participation']['all_workers_active'] and len(r['participation']['workers']) == 6
    assert partial['status'] == 'ERROR'
    counts = Counter(c['name'].split(':')[0] for c in audit['checks'])
    assert counts == {'scitt': 10, 'air-tdx': 1, 'session': 3, 'sello': 3}
    open_topic = '0x8bcf39476f12f2198587ea4025f7f3b26decc7e5b991c682cdab31748a79fea9'
    close_topic = '0x16a8faf1a158dfdc2ee3ddfd2e8bfe4df79f18858372c90b7d467167f2fac57f'
    policies, closes = {}, {}
    for event in events:
        if event.get('removed') or event['address'].lower() != '0x610178da211fef7d417bc0e6fed39f05609ad788':
            continue
        topic = event['topics'][0]
        if topic not in (open_topic, close_topic):
            continue
        rid = int(event['topics'][1], 16)
        words = [int(event['data'][i:i+64], 16) for i in range(2, len(event['data']), 64)]
        if topic == open_topic:
            policies[rid] = {'target': words[0], 'opened': words[1], 'deadline': words[2]}
        else:
            closes[rid] = {'accepted': words[0], 'closed': int(event['blockTimestamp'], 16)}
    observed = [{'round': rid, **policies[rid], **closes[rid],
                 'elapsed_seconds': closes[rid]['closed']-policies[rid]['opened']}
                for rid in range(2, 7)]
    assert [e['elapsed_seconds'] for e in observed] == [45, 45, 14, 13, 11]
    assert all(e['target'] == 5 and e['deadline']-e['opened'] == 45 for e in observed)
    assert [e['accepted'] for e in observed] == [4, 4, 5, 5, 5]
    sources = [RECOVERY / 'result.json', RECOVERY / 'run-artifacts.zip', PARTIAL, OLDER]
    data = {
        'review_date': '2026-09-29', 'run_date': '2026-09-25',
        'recovery_run_id': r['run_id'], 'recovery_status': r['status'],
        'successful_training_rounds': len(rounds), 'active_workers': 6,
        'contribution_counts': [v['input_count'] for v in rounds],
        'successful_attempt_ids': [v['round'] for v in rounds],
        'collection_event_timings': observed,
        'collection_timing_basis': 'archived PolicyOpened data and InputsClosed block timestamps; not direct instrumentation of the local aggregator timer',
        'configured_client_limit': r['condition']['training']['client_limit'],
        'configured_deadline_ms': r['condition']['training']['model_submission_deadline_ms'],
        'quorum_reports': [v['report_count'] for v in q['votes']],
        'required_reports': q['required_reports'], 'eligible_reporters': q['eligible_reporters'],
        'below_quorum_preserved': q['below_quorum_preserved'],
        'reselection_seconds': r['timings_seconds']['aggregator_reselection'],
        'publication_seconds': r['timings_seconds']['recovery_publication'],
        'audit_passed': len(audit['checks']), 'audit_groups': dict(counts),
        'audit_duration_seconds': audit['duration_seconds'], 'audit_limitations': audit['limitations'],
        'partial_scenario_status': partial['status'], 'older_recovery_status': old['status'],
        'unobserved_scenarios': ['early', 'zero', 'security'],
        'unobserved_scope': 'No archived execution found in workspace SMEW/vita-fl-td results, including ignored files; this is not proof that no execution ever occurred.',
        'security_checks_implemented': 16,
        'security_inventory_source': 'presentation/concepts/test-evaluation/ci-external/external-scope.json',
        'sources': {str(p.relative_to(WORKSPACE)): digest(p) for p in sources},
        'new_live_tests_executed': False,
    }
    if security_report is not None:
        security = completed_security_evidence(security_report, r)
        data['security'] = security
        data['sources'].update(security['sources'])
        data['unobserved_scenarios'].remove('security')
        data['security_checks_observed_passed'] = security['checks_passed']
        data['new_live_tests_executed'] = True
        data['live_execution_note'] = 'A separately executed 29 Sep security campaign is consumed; this slide generator launches no cloud run.'
        data['displayed_audit_source'] = security['run_id']
    else:
        data['displayed_audit_source'] = r['run_id']
    if collection_reports is not None:
        collection = completed_collection_evidence(collection_reports, r, data.get('security'))
        data['collection'] = collection
        data['sources'].update(collection['sources'])
        data['unobserved_scenarios'] = [s for s in data['unobserved_scenarios'] if s not in {'early', 'zero'}]
        data['new_live_tests_executed'] = True
        data['live_execution_note'] = 'Separately executed 29 Sep campaigns are consumed; this slide generator launches no cloud run.'
    return data


DATA = {}
NOTES = ''
BASE_NOTES = '''EXTERNAL RESULTS — FOLLOW-UP TO SLIDE 34
Prepared 29 September 2026 from archived results dated 25 September 2026.
This is evidence from the SMEW/SMA/vita-fl-td workflow, separate from the earlier
24-round model-quality experiment. No new campaign was launched for these slides.
The four categories follow the last four external rows on slide 34.

COLLECTION / TIMEOUT
In the successful recovery run, successful training attempts 2,3,4,5,6 had
4,4,5,5,5 accepted client contributions. The configured client limit was five
and the collection deadline 45,000 ms. Attempts 2 and 3 were completed and
published despite having only four contributions. Archived PolicyOpened and
InputsClosed events show closure after 45 s for both four-input attempts, at
their recorded deadlines. The five-input attempts closed after 14,13,11 s,
before the deadline. Timings are reconstructed from recorded chain events,
not direct instrumentation of the aggregator's local timer. These observations
are not standalone passes for the controlled early/partial/zero scenario suite. The separate partial scenario was canceled (SIGTERM), status ERROR,
before reaching the target assertion. No early/zero result archive was found.

AGGREGATOR RECOVERY
A selected aggregator CVM was stopped. Three distinct eligible reporters out
of five produced votes 1,2,3 against a threshold of three. Historical block
states and ordered contract events preserve the old aggregator and un-aborted
round at votes one and two. Vote three triggers abort and reselection in the
same transaction. The failed attempt never published. The replacement published
and five successful training rounds completed, excluding bootstrap and the
aborted attempt. All six workers participated across the run, not necessarily
in every round. The stopped CVM was restored; provider running was confirmed.
Application readiness of that individual CVM was not independently checked;
final deployed-agent inference succeeded.

TIMING
Use the recorded TD durations, rounded for display: reselection 112.38 s,
first recovery publication 169.41 s, both from confirmed VM stop. Both belong to the same recovery interval;
they are not averaged latencies, service guarantees, or the aggregator's 45 s
collection timer. The process diagram is schematic, not a time-scaled axis.

SECURITY / ACCESS AND CONTRACT GUARDS
The 16 checks on slide 34 are implemented checks, not observed pass results.
No archived security scenario result was found in the searched workspace.
The absence of a report is neither a pass nor a failed security assertion.
Read-only eth_call contract probes and authorization negatives would require
their own completed security run before their outcomes can be claimed.

PUBLISHED EVIDENCE AUDIT
The recovery archive contains a successful 17-check audit lasting 13.5826 s:
10 SCITT publisher/inclusion checks, 1 combined AIR/TDX check,
3 attestation-session checks and 3 Sello receipt checks. Checks were re-executed
on published transparency-log evidence with run/model/receiver policy bindings.
The three tools were fetch_latest_verified_tee_model_bundle,
generate_random_tee_chestmnist_image and run_and_verify_tee_inference.
Fresh quote cryptography/collateral appraisal is delegated to Phala. Verifier
primitives are reused from the recorded VITA checkout: independent execution,
not an independently implemented Intel DCAP verifier. Log signatures establish
inclusion under the pinned key, not global log completeness or no split views.
The original TLS/PoP handshake was not replayed. Seventeen checks are assertions
within this one audit, not seventeen independent experiments.

SELECTION AND LIMITS
The slides highlight the latest completed successful recovery run. An earlier
recovery attempt ended ERROR; the partial test also ended ERROR after cancel.
No campaign success percentage, repeated-run reliability estimate, broad
security guarantee or complete Chapter-4 threat coverage is inferred.
Green means observed/passed assertion; amber means incomplete scope;
grey means no archived execution found. The DFL/Agent header chips retain the
main presentation's functional-domain meaning. Chapter 4 and the main deck
are preserved; these are separate editable proposals awaiting selection.

SOURCE MAP
'''
SOURCE_NOTES = '''
presentation/concepts/test-evaluation/ci-external/external-scope.json
vita-fl-td/src/vita_fl_td/runner.py
vita-fl-td/src/vita_fl_td/quorum.py
vita-fl-td/src/vita_fl_td/transparency_audit.py
All numeric slide results are exported with source SHA-256 values in evidence.json.
'''



def notes(data):
    value = BASE_NOTES
    if 'security' in data:
        security = data['security']
        value = value.replace(
            'Prepared 29 September 2026 from archived results dated 25 September 2026.',
            'Prepared from recovery on 25 September and security on 29 September 2026.')
        value = value.replace(
            '24-round model-quality experiment. No new campaign was launched for these slides.',
            '24-round model-quality experiment. The security campaign ran on a newly deployed\nPhala environment; the generator itself launches no campaign. All seven deployed\nimage references match the recovery run, but deployment identities, run dates and\ntest-driver source hashes differ. These are two distinct runs, not a combined run.')
        start = value.index('The 16 checks on slide 34')
        end = value.index('\nPUBLISHED EVIDENCE AUDIT', start)
        value = value[:start] + f'''The completed security run {security['run_id']} passed all 16 probes:
five access/configuration checks (UI missing/wrong credentials, Control
missing/wrong credentials, unchanged configuration); four read-only eth_call
contract guards (stale round, wrong aggregator, unregistered action key, policy
owner); one positive deployed-agent inference control; and six unauthorized
inference requests (missing/spoofed authorization on three endpoints).
Ten requests returned HTTP 401: two UI, two Control and six inference requests,
with the expected authentication challenge or error reason. All four eth_call
probes reverted with the expected reason. The authorized positive control is
one successful agent workflow, not proof of model prediction correctness.
The unchanged-configuration check is an invariant, not a fifth HTTP rejection.
SMEW succeeded with exit code zero; SMA reports ok; TD and its complete security
matrix report PASS. Five training rounds completed with all six workers active
across this security run. No contract attack transaction was broadcast.
These bounded authorization/contract probes do not cover every Chapter-4 threat;
the archived coverage_gaps retain the untested and residual threats.
''' + value[end:]
        value = value.replace(
            'The recovery archive contains a successful 17-check audit lasting 13.5826 s:',
            f"The slide displays the security run's successful 17-check audit lasting {security['audit_duration_seconds']:.4f} s.\nThe earlier recovery audit separately passed 17 checks in 13.5826 s. Counts\nare not pooled; the two audits are separate observations. Each audit contains:")
        value = value.replace(
            'Green means observed/passed assertion; amber means incomplete scope;\ngrey means no archived execution found.',
            'Green means observed/passed assertion; amber means incomplete scope.\nCollection and recovery results are from 25 Sep; security and the displayed\naudit are from 29 Sep on a separate deployment with identical pinned images.')
    if 'collection' in data:
        collection = data['collection']
        value = value.replace('Prepared from recovery on 25 September and security on 29 September 2026.',
                              'Prepared from recovery on 25 September and collection/security on 29 September 2026.')
        value = value.replace('Prepared 29 September 2026 from archived results dated 25 September 2026.',
                              'Prepared from recovery on 25 September and collection on 29 September 2026.')
        start = value.index('In the successful recovery run, successful training attempts')
        end = value.index('\nAGGREGATOR RECOVERY', start)
        lines = []
        for scenario, run in collection['scenarios'].items():
            timing = run['controlled_attempt_timing']
            measured = (f"; archived event interval {timing['elapsed_seconds']} s"
                        if 'elapsed_seconds' in timing else '')
            lines.append(f"{scenario.upper()}: run {run['run_id']}; "
                         f"{run['allowed_clients']} allowed / {run['blocked_clients']} blocked clients; "
                         f"{timing['relation_to_deadline']}{measured}.")
        value = value[:start] + """Three separate controlled scenarios completed on 29 September 2026.
Each used six workers, client limit five, a 45,000 ms collection deadline and
five successful training rounds (bootstrap and aborted attempts excluded).
EARLY: exactly five allowed contributors published before the deadline.
PARTIAL: exactly two allowed contributors published at or after the deadline;
the receiver confirmed blocking of the other three clients.
ZERO: all five clients were blocked, with receiver-confirmed blocked attempts.
The target attempt neither closed nor published, including the observation
window through at least ten seconds after its deadline. Three of five eligible
reports then triggered abort and reselection; votes one and two preserved the
state. The gate was released and five successful training rounds completed.
All six workers participated across each run; this does not mean every worker
contributed to every controlled round. No empty-model publication is claimed.
Each run passed SMEW, SMA and TD completion checks. The slide's 3/3 count means
three distinct scenarios, each executed successfully once, not three checks
or a repeated-run reliability statistic. Gate accounting, actual contributors,
closure timing, zero-input nonpublication and quorum evidence are validated
against the canonical result and archived artifacts before PASS is displayed.
Timing values use archived policy/chain events, not instrumentation of the
aggregator's local timer. Early is also checked by the driver's first observation;
partial uses the closure transaction timestamp as an independent oracle.
""" + '\n'.join(lines) + '\n' + value[end:]
        value = value.replace(
            'Collection and recovery results are from 25 Sep; security and the displayed\naudit are from 29 Sep on a separate deployment with identical pinned images.',
            'Recovery is from 25 Sep; collection, security and the displayed audit are\nfrom separate runs on 29 Sep. Image pins are identical across these runs.')
        value = value.replace('These are two distinct runs, not a combined run.',
                              'Collection, recovery and security are distinct runs, not a combined run.')
        value = value.replace('recovery attempt ended ERROR; the partial test also ended ERROR after cancel.',
                              'recovery attempt ended ERROR; an earlier partial test also ended ERROR after cancel.')
        value = value.replace('Green means observed/passed assertion; amber means incomplete scope.',
                              'Green means observed/passed assertion within the selected scenarios.')
        if collection['runtime_provisioning_config_changed']:
            provisioning = next(r for r in collection['scenarios'].values()
                                if r.get('provisioning_config_changed_after_security_run'))
            value += f'''
RUNTIME PROVISIONING DIFFERENCE
Identical container-image pins do not imply identical runtime configuration.
After the security run, the existing runtime app was updated in place: the
control-api environment now sets TF_CLI_ARGS_apply=-parallelism=2. This reduces
Terraform provisioning concurrency after Phala HTTP 429 rate-limit failures.
Runtime app identity, all seven image references, OS and encrypted environment
were retained. This is a configuration change, not a newly identified deployment.
Old runtime Compose SHA-256: {provisioning['previous_runtime_compose_sha256']}
New runtime Compose SHA-256: {provisioning['runtime_compose_sha256']}
The reviewed/applied plan and its SHA-256 are included in the source map.
The runtime restart also changed the transparency log's SCITT key set. A new
collection audit policy was pinned through the authenticated deployment UI:
{provisioning['audit_policy']['path']}
Policy SHA-256: {provisioning['audit_policy']['sha256']}
SCITT key-set SHA-256: {provisioning['audit_policy']['scitt_keys_sha256']}
The original security audit policy remains preserved for the earlier run.
The platform-measurement allowlist and Sello identity remain unchanged. These
statements are checked against restart-verification.json and both policy files
when the security archive is included. Policy/key-set and verification-file
hashes are retained; the later key set is not substituted into the old audit.
Earlier collection startup attempts ended ERROR (Control HTTP 500 after Phala
429 provisioning failure; one early startup returned HTTP 409). They were not
collection behavior passes and are not erased by successful later attempts.
No campaign-wide 100 percent success rate or same-configuration comparison is
claimed. Per-run provisioning settings and source hashes remain in evidence.json.
'''
        for attempt in collection['earlier_failed_attempts']:
            value += f"\nRetained earlier ERROR: {attempt['scenario']} / {attempt['run_id']} / {attempt['report_directory']}\n"
        value += "\nCOLLECTION SCOPE: The later passing scenarios do not erase earlier aborted attempts.\nThe chosen collection boundaries are covered; arbitrary failures and statistical reliability are not.\n"
    return value + '\n'.join(data['sources']) + SOURCE_NOTES


def security_observed():
    return 'security' in DATA


def collection_observed():
    return 'collection' in DATA


def collection_timing_label():
    scenarios = DATA['collection']['scenarios']
    values = []
    for name, count, fallback in [('early', 5, 'before 45 s'), ('partial', 2, 'at or after 45 s')]:
        timing = scenarios[name]['controlled_attempt_timing']
        label = f"at {timing['elapsed_seconds']} s" if 'elapsed_seconds' in timing else fallback
        values.append(f'{count}/5 {label}')
    return ' · '.join(values)


def audit_seconds():
    return DATA.get('security', DATA)['audit_duration_seconds']


def txt(s, value, x, y, w, h=.32, size=15, color=INK, bold=False, center=False):
    q = b.add_text(s, value, x, y, w, h, size, color, bold,
        align=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0)
    q.name = 'External results: ' + value.replace('\n', ' / ')
    return q


def box(s, x, y, w, h, fill=WHITE, line=LINE):
    return b.add_box(s, x, y, w, h, fill=fill, line=line, radius=False, line_width=.8)


def rule(s, x, y, w):
    b.add_line_segment(s, x, y, x+w, y, color=LINE, width=.7)


def tag(s, value, x, y, w, color, tint):
    box(s, x, y, w, .28, tint, tint)
    txt(s, value, x+.07, y+.01, w-.14, .26, 9.7, color, True, True)


def page(p, letter, title):
    s = b.new_content_slide(p, len(p.slides)+1, title,
        'EVALUATION · EXTERNAL RESULTS · 25 + 29 SEPTEMBER 2026' if security_observed() or collection_observed()
        else 'EVALUATION · OBSERVED EXTERNAL RESULTS · 25 SEPTEMBER 2026')
    b.add_domain_navigation(s, 35, active_domains={'DFL','Agent'})
    for q in s.shapes:
        if q.has_text_frame and q.text.startswith('Page '):
            q.text_frame.paragraphs[0].runs[0].text = 'Variant '+letter
            q.width = b.Inches(1.1)
    b.add_note(s, 'VARIANT '+letter+'\n\n'+NOTES)
    return s


def footer(s, takeaway):
    rule(s, .72, 5.73, 11.9)
    txt(s, takeaway, .75, 5.79, 11.81, .28, 11.9, INK, True)
    if collection_observed():
        detail = ('25 Sep: recovery · 29 Sep: collection/security/audit · Separate runs, same images · Green: passed.'
                  if security_observed() else
                  '25 Sep: recovery/audit · 29 Sep: collection · Separate runs, same images · Green: passed · Grey: no run record.')
        if DATA['collection'].get('runtime_provisioning_config_changed'):
            detail = detail.replace('Separate runs, same images', 'Same images; provisioning config changed')
        txt(s, detail, .75, 6.10, 11.8, .18, 9.0, MUTED)
    elif security_observed():
        txt(s, '25 Sep: collection/recovery · 29 Sep: security/audit · Separate runs, same images · Green: passed · Amber: incomplete.',
            .75, 6.10, 11.8, .18, 9.0, MUTED)
    else:
        txt(s, 'Single recovery run; earlier recovery ERROR. Green: observed/pass · Amber: incomplete scope · Grey: no run record.',
            .75, 6.10, 11.8, .18, 9.0, MUTED)


import revised_layouts as layouts


def design_a(p):
    return layouts.design_a(p, sys.modules[__name__])


def design_b(p):
    return layouts.design_b(p, sys.modules[__name__])


def design_c(p):
    return layouts.design_c(p, sys.modules[__name__])


DESIGNS = [
    ('a-result-cards', 'A · Vier Testfragen', 'Jede Kategorie zeigt die konkrete Testfrage, den Eingriff und die beobachtete Antwort.', design_a),
    ('b-evidence-matrix', 'B · Konkreter Testkatalog', 'Einzelne Testfälle: Was wurde eingegeben oder gesperrt, und welche Eigenschaft wurde überprüft?', design_b),
    ('c-recovery-story', 'C · Vier Prüfpfade', 'Für jede Kategorie ein eigener Weg: Eingriff oder Evidenz → geprüftes System → Ergebnis.', design_c),
]


def main():
    global DATA, NOTES
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--security-report', type=Path, required=True,
        help='Completed canonical SMEW report directory containing smew.json, run.json and subprocess_output/.')
    parser.add_argument('--collection-reports', type=Path, nargs=3, required=True,
        help='Three completed canonical SMEW report directories: one each for early, partial and zero (any order).')
    args = parser.parse_args()
    DATA = evidence(args.security_report, args.collection_reports)
    NOTES = notes(DATA) + '''

REVISED TEST SCOPE
Each visible category identifies its actual intervention or downloaded evidence and the property checked.
Collection numbers describe input closure, not a directly instrumented aggregation-start time.
The configuration probe submits the existing configuration, checks rejection and unchanged state; it does not attempt a changed configuration.
Receipt audit uses authentic published evidence, without receipt mutation attacks.
RTMR3 means event-log replay against each quote, not registration-versus-inference equality; model hashes are bound through manifest, REPORTDATA and AIR.
Sello checks include signature, owner authorization, exact tool input/output and attested receiver/session binding.
No model-accuracy claim, independent cryptographic implementation, original TLS-handshake observation or global-log completeness claim.
'''
    protected=[ROOT/'VITA-FL_Thesis_Presentation_TU_Berlin.pptx',WORKSPACE/'overleaf/chapters/chapter4.tex']
    before={str(p.relative_to(WORKSPACE)):digest(p) for p in protected}
    p=b.prepare_template()
    p.core_properties.title='VITA-FL — External evaluation results: three alternatives'
    for _,_,_,build in DESIGNS: build(p)
    check_layout(p)
    for slide in p.slides:
        for shape in slide.shapes:
            if shape.name.startswith('External results: '):
                assert shape.top + shape.height <= b.Inches(6.29), 'Content enters template footer'
    output=HERE/'VITA-FL_External_Results_Alternatives.pptx'
    p.save(output)
    q=Presentation(output)
    check_layout(q)
    previews=[]
    for slide,(slug,label,desc,build) in zip(q.slides,DESIGNS):
        visible='\n'.join(s.text for s in slide.shapes if s.has_text_frame)
        assert '17 / 17' in visible and '16' in visible
        if security_observed():
            assert '16 / 16' in visible and 'NO RUN RECORD' not in visible
            assert '25 Sep' in visible and '29 Sep' in visible
            assert 'HTTP 401' in visible and 'eth_call' in visible
            assert 'config unchanged' in visible.lower() and 'authorized inference passed' in visible.lower()
        if collection_observed():
            assert '3 / 3' in visible and 'PASS' in visible
            assert 'PARTIAL EVIDENCE' not in visible and 'incomplete' not in visible.lower()
            assert 'zero' in visible.lower() and 'publication' in visible.lower()
        for scope_term in ('uploads', 'credentials', 'stale', 'owner', 'SCITT', 'Sello', 'RTMR3', 'inclusion'):
            assert scope_term.lower() in visible.lower(), f'Missing concrete test scope: {scope_term}'
        assert 'SOURCE MAP' in slide.notes_slide.notes_text_frame.text
        im=renderer.render_slide(q,slide)
        im.save(HERE/(slug+'.png'))
        previews.append(im)
        single=b.prepare_template(); build(single); check_layout(single)
        single.save(HERE/(slug+'.pptx'))
        assert len(Presentation(HERE/(slug+'.pptx')).slides)==1
    previews[0].save(HERE/'VITA-FL_External_Results_Alternatives.pdf',save_all=True,append_images=previews[1:],resolution=120)
    sheet=Image.new('RGB',(1660,1560),'#E8ECEF')
    draw=ImageDraw.Draw(sheet)
    font=ImageFont.truetype(renderer.FONT_BOLD,26)
    body=ImageFont.truetype(renderer.FONT_REGULAR,24)
    descriptions = layouts.DESCRIPTIONS
    for i,(im,(_,label,_,_)) in enumerate(zip(previews,DESIGNS)):
        y=15+i*515
        draw.text((30,y),label,font=font,fill='#263440')
        sheet.paste(im.resize((800,450),Image.Resampling.LANCZOS),(30,y+42))
        for j,t in enumerate(descriptions[i]): draw.text((875,y+125+j*66),t,font=body,fill='#263440' if j==0 else '#65737D')
    sheet.save(HERE/'overview.png')
    cards=[]
    for slug,label,desc,_ in DESIGNS:
        cards.append(f'<article><h2>{html.escape(label)}</h2><p>{html.escape(desc)}</p><a href="{slug}.png"><img src="{slug}.png" alt="{html.escape(label)}"></a><p><a href="{slug}.pptx">PowerPoint</a> · <a href="{slug}.png">Vollbild</a></p></article>')
    (HERE/'index.html').write_text('''<!doctype html><html lang="de"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>VITA-FL – Externe Testergebnisse</title><style>body{font:16px/1.5 system-ui;background:#edf1f4;color:#263440;margin:0}header,main{max-width:1500px;margin:auto;padding:24px}article{background:white;padding:24px;border-radius:12px;margin-bottom:28px}img{width:100%;border:1px solid #dce3e8}a{color:#1764a1}h1,h2{margin:0}nav{display:flex;gap:24px;flex-wrap:wrap}</style><header><h1>Externe Testergebnisse: drei Folienvorschläge</h1><p>'''+(('Recovery: 25.09.2026; Collection/Security/Audit: 29.09.2026. Separate Läufe mit identischen Image-Pins.' if security_observed() else 'Recovery/Audit: 25.09.2026; Collection: 29.09.2026. Drei kontrollierte Szenarien bestanden.') if collection_observed() else 'Recovery/Collection: 25.09.2026; Security/Audit: 29.09.2026. Neues Deployment mit identischen Image-Pins.' if security_observed() else 'Ergebnisse aus gespeicherten Läufen vom 25.09.2026.')+''' Die vier Bereiche entsprechen den letzten vier externen Punkten auf Folie 34. Die Folien stehen als editierbare Alternativen bereit; den Integrationsstand dokumentiert die README.</p><nav><a href="VITA-FL_External_Results_Alternatives.pptx">Alle Varianten als PowerPoint</a><a href="VITA-FL_External_Results_Alternatives.pdf">PDF</a><a href="README.md">Einordnung und Quellen</a></nav></header><main>'''+''.join(cards)+'</main></html>',encoding='utf-8')
    (HERE/'evidence.json').write_text(json.dumps(DATA,indent=2,ensure_ascii=False)+'\n')
    assert before=={str(p.relative_to(WORKSPACE)):digest(p) for p in protected}
    (HERE/'validation.json').write_text(json.dumps({'slides':3,'native_editable_shapes':True,'layout':'passed','pptx_roundtrip':'passed','assertions_against_archived_results':'passed','main_and_chapter_preserved':before,'new_live_tests_executed':DATA['new_live_tests_executed'],'security_archive':DATA.get('security',{}).get('report_directory'),'collection_archives':{name:run['report_directory'] for name,run in DATA.get('collection',{}).get('scenarios',{}).items()},'visual_review':'pending','design_revision':'concrete-test-scope-v2','preview_renderer':'Pillow; native PowerPoint may vary slightly in font metrics'},indent=2)+'\n')
    print('Built three editable proposals; layout and archived-result assertions passed. Main deck and Chapter 4 unchanged.')


if __name__=='__main__': main()
