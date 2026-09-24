#!/usr/bin/env python3
"""Preserve lifecycle A and explore two variations of its circle-and-branch layout."""

from __future__ import annotations

import hashlib
import json
import math
from pathlib import Path
import shutil
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

HERE = Path(__file__).resolve().parent
SOURCE = HERE.parent
ROOT = HERE.parents[2]
sys.path.insert(0, str(SOURCE))
import build_concepts as c

TITLE = 'VITA-FL: from training to verifiable use'


def loop_with_model_exit(s, cx, cy, radius, model_color=c.GREEN):
    """Counterclockwise training -> aggregation -> global model -> training."""
    c.circle(s, cx, cy, radius, c.GREEN, c.WHITE, 2.25)
    # Arrowheads sit on exposed sections of the circle, between the three nodes.
    for angle in [-178, 54, -61]:
        c.arc(s, cx, cy, radius, angle+12, angle, c.GREEN, 2.3)
    for label, angle, width, color, height in [
        ('Local training', -120, 2.16, c.GREEN, .63),
        ('Aggregate', 120, 1.77, c.GREEN, .63),
        ('Global\nmodel', 0, 1.80, model_color, .90),
    ]:
        x=cx+radius*math.cos(math.radians(angle))
        y=cy+radius*math.sin(math.radians(angle))
        c.pill(s,label,x,y,width,color,17,height,
               c.TINT[color] if color==c.PURPLE else c.WHITE)
    c.text(s,'DFL',cx-.76,cy-.25,1.34,.50,26,c.GREEN,True)
    return cx+radius+.94,cy


def agent_above(s, cx, inference_y, agent_y=1.98):
    c.icon_node(s,'agent','Agent',cx,agent_y,c.BLUE,19)
    start_y=agent_y+1.02
    end_y=inference_y-.59
    c.route(s,[(cx-.25,start_y),(cx-.25,end_y)],c.BLUE,2.1)
    label_y=(start_y+end_y)/2-.15
    c.text(s,'MCP',cx-.98,label_y,.65,.30,16,c.BLUE,True)
    c.route(s,[(cx+.27,end_y),(cx+.27,start_y)],c.BLUE,1.5)
    c.text(s,'Result',cx+.41,label_y,.79,.30,14,c.MUTED)


def variant_a1(prs):
    s=c.page(prs,'A1',TITLE,
             'A1: The global model is the visible exit of the training circle. '
             'A separate model symbol denotes the published version handed to inference. '
             'This keeps the original layout while making the publication point explicit.')
    y=4.15
    end=loop_with_model_exit(s,2.25,y,1.25)
    c.route(s,[end,(5.49,y)],c.PURPLE,2.2)
    c.icon_node(s,'model','Published\nmodel',6.03,y,c.PURPLE,19)
    c.route(s,[(6.56,y),(7.76,y)],c.PURPLE,2.2)
    c.pill(s,'Attested\ninference',8.80,y,1.98,c.BLUE,20,1.05)
    c.route(s,[(9.84,y),(11.18,y)],c.BLUE,2.1)
    c.text(s,'Tool receipt',9.88,y-.49,1.23,.33,13.5,c.BLUE)
    c.icon_node(s,'log','Transparency\nlog',11.77,y,c.BLUE,18)
    agent_above(s,8.80,y)
    return s


def variant_a2(prs):
    s=c.page(prs,'A2',TITLE,
             'A2: Compact boundary handoff. The global-model node lies on the circle boundary. '
             'The outgoing Publish arrow represents selecting a finalized model version for use. '
             'There is no second published-model box. The purple node and arrow represent the model handoff, not a new actor.')
    y=4.25
    end=loop_with_model_exit(s,2.68,y,1.45,c.PURPLE)
    c.route(s,[end,(7.60,y)],c.PURPLE,2.4)
    c.text(s,'Publish',5.70,y-.49,1.18,.35,17,c.PURPLE)
    c.pill(s,'Attested\ninference',8.64,y,1.98,c.BLUE,20,1.05)
    c.route(s,[(9.68,y),(11.06,y)],c.BLUE,2.1)
    c.text(s,'Tool receipt',9.74,y-.49,1.23,.33,13.5,c.BLUE)
    c.icon_node(s,'log','Transparency\nlog',11.65,y,c.BLUE,18)
    agent_above(s,8.64,y)
    return s


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    protected=[p for p in SOURCE.iterdir() if p.is_file()]
    protected.append(ROOT/'VITA-FL_Thesis_Presentation_TU_Berlin.pptx')
    before={p:digest(p) for p in protected}
    # Load the actual previous deck: original A's shapes, text, and notes are retained.
    prs=Presentation(SOURCE/'VITA-FL_Lifecycle_Concepts.pptx')
    original_xml=str(prs.slides[0]._element.xml)
    while len(prs.slides)>1:
        c.b.delete_slide(prs,len(prs.slides)-1)
    prs.core_properties.title='VITA-FL — Original A and lifecycle variations'
    slides=[prs.slides[0],variant_a1(prs),variant_a2(prs)]
    assert str(slides[0]._element.xml)==original_xml
    destination=HERE/'VITA-FL_Lifecycle_A_Variants.pptx'
    prs.save(destination)
    restored=Presentation(destination)
    assert str(restored.slides[0]._element.xml)==original_xml

    names=['a0-original','a1-model-handoff','a2-compact']
    previews=[]
    for i,(name,slide) in enumerate(zip(names,slides)):
        image=c.renderer.render_slide(prs,slide)
        if i==0:
            original=Image.open(SOURCE/'a-loop-and-branch.png').convert('RGB')
            assert image.tobytes()==original.tobytes(), 'Original A preview changed'
            shutil.copy2(SOURCE/'a-loop-and-branch.png',HERE/(name+'.png'))
        else:
            image.save(HERE/(name+'.png'))
        previews.append(image)

    w,h,gutter,caption=800,450,30,53
    sheet=Image.new('RGB',(3*w+4*gutter,h+caption+2*gutter),'#E9EDF0')
    draw=ImageDraw.Draw(sheet)
    font=ImageFont.truetype('/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf',25)
    labels=['A0 · Original','A1 · Modellübergabe mit Symbol',
            'A2 · Kompakter Übergang']
    for i,(preview,label) in enumerate(zip(previews,labels)):
        x=gutter+i*(w+gutter)
        y=gutter
        draw.text((x+5,y+8),label,font=font,fill='#263440')
        sheet.paste(preview.resize((w,h),Image.Resampling.LANCZOS),(x,y+caption))
    sheet.save(HERE/'overview.png')
    assert all(digest(p)==value for p,value in before.items()), 'A protected original changed'
    (HERE/'preserved-originals.json').write_text(json.dumps(
        {str(p.relative_to(ROOT)):value for p,value in before.items()},indent=2)+'\n')
    print('Created original A + two editable variants; original files and first-slide XML/preview verified unchanged.')
    print(destination)


if __name__=='__main__':
    main()
