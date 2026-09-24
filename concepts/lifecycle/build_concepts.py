#!/usr/bin/env python3
"""Four editable lifecycle proposals, separate from the main presentation."""

from __future__ import annotations

import hashlib
import math
from pathlib import Path
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))
import build_presentation as b
import render_preview as renderer

GREEN, BLUE, PURPLE = '207548', '1764A1', '7446A6'
INK, MUTED, LINE = '434343', '67747C', 'DAE1E5'
WHITE = 'FFFFFF'
TINT = {GREEN: 'EDF6F0', BLUE: 'EDF4FA', PURPLE: 'F3EEF8'}


def text(s, value, x, y, w, h=.38, size=18, color=INK, bold=False, center=True):
    shape = b.add_text(s, value, x, y, w, h, size, color, bold,
                       align=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT,
                       valign=MSO_ANCHOR.MIDDLE, margin=0)
    shape.name = 'Label: ' + value.replace('\n', ' / ')
    return shape


def box(s, x, y, w, h, color, fill=None, width=1.4):
    return b.add_box(s, x, y, w, h, fill=fill or TINT[color],
                     line=color, line_width=width)


def circle(s, cx, cy, radius, color, fill=WHITE, width=1.6):
    return b.add_auto_shape(s, MSO_AUTO_SHAPE_TYPE.OVAL,
                           cx-radius, cy-radius, 2*radius, 2*radius,
                           fill=fill, line=color, line_width=width)


def route(s, points, color=BLUE, width=2.1, arrow=True, name='Flow'):
    for i, (a, z) in enumerate(zip(points, points[1:])):
        part = b.add_line_segment(s, *a, *z, color, width,
                                  arrow=arrow and i == len(points)-2)
        part.name = name


def arc(s, cx, cy, radius, start, end, color, width=2.4, arrow=True):
    steps = max(8, round(abs(end-start)/4))
    points = [(cx+radius*math.cos(math.radians(start+(end-start)*i/steps)),
               cy+radius*math.sin(math.radians(start+(end-start)*i/steps)))
              for i in range(steps+1)]
    route(s, points, color, width, arrow, name='Editable circular flow')


def pill(s, value, cx, cy, w, color, size=17, h=.62, fill=None):
    shape = box(s, cx-w/2, cy-h/2, w, h, color, fill)
    shape.name = 'Node: ' + value.replace('\n', ' / ')
    text(s, value, cx-w/2+.07, cy-h/2+.03, w-.14, h-.06,
         size, color, True)
    return shape


def model_icon(s, cx, cy, scale=1, color=PURPLE):
    layers = [[(cx-.29*scale, cy-.22*scale), (cx-.29*scale, cy+.22*scale)],
              [(cx, cy-.31*scale), (cx, cy), (cx, cy+.31*scale)],
              [(cx+.29*scale, cy-.15*scale), (cx+.29*scale, cy+.15*scale)]]
    for a, z in zip(layers, layers[1:]):
        for p in a:
            for q in z:
                route(s, [p, q], color, 1.1, False)
    for layer in layers:
        for x, y in layer:
            circle(s, x, y, .055*scale, color, WHITE, 1.2)


def agent_icon(s, cx, cy, scale=1):
    box(s, cx-.26*scale, cy-.20*scale, .52*scale, .39*scale, BLUE, WHITE)
    for dx in [-.105, .105]:
        circle(s, cx+dx*scale, cy-.035*scale, .033*scale, BLUE, BLUE, .8)
    route(s, [(cx-.09*scale, cy+.10*scale), (cx+.09*scale, cy+.10*scale)], BLUE, 1.2, False)
    route(s, [(cx,cy-.20*scale),(cx,cy-.31*scale)], BLUE, 1.2, False)
    circle(s, cx, cy-.35*scale, .04*scale, BLUE, WHITE, 1.1)


def log_icon(s, cx, cy, scale=1):
    for dx, dy in [(.10,-.10),(.05,-.05),(0,0)]:
        b.add_box(s, cx-.25*scale+dx*scale, cy-.31*scale+dy*scale,
                  .50*scale, .62*scale, WHITE, BLUE, radius=False, line_width=1.2)
    for dy in [-.14,0,.14]:
        route(s, [(cx-.15*scale,cy+dy*scale),(cx+.14*scale,cy+dy*scale)],BLUE,1.2,False)


def icon_node(s, kind, label, cx, cy, color=BLUE, size=19):
    circle(s, cx, cy, .48, color, TINT[color], 1.2)
    {'agent': agent_icon, 'model': model_icon, 'log': log_icon}[kind](s,cx,cy,.85)
    text(s, label, cx-1.15, cy+.61, 2.3, .72 if '\n' in label else .38,
         size, color, True)


def dfl_loop(s, cx, cy, radius=1.23, compact=False):
    """Clockwise: local training -> aggregation -> global model -> training."""
    if compact:
        for start in [-77,43,163]:
            arc(s,cx,cy,radius,start,start+92,GREEN,2.1)
        for angle in [-90,30,150]:
            x=cx+radius*math.cos(math.radians(angle))
            y=cy+radius*math.sin(math.radians(angle))
            circle(s,x,y,.15,GREEN,WHITE,1.4)
        text(s,'DFL',cx-.62,cy-.23,1.24,.46,20,GREEN,True)
        return
    for start in [-72,52]:
        arc(s,cx,cy,radius,start,start+77,GREEN,2.5)
    arc(s,cx,cy,radius,172,223,GREEN,2.5)
    arc(s,cx,cy,radius,223,249,GREEN,2.5,False)
    for label, angle, w in [('Local training',-90,1.88),
                            ('Aggregate',30,1.68),('Global model',150,1.83)]:
        x=cx+radius*math.cos(math.radians(angle))
        y=cy+radius*math.sin(math.radians(angle))
        pill(s,label,x,y,w,GREEN,16,.61,WHITE)
    text(s,'DFL',cx-.74,cy-.23,1.48,.48,26,GREEN,True)
    return cx+radius*math.cos(math.radians(30))+.84, cy+radius*.5


COMMON_NOTES = (
    'Conceptual lifecycle overview, intended before the existing architecture slide (currently page 9). '
    'Green is decentralized model production; purple is the published model handoff; blue is agent-mediated use and evidence. '
    'The DFL loop represents repeated local training, aggregation, and distribution of the next global model. '
    'Aggregation is a participant-assigned role, not a permanent central server. '
    'A finalized publication can be used for inference while training can continue in later rounds. '
    'The model is consumed by the inference service. The agent invokes tools through MCP; the agent does not execute the trained model itself. '
    'The inference service executes the admitted native workload. Its returned result and domain evidence bind model, request, and response. '
    'The transparency log records a receiver-issued tool receipt for the inference call; separate domain evidence binds the numeric inference input and output. '
    'The drawing does not claim that every training event is written to this log or that log inclusion alone verifies computation. '
    'No automatic feedback or retraining from inference is implied. Quote, RTMR3, ledger, IPFS, and deployment boundaries are deliberately abstracted. '
    'Admission and retained signing identities underpin attested execution; the overview does not imply a fresh full DCAP check on each inference. '
    'Arrows labelled Result show a returned tool result, not an assertion that patient records enter the language-model context.'
)


def page(prs, letter, title, description):
    s = b.new_content_slide(prs, 9, title, '')
    b.add_domain_navigation(s, 9)
    for shape in s.shapes:
        if getattr(shape, 'has_text_frame', False):
            if shape.text.startswith('Page '):
                shape.text_frame.paragraphs[0].runs[0].text = 'Concept ' + letter
            elif shape.text == '19 August 2026':
                shape.text_frame.paragraphs[0].runs[0].text = '20 September 2026'
    b.add_note(s, description+'\n\n'+COMMON_NOTES)
    return s


def proposal_a(prs):
    s=page(prs,'A','VITA-FL: from training to verifiable use',
           'A: Training loop with a branch. Closest to the requested circle-and-path idea. '
           'Read left to right: repeat DFL rounds, publish a model, invoke its inference using MCP, and retain a tool receipt.')
    end=dfl_loop(s,2.35,3.62,1.30)
    my=end[1]
    route(s,[end,(5.08,my)],PURPLE)
    pill(s,'Published\nmodel',5.97,my,1.72,PURPLE,19,1.0)
    route(s,[(6.88,my),(7.77,my)],PURPLE,2.3)
    pill(s,'Attested\ninference',8.80,my,1.98,BLUE,20,1.05)
    route(s,[(9.83,my),(11.15,my)],BLUE)
    text(s,'Tool receipt',9.87,my-.47,1.23,.33,13.5,BLUE)
    icon_node(s,'log','Transparency\nlog',11.74,my,BLUE,18)
    icon_node(s,'agent','Agent',8.80,1.98,BLUE,19)
    route(s,[(8.55,2.99),(8.55,my-.59)],BLUE)
    text(s,'MCP',7.82,3.25,.64,.3,16,BLUE,True)
    route(s,[(9.07,my-.59),(9.07,2.99)],BLUE,1.6)
    text(s,'Result',9.20,3.25,.78,.3,14,MUTED)
    return s


def proposal_b(prs):
    s=page(prs,'B','One training loop. Many model uses.',
           'B: Two separate loops. Green is repeated training; blue is repeated request/response use. '
           'The model connects the loops in one direction. Receipt logging branches from inference, not back into training.')
    end=dfl_loop(s,2.18,3.35,1.20)
    route(s,[end,(4.34,end[1]),(4.34,4.66),(4.93,4.66)],PURPLE)
    pill(s,'Published\nmodel',5.82,4.66,1.70,PURPLE,18,1.0)
    route(s,[(6.71,4.66),(8.32,4.66)],PURPLE,2.3)
    cx,cy,r=9.32,3.35,1.31
    arc(s,cx,cy,r,-72,16,BLUE,2.4)
    arc(s,cx,cy,r,16,70,BLUE,2.4,False)
    arc(s,cx,cy,r,109,198,BLUE,2.0)
    arc(s,cx,cy,r,198,250,BLUE,2.0,False)
    pill(s,'Agent',cx,cy-r,1.82,BLUE,20,.75)
    pill(s,'Attested\ninference',cx,cy+r,1.94,BLUE,19,1.0)
    text(s,'Use',cx-.58,cy-.21,1.16,.42,24,BLUE,True)
    text(s,'MCP',10.88,3.16,.84,.34,17,BLUE,True)
    text(s,'Result',7.08,3.16,.95,.34,15,MUTED)
    route(s,[(10.33,4.66),(11.33,4.66)],BLUE)
    text(s,'Tool receipt',10.26,4.01,1.33,.3,13.3,BLUE)
    icon_node(s,'log','Transparency\nlog',11.9,4.66,BLUE,17)
    return s


def proposal_c(prs):
    s=page(prs,'C','Train. Publish. Use. Record.',
           'C: Four-scene storyboard. Arrows between the headings denote narrative stages, not deployed network links. '
           'Within Use, the MCP arrow runs from the agent to inference. The final scene records the tool receipt.')
    centers=[2.05,5.10,8.15,11.22]
    headings=['Train','Publish','Use','Record']
    colors=[GREEN,PURPLE,BLUE,BLUE]
    for i,(cx,heading,color) in enumerate(zip(centers,headings,colors)):
        circle(s,cx-.82,2.03,.16,color,color,1)
        text(s,str(i+1),cx-.98,1.87,.32,.32,12,WHITE,True)
        text(s,heading,cx-.48,1.78,1.80,.50,24,color,True,False)
        if i<3:
            route(s,[(cx+1.27,2.04),(centers[i+1]-1.15,2.04)],MUTED,1.5)
    for x in [3.57,6.63,9.70]:
        route(s,[(x,2.67),(x,5.62)],LINE,1,False)
    dfl_loop(s,centers[0],3.85,.75,True)
    text(s,'Training rounds',.78,5.08,2.54,.46,18,GREEN,True)
    circle(s,centers[1],3.85,.64,PURPLE,TINT[PURPLE],1.2)
    model_icon(s,centers[1],3.85,1.3)
    text(s,'Published model',3.80,5.08,2.60,.46,18,PURPLE,True)
    agent_icon(s,centers[2],3.12,1.15)
    text(s,'Agent',8.56,2.96,1.04,.34,17,BLUE,True)
    route(s,[(8.15,3.49),(8.15,4.06)],BLUE,2.0)
    text(s,'MCP',8.34,3.61,.73,.32,15,BLUE,True)
    pill(s,'Inference',8.15,4.43,2.04,BLUE,20,.65)
    text(s,'Attested execution',6.85,5.08,2.60,.46,18,BLUE,True)
    circle(s,centers[3],3.85,.64,BLUE,TINT[BLUE],1.2)
    log_icon(s,centers[3],3.85,1.23)
    text(s,'Transparency log',9.88,5.08,2.68,.46,18,BLUE,True)
    return s


def proposal_d(prs):
    s=page(prs,'D','The model and the request meet at inference',
           'D: Two converging paths. Upper path is model provenance; lower path is agent orchestration. '
           'The two paths join at the inference service, which publishes a receiver-issued tool receipt. '
           'The return path to the agent is omitted to keep the distinction between artifact flow and invocation clear.')
    # Two wide, lightly tinted lanes stop before their shared execution point.
    box(s,.75,1.77,6.26,1.77,GREEN,'F6FAF7',.7)
    box(s,.75,4.10,6.26,1.70,BLUE,'F5F9FC',.7)
    dfl_loop(s,1.97,2.65,.50,True)
    route(s,[(2.77,2.65),(4.43,2.65)],PURPLE,2.1)
    text(s,'Publish',3.03,2.18,1.15,.3,14,PURPLE)
    pill(s,'Published\nmodel',5.39,2.65,1.83,PURPLE,19,.98)
    route(s,[(6.35,2.65),(7.41,2.65),(7.41,3.45),(8.14,3.45)],PURPLE,2.3)
    agent_icon(s,1.97,4.72,1.02)
    text(s,'Agent',1.31,5.13,1.32,.35,19,BLUE,True)
    route(s,[(2.54,4.95),(4.48,4.95)],BLUE,2.1)
    pill(s,'MCP',5.39,4.95,1.70,BLUE,21,.66)
    route(s,[(6.29,4.95),(7.41,4.95),(7.41,4.14),(8.14,4.14)],BLUE,2.3)
    pill(s,'Attested\ninference',9.17,3.79,1.98,BLUE,20,1.26)
    route(s,[(10.21,3.79),(11.38,3.79)],BLUE,2.1)
    text(s,'Tool receipt',10.22,3.23,1.2,.35,13.3,BLUE)
    icon_node(s,'log','Transparency\nlog',11.99,3.79,BLUE,18)
    return s


def main():
    main_deck=ROOT/'VITA-FL_Thesis_Presentation_TU_Berlin.pptx'
    before=hashlib.sha256(main_deck.read_bytes()).digest()
    prs=b.prepare_template()
    prs.core_properties.title='VITA-FL — Lifecycle visualization proposals'
    slides=[proposal_a(prs),proposal_b(prs),proposal_c(prs),proposal_d(prs)]
    prs.save(HERE/'VITA-FL_Lifecycle_Concepts.pptx')
    names=['a-loop-and-branch','b-two-loops','c-four-stages','d-two-paths']
    previews=[]
    for name,slide in zip(names,slides):
        preview=renderer.render_slide(prs,slide)
        preview.save(HERE/(name+'.png'))
        previews.append(preview)
    width,height,gutter,caption=1000,563,30,53
    sheet=Image.new('RGB',(width*2+gutter*3,(height+caption)*2+gutter*3),'#E9EDF0')
    draw=ImageDraw.Draw(sheet)
    font=ImageFont.truetype('/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf',25)
    labels=['A · Kreis mit Abzweig','B · Zwei Kreisläufe','C · Vier Stationen','D · Zwei Wege zur Inferenz']
    for i,(preview,label) in enumerate(zip(previews,labels)):
        x=gutter+(i%2)*(width+gutter)
        y=gutter+(i//2)*(height+caption+gutter)
        draw.text((x+5,y+8),label,font=font,fill='#263440')
        sheet.paste(preview.resize((width,height),Image.Resampling.LANCZOS),(x,y+caption))
    sheet.save(HERE/'overview.png')
    assert hashlib.sha256(main_deck.read_bytes()).digest()==before
    print('Created four editable proposal slides and five PNG previews.')
    print(HERE/'VITA-FL_Lifecycle_Concepts.pptx')


if __name__=='__main__':
    main()
