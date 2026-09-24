#!/usr/bin/env python3
"""Functionally integrated DFL infrastructure, checked against the prototype."""

from __future__ import annotations

import hashlib
import importlib.util
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

HERE = Path(__file__).resolve().parent
SOURCE = HERE.parent
ROOT = HERE.parents[2]
spec = importlib.util.spec_from_file_location('inside_dfl', SOURCE/'inside-circle'/'build_inside.py')
inside = importlib.util.module_from_spec(spec)
spec.loader.exec_module(inside)
c, a, infra = inside.c, inside.a, inside.infra

PROTOTYPE_NOTES = (
    'Verified against vita-fl/dfl/node_server/src/server.ts and ipfs.ts. '
    'Workers obtain the finalized global-model CID tuple from blockchain, then retrieve '
    'and decrypt its bundle from IPFS (ipfs.ts:380-435). They train locally and send '
    'encrypted local updates directly to the current aggregator over authenticated HTTPS '
    '(server.ts:463-533, 2110-2180). Local updates are not exchanged through IPFS. '
    'The aggregator records accepted submission commitments on blockchain (server.ts:672-702), '
    'aggregates updates (server.ts:2313-2330), uploads the encrypted global bundle, signature, '
    'and key bundle to IPFS, then commits the CID tuple on blockchain to finalize the round '
    '(ipfs.ts:315-346; GMStorage.sol:180-248). AggregatorSelection.sol:171-195 selects the '
    'next aggregator; it can select the same participant again. The model is loaded for '
    'the next round. Publication arrows summarize operations performed by the aggregator; '
    'IPFS does not itself submit blockchain transactions. Blockchain carries references '
    'and round state, not model weights. The inference receiver resolves the finalized '
    'reference and retrieves the model bundle independently. Source-to-inference arrows '
    'show information consumed; the receiver initiates retrieval. Placement in a circle '
    'does not define a common TEE or trust boundary. The transparency log is separate '
    'from blockchain and records the receiver-issued tool receipt.'
)


def page(prs, code, detail):
    s = c.page(prs, code, a.TITLE, detail)
    common = c.COMMON_NOTES.replace(
        'Quote, RTMR3, ledger, IPFS, and deployment boundaries are deliberately abstracted.',
        'Quote, RTMR3, and deployment boundaries are deliberately abstracted.')
    c.b.add_note(s, detail+'\n\n'+common+'\n\n'+PROTOTYPE_NOTES)
    return s


def node(s, title, detail, cx, cy, width=2.10, color=c.GREEN, height=.88):
    box = c.box(s, cx-width/2, cy-height/2, width, height, color,
                c.TINT[color] if color == c.PURPLE else c.WHITE, 1.35)
    box.name = 'Functional node: '+title+' / '+detail
    c.text(s, title, cx-width/2+.06, cy-.32, width-.12, .36, 18, color, True)
    c.text(s, detail, cx-width/2+.06, cy+.13, width-.12, .24, 13, color)


def point(cx, cy, r, angle):
    return cx+r*math.cos(math.radians(angle)), cy+r*math.sin(math.radians(angle))


def ring(s, cx, cy, r, arrows):
    c.circle(s, cx, cy, r, c.GREEN, c.WHITE, 2.25)
    for angle in arrows:
        c.arc(s, cx, cy, r, angle+12, angle, c.GREEN, 2.35)


def process_ring(prs):
    s = page(prs, 'F1', 'F1: Workflow ring. Workers train, the aggregator combines their direct '
             'updates, stores the global model in IPFS, and finalizes its reference on blockchain. '
             'The CID label denotes reference registration by the aggregator, not a transfer of '
             'model bytes or an autonomous action by IPFS. The Next round transition summarizes '
             'reference resolution AND loading the model from IPFS. Separate inference sources remain visible.')
    cx, cy, r = 3.23, 3.80, 1.80
    ring(s, cx, cy, r, [-180, 90, 0, -90])
    coordinates = {}
    for title, detail, angle, color in [
        ('Workers', 'Local training', -135, c.GREEN),
        ('Aggregator', 'Combine updates', 135, c.GREEN),
        ('IPFS', 'Store global model', 45, c.PURPLE),
        ('Blockchain', 'Finalize round', -45, c.PURPLE),
    ]:
        x, y = point(cx, cy, r, angle)
        coordinates[title] = (x,y)
        node(s, title, detail, x, y, color=color)
    c.text(s, 'DFL', cx-.68, cy-.27, 1.36, .54, 27, c.GREEN, True)
    c.text(s, 'Next round', 2.39, 1.50, 1.68, .34, 15, c.GREEN)
    c.text(s, 'Publish model', 2.20, 5.80, 2.06, .32, 14, c.GREEN)
    c.text(s, 'Register\nCID', 5.15, 3.40, 1.16, .65, 15, c.PURPLE)
    bx, by = coordinates['Blockchain']
    ix, iy = coordinates['IPFS']
    c.route(s, [(bx+1.11,by),(6.65,by),(6.65,3.83),(8.03,3.83)], c.PURPLE, 2.1)
    c.text(s, 'Reference', 5.67, by-.45, 1.49, .33, 15, c.PURPLE)
    c.route(s, [(ix+1.11,iy),(6.65,iy),(6.65,4.39),(8.03,4.39)], c.PURPLE, 2.1)
    c.text(s, 'Model files', 5.66, iy+.20, 1.70, .33, 15, c.PURPLE)
    infra.use_path(s, 9.07, 4.10, 11.89)
    return s


def coordination_core(prs):
    s = page(prs, 'F2', 'F2: Functional ring. Workers send local updates directly to the aggregator. '
             'The aggregator stores the global model in IPFS; workers load it from IPFS for the next round. '
             'Dashed blockchain links represent commitment/finalization and round/reference reads. '
             'The aggregator also reads current round state; the drawing retains only the principal '
             'directions to stay readable. Reference and model-file retrieval by inference are explicit.')
    cx, cy, r = 3.62, 3.88, 1.73
    ring(s, cx, cy, r, [-180, 60, -55])
    wx, wy = point(cx,cy,r,-125)
    ax, ay = point(cx,cy,r,125)
    px, py = point(cx,cy,r,0)
    node(s, 'Workers', 'Local training', wx, wy, 2.06)
    node(s, 'Aggregator', 'Combine updates', ax, ay, 2.06)
    node(s, 'IPFS', 'Global model', px, py, 2.06, c.PURPLE)
    node(s, 'Blockchain', 'Round + refs', 3.08, cy, 1.94, c.PURPLE)
    c.text(s, 'DFL', 3.86, 2.63, 1.12, .40, 24, c.GREEN, True)
    c.text(s, 'Local\nupdates', .62, 3.51, 1.21, .67, 15, c.GREEN)
    c.text(s, 'Load model', 3.74, 1.68, 1.72, .34, 15, c.GREEN)
    c.text(s, 'Store model', 3.54, 5.82, 1.88, .32, 15, c.GREEN)
    # Actual blockchain interactions, distinct from model/update transport.
    c.b.add_line_segment(s, 3.15,3.37, 2.90,wy+.50, c.PURPLE,1.6,arrow=True,dashed=True)
    c.text(s, 'Round + refs', 3.24, 3.05, 1.54, .28, 13, c.PURPLE)
    c.b.add_line_segment(s, 2.90,ay-.50, 3.15,4.39, c.PURPLE,1.6,arrow=True,dashed=True)
    c.text(s, 'Commit', 3.26, 4.65, 1.08, .28, 14, c.PURPLE)
    c.route(s, [(px+1.09,py),(7.54,py),(7.54,3.84),(8.03,3.84)], c.PURPLE,2.1)
    c.text(s, 'Model files', 6.51,3.36,1.47,.33,15,c.PURPLE)
    c.route(s, [(4.11,4.15),(4.60,4.68),(7.30,4.68),(7.30,4.39),(8.03,4.39)], c.PURPLE,2.1)
    c.text(s, 'Reference', 6.05,4.91,1.42,.32,15,c.PURPLE)
    infra.use_path(s, 9.07,4.10,11.89)
    return s


def publication_arc(prs):
    s = page(prs, 'F3', 'F3: Compact publication arc, closest to the original A silhouette. '
             'The aggregator receives direct worker updates. Its Publish transition stores model '
             'files in IPFS and finalizes their references on blockchain. The Load transition '
             'resolves the finalized reference and retrieves the model from IPFS for the next round. '
             'Both source components form a functional section of the loop. The outgoing '
             'Published model path summarizes their joint use by inference. Blockchain round '
             'coordination and accepted submission commitments are abstracted in this overview.')
    cx,cy,r = 3.56,3.88,1.73
    ring(s,cx,cy,r,[-180,58,-56])
    wx,wy = point(cx,cy,r,-125)
    ax,ay = point(cx,cy,r,125)
    node(s,'Workers','Local training',wx,wy,2.10)
    node(s,'Aggregator','Combine updates',ax,ay,2.10)
    c.text(s,'DFL',2.97,3.65,1.22,.46,26,c.GREEN,True)
    c.text(s,'Local\nupdates',.62,3.51,1.18,.67,15,c.GREEN)
    c.text(s,'Load model',3.67,1.68,1.82,.34,15,c.GREEN)
    c.text(s,'Publish model',3.47,5.82,2.13,.32,15,c.GREEN)
    # This paired segment sits ON the training cycle and has both incoming and
    # outgoing loop paths: publish a model, then load it for the next round.
    px = cx+r
    panel = c.box(s,px-1.10,cy-1.14,2.20,2.28,c.PURPLE,'F8F5FB',1.3)
    panel.name = 'Functional publication segment on DFL cycle'
    c.text(s,'Global model',px-1.02,cy-1.02,2.04,.35,18,c.PURPLE,True)
    for title,detail,y in [('IPFS','Model files',cy-.18),('Blockchain','Model reference',cy+.61)]:
        c.box(s,px-.99,y-.33,1.98,.66,c.PURPLE,c.WHITE,1.0)
        c.text(s,title,px-.91,y-.28,1.82,.32,16,c.PURPLE,True)
        c.text(s,detail,px-.91,y+.09,1.82,.21,12.5,c.PURPLE)
    c.route(s,[(px+1.16,cy),(8.03,cy)],c.PURPLE,2.3)
    c.text(s,'Published\nmodel',6.58,cy-.83,1.37,.65,16,c.PURPLE)
    infra.use_path(s,9.07,cy,11.89,1.98)
    return s


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    protected = [p for p in SOURCE.rglob('*') if p.is_file()
                 and HERE not in p.parents and '__pycache__' not in p.parts]
    protected.append(ROOT/'VITA-FL_Thesis_Presentation_TU_Berlin.pptx')
    before = {p:digest(p) for p in protected}
    prs = c.b.prepare_template()
    prs.core_properties.title = 'VITA-FL — Blockchain and IPFS integrated into DFL'
    slides = [process_ring(prs),coordination_core(prs),publication_arc(prs)]
    destination = HERE/'VITA-FL_Lifecycle_Integrated_DFL.pptx'
    prs.save(destination)
    names = ['f1-process-ring','f2-functional-links','f3-publication-arc']
    previews = []
    for name,slide in zip(names,slides):
        preview = c.renderer.render_slide(prs,slide)
        preview.save(HERE/(name+'.png'))
        previews.append(preview)
    w,h,gap,caption = 1000,563,25,53
    sheet = Image.new('RGB',(w+gap*2,3*(h+caption)+gap*4),'#E9EDF0')
    draw = ImageDraw.Draw(sheet)
    font = ImageFont.truetype('/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf',25)
    labels = ['F1 · Infrastruktur als Arbeitsschritte','F2 · Konkrete Verbindungen im Kreis',
              'F3 · Gemeinsamer Abschnitt für Veröffentlichung']
    for i,(preview,label) in enumerate(zip(previews,labels)):
        y=gap+i*(h+caption+gap)
        draw.text((gap+5,y+8),label,font=font,fill='#263440')
        sheet.paste(preview.resize((w,h),Image.Resampling.LANCZOS),(gap,y+caption))
    sheet.save(HERE/'overview.png')
    reopened=Presentation(destination)
    assert len(reopened.slides)==3
    for slide in reopened.slides:
        contents='\n'.join(sh.text for sh in slide.shapes if sh.has_text_frame)
        assert all(t in contents for t in ['Workers','Aggregator','IPFS','Blockchain','MCP'])
        for sh in slide.shapes:
            assert sh.left>=0 and sh.top>=0
            assert sh.left+sh.width<=reopened.slide_width+1
            assert sh.top+sh.height<=reopened.slide_height+1
    assert all(digest(p)==v for p,v in before.items()), 'A retained original changed'
    print('Created and checked three editable integrated variants; retained proposals/main deck unchanged.')
    print(destination)


if __name__=='__main__':
    main()
