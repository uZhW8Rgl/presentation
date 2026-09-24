#!/usr/bin/env python3
"""Editable lifecycle proposals with Blockchain and IPFS inside the DFL loop."""

from __future__ import annotations

import hashlib
import importlib.util
import math
from pathlib import Path
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

HERE = Path(__file__).resolve().parent
SOURCE = HERE.parent
ROOT = HERE.parents[2]
spec = importlib.util.spec_from_file_location('lifecycle_infrastructure', SOURCE / 'infrastructure' / 'build_variants.py')
infra = importlib.util.module_from_spec(spec)
spec.loader.exec_module(infra)

c = infra.c
a = infra.a

NOTES = (
    'Blockchain and IPFS are shared infrastructure inside the conceptual DFL loop, '
    'not sequential training stages. Blockchain coordinates membership and rounds and '
    'provides authoritative model references. IPFS carries encrypted model bundles and '
    'training artifacts. The active aggregator publishes a finalized model reference and '
    'the corresponding artifacts. Placement inside the loop is functional grouping, not '
    'a common TEE, trust boundary, or restriction to training-only access. '
    'The inference receiver resolves the reference and retrieves the artifacts. '
    'Purple arrows show information consumed by the receiver; the receiver initiates retrieval. '
    'A single Published model path abstracts both operations, rather than a direct file '
    'transfer that bypasses the infrastructure. The two-source variant makes them explicit. '
    'The transparency log is separate from the blockchain.'
)


def page(prs, code, description):
    s = c.page(prs, code, a.TITLE, description)
    common = c.COMMON_NOTES.replace(
        'Quote, RTMR3, ledger, IPFS, and deployment boundaries are deliberately abstracted.',
        'Quote, RTMR3, and deployment boundaries are deliberately abstracted.')
    c.b.add_note(s, description + '\n\n' + common + '\n\n' + NOTES)
    return s


def loop(s, cx, cy, r, layout='right'):
    c.circle(s, cx, cy, r, c.GREEN, c.WHITE, 2.25)
    if layout == 'right':
        nodes = [('Local training', -90, 2.16, .63, c.GREEN),
                 ('Aggregate', 125, 1.77, .63, c.GREEN),
                 ('Global\nmodel', 0, 1.74, .88, c.PURPLE)]
        arrows = [-160, 59, -45]
    else:
        nodes = [('Local training', -90, 2.16, .63, c.GREEN),
                 ('Aggregate', 145, 1.77, .63, c.GREEN),
                 ('Global\nmodel', 35, 1.74, .88, c.PURPLE)]
        arrows = [-151, 92, -16]
    for angle in arrows:
        c.arc(s, cx, cy, r, angle + 13, angle, c.GREEN, 2.3)
    for label, angle, width, height, color in nodes:
        x = cx + r * math.cos(math.radians(angle))
        y = cy + r * math.sin(math.radians(angle))
        c.pill(s, label, x, y, width, color, 17, height,
               c.TINT[color] if color == c.PURPLE else c.WHITE)
        if label.startswith('Global'):
            exit_point = (x + width / 2 + .06, y)
    return exit_point


def store(s, title, detail, cx, cy, w=1.94, h=.74):
    shape = c.box(s, cx-w/2, cy-h/2, w, h, c.PURPLE, c.TINT[c.PURPLE], 1.2)
    shape.name = 'Internal infrastructure: ' + title
    c.text(s, title, cx-w/2+.05, cy-.28, w-.10, .33, 17, c.PURPLE, True)
    c.text(s, detail, cx-w/2+.05, cy+.09, w-.10, .23, 12.5, c.PURPLE)


def blockchain_icon(s, cx, cy):
    positions = [(cx-.31, cy+.07), (cx, cy-.07), (cx+.31, cy+.07)]
    for p, q in zip(positions, positions[1:]):
        c.route(s, [p, q], c.PURPLE, 1.8, False)
    for x, y in positions:
        c.box(s, x-.10, y-.10, .20, .20, c.PURPLE, c.WHITE, 1.2)


def ipfs_icon(s, cx, cy):
    top = (cx, cy-.26)
    left = (cx-.25, cy-.11)
    right = (cx+.25, cy-.11)
    center = (cx, cy+.04)
    bottom = (cx, cy+.31)
    c.route(s, [top, right, (cx+.25, cy+.16), bottom,
                (cx-.25, cy+.16), left, top], c.PURPLE, 1.5, False)
    for points in [[left, center, right], [center, bottom]]:
        c.route(s, points, c.PURPLE, 1.5, False)


def stacked(prs):
    s = page(prs, 'I1', 'I1: Blockchain and IPFS form a stacked core inside the DFL loop. '
             'A separate published-model symbol summarizes their handoff to inference.')
    y = 4.10
    end = loop(s, 3.02, y, 1.73)
    c.text(s, 'DFL', 2.38, 2.76, 1.28, .42, 25, c.GREEN, True)
    store(s, 'Blockchain', 'Rounds & references', 2.87, 3.60)
    store(s, 'IPFS', 'Model files', 2.87, 4.62)
    c.route(s, [end, (6.02, y)], c.PURPLE, 2.2)
    c.icon_node(s, 'model', 'Published\nmodel', 6.56, y, c.PURPLE, 18)
    c.route(s, [(7.10, y), (8.03, y)], c.PURPLE, 2.2)
    infra.use_path(s, 9.07, y, 11.89)
    return s


def side_by_side(prs):
    s = page(prs, 'I2', 'I2: Two infrastructure tiles sit side by side inside the training cycle. '
             'The global model at the circle boundary has a compact, abstracted published-model path to inference.')
    end = loop(s, 3.30, 3.70, 1.80, 'triangle')
    c.text(s, 'DFL', 2.65, 2.47, 1.30, .45, 25, c.GREEN, True)
    for title, x, icon in [('Blockchain', 2.43, blockchain_icon), ('IPFS', 4.17, ipfs_icon)]:
        shape = c.box(s, x-.79, 3.07, 1.58, 1.09, c.PURPLE, c.TINT[c.PURPLE], 1.2)
        shape.name = 'Internal infrastructure: ' + title
        icon(s, x, 3.39)
        c.text(s, title, x-.75, 3.79, 1.50, .30, 16, c.PURPLE, True)
    y = end[1]
    c.route(s, [end, (8.03, y)], c.PURPLE, 2.3)
    c.text(s, 'Published model', 5.85, y-.56, 2.08, .36, 16, c.PURPLE)
    infra.use_path(s, 9.07, y, 11.89, 2.08)
    return s


def two_outputs(prs):
    s = page(prs, 'I3', 'I3: Blockchain and IPFS remain inside the DFL circle. '
             'Two explicit source paths carry the model reference and model files to inference. '
             'The global model stays on the training loop; the paths are retrieval, not extra training phases.')
    loop(s, 3.30, 3.90, 1.85, 'triangle')
    c.text(s, 'DFL', 2.65, 2.46, 1.30, .40, 24, c.GREEN, True)
    store(s, 'Blockchain', 'Rounds & references', 3.20, 3.28, 2.02)
    store(s, 'IPFS', 'Model files', 3.20, 4.15, 2.02)
    c.route(s, [(4.27, 3.28), (7.28, 3.28), (7.28, 3.82), (8.03, 3.82)], c.PURPLE, 2.1)
    c.route(s, [(4.27, 4.15), (7.28, 4.15), (7.28, 4.38), (8.03, 4.38)], c.PURPLE, 2.1)
    c.text(s, 'Model reference', 5.34, 2.84, 1.96, .32, 15, c.PURPLE)
    c.text(s, 'Model files', 5.39, 4.36, 1.78, .32, 15, c.PURPLE)
    infra.use_path(s, 9.07, 4.10, 11.89)
    return s


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    protected = [p for p in SOURCE.rglob('*')
                 if p.is_file() and HERE not in p.parents and '__pycache__' not in p.parts]
    protected.append(ROOT / 'VITA-FL_Thesis_Presentation_TU_Berlin.pptx')
    before = {p: digest(p) for p in protected}
    prs = c.b.prepare_template()
    prs.core_properties.title = 'VITA-FL — Blockchain and IPFS inside the DFL loop'
    slides = [stacked(prs), side_by_side(prs), two_outputs(prs)]
    destination = HERE / 'VITA-FL_Lifecycle_Infrastructure_Inside_DFL.pptx'
    prs.save(destination)
    names = ['i1-stacked-core', 'i2-side-by-side', 'i3-two-outputs']
    previews = []
    for name, slide in zip(names, slides):
        preview = c.renderer.render_slide(prs, slide)
        preview.save(HERE / (name + '.png'))
        previews.append(preview)
    w, h, gutter, caption = 1000, 563, 25, 53
    sheet = Image.new('RGB', (w + gutter*2, 3*(h+caption)+gutter*4), '#E9EDF0')
    draw = ImageDraw.Draw(sheet)
    font = ImageFont.truetype('/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf', 25)
    labels = ['I1 · Infrastruktur übereinander', 'I2 · Infrastruktur nebeneinander',
              'I3 · Zwei sichtbare Wege zur Inferenz']
    for i, (preview, label) in enumerate(zip(previews, labels)):
        x, y = gutter, gutter + i*(h+caption+gutter)
        draw.text((x+5, y+8), label, font=font, fill='#263440')
        sheet.paste(preview.resize((w,h), Image.Resampling.LANCZOS), (x, y+caption))
    sheet.save(HERE / 'overview.png')
    reopened = Presentation(destination)
    assert len(reopened.slides) == 3
    for slide in reopened.slides:
        contents = '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
        assert all(word in contents for word in ['Blockchain', 'IPFS', 'MCP', 'DFL'])
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left+shape.width <= reopened.slide_width+1
            assert shape.top+shape.height <= reopened.slide_height+1
    assert all(digest(p) == value for p, value in before.items()), 'Retained proposal or main deck changed'
    print('Created three editable variants. Slide bounds/content checked; prior proposals and main deck unchanged.')
    print(destination)


if __name__ == '__main__':
    main()
