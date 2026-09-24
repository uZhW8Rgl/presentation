#!/usr/bin/env python3
"""Two separate slide-6 proposals with ledger and IPFS in the purple domain."""

from pathlib import Path
import hashlib
import sys

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'concepts' / 'color-navigation'))
import build_presentation as b
import render_preview as renderer
from build_concepts import check_layout

GREEN, PURPLE, BLUE = '207548', '7446A6', '1764A1'
INK, GRAY = '434343', '89959C'
TINT = {GREEN: 'EDF6F0', PURPLE: 'F3EEF8', BLUE: 'EDF4FA'}


def text(s, value, x, y, w, h, size=12, color=INK, bold=False, center=False):
    return b.add_text(s, value, x, y, w, h, size, color, bold,
                      align=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT,
                      valign=MSO_ANCHOR.MIDDLE, margin=0)


def panel(s, x, y, w, h, color):
    return b.add_box(s, x, y, w, h, fill=TINT[color], line=color, line_width=1.1)


def card(s, name, heading, detail, x, y, w, h, color, compact=False, heading_size=12):
    shape = b.add_box(s, x, y, w, h, fill='FFFFFF', line=color, line_width=1.2)
    shape.name = name
    text(s, heading, x + .10, y + .10, w - .20, .30, heading_size, color, True, True)
    text(s, detail, x + .10, y + (.43 if compact else .51), w - .20,
         .29 if compact else h - .60, 10.5, INK, False, True)
    return shape


def route(s, name, points, color, width=1.8):
    for i, (a, z) in enumerate(zip(points, points[1:])):
        arrow = i == len(points) - 2
        shape = b.add_line_segment(s, *a, *z, color, width, arrow=arrow)
        shape.name = f'{name} / {i + 1}'


def shared_group_frames(s):
    """Overlapping outlines include the shared infrastructure in both groups."""
    for name, x1, y1, x2, y2 in (
        ('DFL and infrastructure', .48, 1.51, 6.99, 5.74),
        ('Infrastructure and agent', 4.14, 1.42, 12.86, 5.88),
    ):
        corners = [(x1, y1), (x2, y1), (x2, y2), (x1, y2), (x1, y1)]
        for index, (a, z) in enumerate(zip(corners, corners[1:]), start=1):
            line = b.add_line_segment(s, *a, *z, '000000', 1.1, dashed=True)
            line.name = f'Shared group frame: {name} / {index}'


def footer(s):
    b.add_box(s, 416/120, 852/120, 711/120, 41/120,
              fill='FFFFFF', line='FFFFFF', radius=False, line_width=0)
    x = 422
    for name, width, color in (
        ('DFL', 92, GREEN),
        ('Blockchain / IPFS', 230, PURPLE),
        ('Agent / MCP / Inference / Log', 352, BLUE),
    ):
        b.add_box(s, x/120, 858/120, width/120, 28/120,
                  fill=color, line=color, radius=False, line_width=.55)
        text(s, name, (x+10)/120, 861/120, (width-20)/120, 22/120,
             8.4, 'FFFFFF', True, True)
        x += width + 12


def page(prs, variant):
    s = b.new_content_slide(prs, 6, 'VITA-FL: two responsibility blocks', '')
    for shape in s.shapes:
        if getattr(shape, 'has_text_frame', False) and shape.text.startswith('Page '):
            shape.text_frame.paragraphs[0].runs[0].text = 'Page 6'
    footer(s)
    b.add_note(s,
        f'Layout proposal {variant} for slide 6. '
        'Green denotes model production; purple denotes blockchain and IPFS; '
        'blue denotes agent orchestration, inference and evidence logging. '
        'The DFL and agent/inference responsibilities remain separate; shared infrastructure '
        'is not a third execution responsibility. These boxes are functional groups, not '
        'deployment or attestation boundaries.\n\n'
        'Ledger -> Receiver TEE supplies the finalized model reference: CIDs, round and '
        'publisher key. IPFS -> Receiver TEE supplies the signed, encrypted model bundle '
        'and associated artifacts. The receiver independently resolves, retrieves and verifies '
        'them. Purple arrowheads depict incoming information; the receiver initiates retrieval. '
        'The ledger does not upload model bytes to IPFS. DFL publication arrows summarize '
        'the active aggregator publishing artifacts and their references.\n\n'
        'Local code references: vita-fl/tee_inference/service/model_source.py:218; '
        'vita-fl/agent/blockchain_source.py:354,544. '
        'The MCP/log chain is the existing slide\'s compact view of tools and evidence, '
        'not a detailed call-sequence diagram.'
    )
    return s


def between_blocks(prs):
    s = page(prs, 'A: infrastructure between the two blocks')
    panel(s, .63, 1.65, 3.40, 3.50, GREEN)
    panel(s, 4.25, 1.65, 2.66, 3.95, PURPLE)
    panel(s, 7.13, 1.65, 5.57, 3.50, BLUE)
    shared_group_frames(s)
    text(s, 'DFL BLOCK', .87, 1.88, 2.90, .30, 13.5, GREEN, True)
    text(s, 'BLOCKCHAIN / IPFS', 4.45, 1.88, 2.26, .30, 11.5, PURPLE, True, True)
    text(s, 'AGENT AND INFERENCE BLOCK', 7.36, 1.88, 5.10, .30, 13.5, BLUE, True)

    card(s, 'Data', 'DATA', 'Signed\ninput', .87, 2.56, 1.24, 1.15, GREEN)
    card(s, 'Worker', 'WORKER TEE', 'Attested\ntraining', 2.37, 2.56, 1.39, 1.15, GREEN, heading_size=10.5)
    card(s, 'Ledger', 'LEDGER', 'Finalized\nmodel reference', 4.48, 2.56, 2.20, 1.15, PURPLE)
    card(s, 'IPFS', 'IPFS', 'Signed, encrypted\nmodel bundle', 4.48, 4.03, 2.20, 1.20, PURPLE)
    card(s, 'Receiver', 'RECEIVER TEE', 'Resolve\nand infer', 7.41, 2.56, 2.11, 1.15, BLUE)
    card(s, 'MCP', 'MCP', 'Bounded\ntools', 9.82, 2.56, 1.14, 1.15, BLUE)
    card(s, 'Log', 'LOG', 'Receipts /\nevidence', 11.25, 2.56, 1.18, 1.15, BLUE)
    route(s, 'Data to worker', [(2.14, 3.135), (2.34, 3.135)], GREEN, 1.4)
    route(s, 'Publish reference', [(3.79, 3.135), (4.44, 3.135)], GREEN)
    route(s, 'Publish artifacts', [(3.065, 3.75), (3.065, 4.63), (4.44, 4.63)], GREEN)
    route(s, 'Ledger to receiver', [(6.71, 3.135), (7.37, 3.135)], PURPLE, 2.2)
    route(s, 'IPFS to receiver', [(6.71, 4.63), (7.06, 4.63), (7.06, 3.46), (7.37, 3.46)], PURPLE, 2.2)
    route(s, 'Receiver to MCP', [(9.55, 3.135), (9.78, 3.135)], BLUE, 1.4)
    route(s, 'MCP to log', [(10.99, 3.135), (11.21, 3.135)], BLUE, 1.4)
    text(s, 'Produces the model', .87, 5.32, 2.90, .27, 11.5, GREEN, True, True)
    text(s, 'The agent orchestrates tools and evidence.', 7.44, 5.31, 4.96, .30, 12, BLUE, True, True)
    return s


def infrastructure_band(prs):
    s = page(prs, 'B: shared infrastructure below the two blocks')
    panel(s, .63, 1.65, 5.08, 2.08, GREEN)
    panel(s, 5.95, 1.65, 6.75, 2.08, BLUE)
    panel(s, .63, 4.68, 12.07, 1.26, PURPLE)
    text(s, 'DFL BLOCK', .87, 1.87, 4.50, .30, 13.5, GREEN, True)
    text(s, 'AGENT AND INFERENCE BLOCK', 6.19, 1.87, 6.15, .30, 13.5, BLUE, True)
    text(s, 'BLOCKCHAIN / IPFS', .91, 4.90, 3.00, .30, 13.5, PURPLE, True)
    text(s, 'Shared infrastructure', .91, 5.35, 3.00, .24, 11, PURPLE)
    card(s, 'Data', 'DATA', 'Signed medical input', .93, 2.35, 1.84, 1.10, GREEN)
    card(s, 'Worker', 'WORKER TEE', 'Attested training', 3.27, 2.35, 2.10, 1.10, GREEN)
    card(s, 'Receiver', 'RECEIVER TEE', 'Resolve and infer', 6.27, 2.35, 2.22, 1.10, BLUE)
    card(s, 'MCP', 'MCP', 'Bounded\ntools', 9.08, 2.35, 1.35, 1.10, BLUE)
    card(s, 'Log', 'LOG', 'Receipts /\nevidence', 11.07, 2.35, 1.33, 1.10, BLUE)
    card(s, 'Ledger', 'LEDGER', 'Finalized model reference', 4.29, 4.88, 2.42, .84, PURPLE, compact=True)
    card(s, 'IPFS', 'IPFS', 'Signed, encrypted model bundle', 8.78, 4.88, 3.57, .84, PURPLE, compact=True)
    route(s, 'Data to worker', [(2.81, 2.93), (3.23, 2.93)], GREEN, 1.5)
    route(s, 'Receiver to MCP', [(8.53, 2.93), (9.04, 2.93)], BLUE, 1.5)
    route(s, 'MCP to log', [(10.47, 2.93), (11.03, 2.93)], BLUE, 1.5)
    route(s, 'Publish model and reference', [(4.32, 3.49), (4.32, 4.08), (3.75, 4.08), (3.75, 4.64)], GREEN)
    text(s, 'Publish model + reference', .94, 4.12, 2.50, .30, 10.5, GREEN, True)
    route(s, 'Ledger to receiver', [(5.50, 4.84), (5.50, 4.00), (6.70, 4.00), (6.70, 3.49)], PURPLE, 2.2)
    route(s, 'IPFS to receiver', [(10.565, 4.84), (10.565, 4.20), (8.03, 4.20), (8.03, 3.49)], PURPLE, 2.2)
    text(s, 'Model reference', 5.67, 4.21, 1.90, .24, 10, PURPLE, True)
    text(s, 'Model bundle', 8.87, 3.86, 1.60, .24, 10, PURPLE, True)
    return s


def main():
    original_paths = [ROOT / 'VITA-FL_Thesis_Presentation_TU_Berlin.pptx', *sorted((ROOT/'preview').glob('*.png'))]
    original_hashes = {p: hashlib.sha256(p.read_bytes()).digest() for p in original_paths}
    prs = b.prepare_template()
    prs.core_properties.title = 'VITA-FL — Slide 6 ledger layout proposals'
    slides = [between_blocks(prs), infrastructure_band(prs)]
    check_layout(prs)
    for s in slides:
        ledger = next(a for a in s.shapes if a.name == 'Ledger')
        assert str(ledger.line.color.rgb) == PURPLE
        for prefix in ('Ledger to receiver', 'IPFS to receiver'):
            parts = [a for a in s.shapes if a.name.startswith(prefix)]
            assert parts and all(str(a.line.color.rgb) == PURPLE for a in parts)
            assert parts[-1]._element.xpath('.//a:tailEnd[@type="triangle"]')
    prs.save(HERE/'VITA-FL_Slide06_Architecture_Drafts.pptx')
    images = []
    for s, filename in zip(slides, ('draft-a.png', 'draft-b.png')):
        preview = renderer.render_slide(prs, s)
        preview.save(HERE/filename)
        images.append(preview)
    images[0].save(HERE/'VITA-FL_Slide06_Architecture_Drafts.pdf', save_all=True, append_images=images[1:], resolution=144)
    assert all(hashlib.sha256(p.read_bytes()).digest() == digest for p, digest in original_hashes.items())
    print('Two editable drafts rendered; ledger and both receiver paths are purple. Main deck and all main previews unchanged.')


if __name__ == '__main__':
    main()
