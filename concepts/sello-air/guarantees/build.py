#!/usr/bin/env python3
"""Export four editable single-slide approaches to Sello and AIR guarantees."""
from pathlib import Path
import hashlib
import html
import importlib.util
import json

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

import slidekit as k

DESIGNS = [
    ('a-two-evidence-layers', 'A · Direkter Vergleich', 'Aufgabe und Garantie unmittelbar nebeneinander.', 'design_a.py'),
    ('b-verification-workflow', 'B · Ablauf', 'Vom autorisierten Aufruf bis zur geprüften Veröffentlichung.', 'design_b.py'),
    ('c-checking-scopes', 'C · Prüfbereiche', 'Tool-Interaktion und numerische Inferenz als getrennte Bereiche.', 'design_c.py'),
    ('d-detected-changes', 'D · Manipulationen', 'Konkrete Szenarien mit Prüfung und erkanntem Verstoß.', 'design_d.py'),
]
STEM = 'VITA-FL_Sello_AIR_Guarantees'


def module_for(filename):
    spec = importlib.util.spec_from_file_location('guarantees_' + Path(filename).stem, k.HERE / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def verify_deck(prs):
    k.c.check_layout(prs)
    for slide in prs.slides:
        ids = [shape.shape_id for shape in slide.shapes]
        assert len(ids) == len(set(ids)), 'duplicate shape IDs'
        notes = slide.notes_slide.notes_text_frame.text
        assert 'SOURCE:' in notes and 'EXACT TASKS AND ENFORCED CHECKS' in notes
        assert 'dcap_collateral_verified=False' in notes
        assert 'HPKE privacy applies only' in notes
        links = [run.hyperlink.address for shape in slide.shapes if shape.has_text_frame
                 for paragraph in shape.text_frame.paragraphs for run in paragraph.runs
                 if run.hyperlink.address]
        assert k.c.SELLO_URL in links and k.c.AIR_URL in links


def gallery():
    cards = '\n'.join(
        f'<article><div class="label">{html.escape(title)}</div><p>{html.escape(detail)}</p>'
        f'<a href="{slug}.png"><img src="{slug}.png" alt="{html.escape(title)}"></a>'
        f'<div class="links"><a href="{slug}.pptx">Einzelfolie · PPTX</a>'
        f'<a href="{slug}.png">Große Vorschau</a></div></article>'
        for slug,title,detail,_ in DESIGNS
    )
    markup = '''<!doctype html><html lang="de"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>VITA-FL · Sello und AIR · Aufgaben und Garantien</title>
<style>
:root{color-scheme:light}*{box-sizing:border-box}body{margin:0;background:#eef1f3;color:#263440;font:17px/1.5 system-ui,sans-serif}
header,main,footer{max-width:1800px;margin:auto;padding:28px}header{padding-bottom:12px}
h1{font-size:clamp(24px,3vw,38px);line-height:1.2;margin:0 0 10px}header p{margin:0 0 16px;max-width:900px}
a{color:#1764a1;text-decoration-thickness:1px;text-underline-offset:3px}.downloads{display:flex;gap:20px;flex-wrap:wrap}
main{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:28px;padding-top:12px}article{background:white;border:1px solid #d9e1e6;padding:18px}
.label{font-size:21px;font-weight:700}article p{font-size:15px;margin:4px 0 16px;color:#5e6d78}img{width:100%;height:auto;border:1px solid #e4e9ed;display:block}
.links{display:flex;gap:20px;margin-top:12px;font-size:15px}footer{font-size:14px;padding-top:0;color:#5e6d78}
@media(max-width:1000px){main{grid-template-columns:1fr}header,main,footer{padding-left:16px;padding-right:16px}}
</style></head><body><header><h1>Sello und AIR: Aufgaben und Garantien</h1>
<p>Vier alternative Einzelfolien im TU-Berlin-Design. Englische Folientexte, editierbare Diagramme und ausführliche Sprechernotizen zum aktuellen Prototyp.</p>
<div class="downloads"><a href="VITA-FL_Sello_AIR_Guarantees.pptx">Alle Varianten · PowerPoint</a>
<a href="VITA-FL_Sello_AIR_Guarantees.pdf">PDF-Vorschau</a><a href="overview.png">Gesamtübersicht</a></div>
</header><main>'''+cards+'''</main><footer>Die Folien beschreiben die bestehende Implementierung. Die geplante Erweiterung mit RA-TLS und Phala-Verifikation ist darin nicht als umgesetzt dargestellt. PNG/PDF sind lokale Layoutvorschauen; die PPTX enthält die editierbaren Originale.</footer></body></html>'''
    (k.HERE/'index.html').write_text(markup,encoding='utf-8')


def main():
    main_deck = k.ROOT / 'VITA-FL_Thesis_Presentation_TU_Berlin.pptx'
    before = hashlib.sha256(main_deck.read_bytes()).hexdigest()
    prs=k.b.prepare_template()
    prs.core_properties.title='VITA-FL — Sello and AIR: tasks and guarantees'
    prs.core_properties.subject='Four alternative combined single slides: comparison, workflow, scope map and manipulation matrix'
    prs.core_properties.author='Ramon Mehrpoya'
    for slug,title,detail,filename in DESIGNS:
        module = module_for(filename)
        module.build(prs)
        single=k.b.prepare_template()
        single.core_properties.title=title+' — VITA-FL Sello / AIR'
        single.core_properties.author='Ramon Mehrpoya'
        module.build(single)
        verify_deck(single)
        single.save(k.HERE/(slug+'.pptx'))
    verify_deck(prs)
    output=k.HERE/(STEM+'.pptx')
    prs.save(output)
    reopened=Presentation(output)
    assert len(reopened.slides)==4
    verify_deck(reopened)
    images=[]
    for slide,(slug,*_) in zip(reopened.slides,DESIGNS):
        image=k.renderer.render_slide(reopened,slide)
        image.save(k.HERE/(slug+'.png'))
        images.append(image)
    images[0].save(k.HERE/(STEM+'.pdf'),save_all=True,append_images=images[1:],resolution=120)
    overview=Image.new('RGB',(1660,1040),'#E8ECEF')
    draw=ImageDraw.Draw(overview)
    label_font=ImageFont.truetype(k.renderer.FONT_BOLD,22)
    for i,(image,(_,title,_,_)) in enumerate(zip(images,DESIGNS)):
        x,y=20+(i%2)*820,18+(i//2)*510
        draw.text((x+4,y),title,font=label_font,fill='#263440')
        overview.paste(image.resize((800,450),Image.Resampling.LANCZOS),(x,y+39))
    overview.save(k.HERE/'overview.png')
    gallery()
    after=hashlib.sha256(main_deck.read_bytes()).hexdigest()
    assert before==after, 'Main presentation unexpectedly changed'
    validation={
        'slides':len(reopened.slides), 'individual_pptx_files':4,
        'editable_native_shapes':True,'layout_check':'passed',
        'source_links_and_speaker_notes':'passed',
        'main_deck_unchanged':True,'main_deck_sha256':before,
        'preview_renderer':'local Pillow layout renderer; PDF is a preview',
        'visual_review':'pending',
        'variants':[{'file':slug+'.pptx','shapes':len(slide.shapes)}
                    for slide,(slug,*_) in zip(reopened.slides,DESIGNS)]
    }
    (k.HERE/'validation.json').write_text(json.dumps(validation,indent=2,ensure_ascii=False)+'\n')
    print('Built 4 editable variants, individual PPTX/PNG files, combined PPTX/PDF, overview and gallery. Layout and source checks passed.')


if __name__=='__main__':
    main()
