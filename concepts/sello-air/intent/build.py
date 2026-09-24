#!/usr/bin/env python3
"""Build revised purpose-led Sello/AIR proposals using editable shapes."""
from pathlib import Path
import hashlib
import json
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
import common as c
import designs

DESIGNS = [
    ('a-sello-intention', 'A · Sello · Absicht und Übernahme', designs.a_sello),
    ('a-air-intention', 'A · AIR · Absicht und Übernahme', designs.a_air),
    ('b-sello-scenario', 'B · Sello · Am Agentenbeispiel', designs.b_sello),
    ('b-air-scenario', 'B · AIR · Am Agentenbeispiel', designs.b_air),
]


def main():
    main_deck = c.ROOT / 'VITA-FL_Thesis_Presentation_TU_Berlin.pptx'
    before = hashlib.sha256(main_deck.read_bytes()).hexdigest()
    prs = c.b.prepare_template()
    prs.core_properties.title = 'VITA-FL — Sello and AIR: intention and implementation'
    for _,_,build in DESIGNS:
        build(prs)
    c.check_layout(prs)
    path = HERE / 'VITA-FL_Sello_AIR_Purpose.pptx'
    prs.save(path)
    reopened = Presentation(path)
    previews = []
    for slide,(slug,_,_) in zip(reopened.slides, DESIGNS):
        preview = c.renderer.render_slide(reopened,slide)
        preview.save(HERE / (slug+'.png'))
        previews.append(preview)
        assert 'SOURCE:' in slide.notes_slide.notes_text_frame.text
    previews[0].save(HERE/'VITA-FL_Sello_AIR_Purpose.pdf', save_all=True,
                     append_images=previews[1:], resolution=120)
    overview = Image.new('RGB',(1660,1040),'#E8ECEF')
    draw = ImageDraw.Draw(overview)
    font = ImageFont.truetype(c.renderer.FONT_BOLD,21)
    for i,(preview,(_,label,_)) in enumerate(zip(previews,DESIGNS)):
        x,y = 20+(i%2)*820, 16+(i//2)*515
        draw.text((x+4,y),label,font=font,fill='#263440')
        overview.paste(preview.resize((800,450),Image.Resampling.LANCZOS),(x,y+37))
    overview.save(HERE/'overview.png')
    assert before == hashlib.sha256(main_deck.read_bytes()).hexdigest()
    (HERE/'validation.json').write_text(json.dumps({
        'slides':4,'layout':'passed','source_notes':'present',
        'main_deck_unchanged':True,'main_deck_sha256':before},indent=2)+'\n')
    print('Built four purpose-led slides, previews and PDF; layout checked.')


if __name__ == '__main__':
    main()
