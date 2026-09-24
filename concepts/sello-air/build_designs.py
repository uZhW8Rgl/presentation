#!/usr/bin/env python3
"""Build three pairs of editable proposals and visual previews."""
import hashlib
import html
import json

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

import common as c
import design_a
import design_b
import design_c

DESIGNS = [
    ('a-sello-flow', 'A · Sello · Ablauf', design_a.build_sello),
    ('a-air-flow', 'A · AIR · Ablauf', design_a.build_air),
    ('b-sello-comparison', 'B · Sello · Paper und Umsetzung', design_b.build_sello),
    ('b-air-comparison', 'B · AIR · Draft und Umsetzung', design_b.build_air),
    ('c-sello-receipt', 'C · Sello · Aufbau des Belegs', design_c.build_sello),
    ('c-air-bundle', 'C · AIR · Aufbau des Nachweises', design_c.build_air),
]


def main():
    main_deck = c.ROOT / 'VITA-FL_Thesis_Presentation_TU_Berlin.pptx'
    before = hashlib.sha256(main_deck.read_bytes()).hexdigest()
    prs = c.b.prepare_template()
    prs.core_properties.title = 'VITA-FL — Sello and AIR slide alternatives'
    for _, _, build in DESIGNS:
        build(prs)
    c.check_layout(prs)
    destination = c.HERE / 'VITA-FL_Sello_AIR_Alternatives.pptx'
    prs.save(destination)
    reopened = Presentation(destination)
    assert len(reopened.slides) == 6
    images = []
    for slide, (slug, title, _) in zip(reopened.slides, DESIGNS):
        ids = [shape.shape_id for shape in slide.shapes]
        assert len(ids) == len(set(ids))
        assert 'SOURCE:' in slide.notes_slide.notes_text_frame.text
        assert 'DEVIATIONS / LIMITS:' in slide.notes_slide.notes_text_frame.text
        assert any(run.hyperlink.address for shape in slide.shapes if shape.has_text_frame
                   for paragraph in shape.text_frame.paragraphs for run in paragraph.runs)
        preview = c.renderer.render_slide(reopened, slide)
        preview.save(c.HERE / (slug + '.png'))
        images.append(preview)
    images[0].save(c.HERE / 'VITA-FL_Sello_AIR_Alternatives.pdf',
                   save_all=True, append_images=images[1:], resolution=120)
    overview = Image.new('RGB', (1660, 1540), '#E8ECEF')
    draw = ImageDraw.Draw(overview)
    font = ImageFont.truetype(c.renderer.FONT_BOLD, 21)
    for index, (preview, (_, title, _)) in enumerate(zip(images, DESIGNS)):
        x, y = 20 + (index % 2) * 820, 18 + (index // 2) * 510
        draw.text((x+6,y+3), title, font=font, fill='#263440')
        overview.paste(preview.resize((800,450), Image.Resampling.LANCZOS), (x,y+39))
    overview.save(c.HERE / 'overview.png')
    for number, letter in enumerate('abc'):
        pair = Image.new('RGB', (1660, 515), '#E8ECEF')
        pair.paste(overview.crop((0, number*510, 1660, number*510+510)), (0,0))
        pair.save(c.HERE / f'{letter}-pair.png')
    cards = '\n'.join(
        f'<article><h2>{html.escape(title)}</h2><a href="{slug}.png">'
        f'<img src="{slug}.png" alt="{html.escape(title)}"></a></article>'
        for slug,title,_ in DESIGNS
    )
    (c.HERE / 'index.html').write_text('''<!doctype html><html lang="de"><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>Sello / AIR – Entwürfe</title>
<style>body{font:17px system-ui;color:#263440;background:#e8ecef;margin:24px}h1{margin-bottom:8px}
main{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:22px}article{background:white;padding:14px}
img{width:100%;height:auto}h2{font-size:19px}a{color:#1764a1}@media(max-width:850px){main{grid-template-columns:1fr}}</style>
<h1>Sello und AIR – drei Varianten pro Folie</h1><p>A: Ablauf · B: Quelle und Umsetzung · C: Aufbau des Nachweises</p>
<p><a href="VITA-FL_Sello_AIR_Alternatives.pptx">Editierbare PowerPoint</a> ·
<a href="VITA-FL_Sello_AIR_Alternatives.pdf">PDF</a> · <a href="README.md">Quellen und Einordnung</a></p><main>'''
        + cards + '</main>', encoding='utf-8')
    assert before == hashlib.sha256(main_deck.read_bytes()).hexdigest()
    (c.HERE / 'validation.json').write_text(json.dumps({
        'slides': 6, 'editable_shapes': True, 'layout_check': 'passed',
        'source_links_and_notes': 'passed', 'main_deck_sha256': before,
        'main_deck_unchanged': True,
    }, indent=2) + '\n')
    print('Built 6 editable proposals, individual PNGs, pair previews, overview, PDF and HTML. Main deck unchanged.')


if __name__ == '__main__':
    main()
