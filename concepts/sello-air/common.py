"""Shared editable shapes and source notes for Sello/AIR slide proposals."""
from pathlib import Path
import sys

from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'concepts' / 'color-navigation'))
import build_presentation as b
import render_preview as renderer
from build_concepts import check_layout
from technology_icons import add_icon

INK, MUTED, WHITE = '263440', '67747C', 'FFFFFF'
BLUE, GREEN, PURPLE, AMBER = '1764A1', '207548', '7446A6', '986021'
TINT = {INK:'F3F5F6', BLUE:'EDF4FA', GREEN:'EDF6F0', PURPLE:'F3EEF8', AMBER:'FCF4E9'}
SELLO_URL = 'https://arxiv.org/abs/2606.04193'
AIR_URL = 'https://www.ietf.org/archive/id/draft-tsyrulnikov-rats-attested-inference-receipt-02.html'

SELLO_NOTES = '''SOURCE: Juan Figuera, Notarized Agents: Receiver-Attested Confidential Receipts for AI Agent Actions (2026), arXiv:2606.04193v1. https://arxiv.org/abs/2606.04193
The paper moves receipt production to the receiving service, encrypts receipt contents for an independent owner, and combines witnessed logging with owner discovery by token reference.

CURRENT VITA-FL IMPLEMENTATION (23 September 2026): The three fixed operations are model-bundle fetch, ChestMNIST image selection, and inference. The receiver hashes the exact canonical tool input and raw response, HPKE-encrypts the receipt body, and signs its COSE_Sign1 envelope with a dedicated Ed25519 key. The receiver publishes to SCITT before releasing a successful response. The agent verifies the signature, decrypts the receipt, checks byte commitments and verifies the SCITT publication bundle. A log failure yields HTTP 503; already performed computation is not rolled back. The Sello inference input is a job_id object, not the image pixels. Pixel-level evidence is in the separately published AIR bundle.

VITA-FL ADDITIONS: dstack-derived receiver key, bound at DCAP admission to Worker 0 and its endpoint in DeviceRegistry; mandatory mTLS; short-lived JWS with audience, action scopes, certificate thumbprint and matching subject. This is transport identity, not attestation of agent code. A random jti is not a one-use store and the token is not bound to the request Digest.

DEVIATIONS / LIMITS: SCITT/CCF runs as one virtual node with no witness cosigning or independent log-key pin. Retrieval follows a returned publication URL; there is no owner discovery service indexed by token reference. Owner HPKE private key and token-issuer signing seed are configured in the agent runtime: the independent-owner threat boundary is therefore not preserved against compromise of that runtime. Domain evidence remains a separate layer. These slides describe the local implementation, not a newly deployed or evaluated run.

CODE: vita-fl/agent_receipts/sello_v1.py; agent_receipts/environment.py; agent_receipts/receiver_log.py; agent_receipts/scitt.py; agent_receipts/dstack_key.py; agent/sello_client.py; tee_inference/service/app.py; transport_security/peer.py; phala/dstack-compose.contracts.phala.tftpl.
'''

AIR_NOTES = '''SOURCE: B. Tsyrulnikov, Attested Inference Receipt (AIR): A COSE/CWT Profile for Confidential AI Inference, Internet-Draft 02, 5 July 2026. https://www.ietf.org/archive/id/draft-tsyrulnikov-rats-attested-inference-receipt-02.html
This is a work-in-progress draft, not a peer-reviewed paper or RFC. It defines signed single-inference evidence and distinguishes receipt validation from platform attestation appraisal.

CURRENT VITA-FL IMPLEMENTATION (23 September 2026): Native ChestMNIST inference uses deterministic CBOR request, response and model-manifest bytes. A dedicated dstack-derived Ed25519 key signs the AIR-style COSE/CWT receipt. The model Digest is the manifest Digest; that manifest binds the separately available DFL model artifact. The receipt contains commitments, not the original model weights or pixels. The returned domain bundle additionally contains exact numeric request/response bytes, manifest, receipt/key, quote, REPORTDATA, ordered event log and measured app_compose. The agent verifies the bundle and separately publishes it to SCITT. This published bundle contains the input pixels in plaintext; receipt hashing does not make this bundle confidential.

KEY / WORKLOAD BINDING: The 64-byte TDX REPORTDATA contains a domain-separated hash of the AIR public key and model-manifest Digest, followed by the request Digest. A fresh quote is requested per inference; its Digest is included in the AIR receipt. The verifier checks the AIR signature, request/response/manifest/nonce bindings, REPORTDATA, ordered SHA-384 RTMR3 replay, measured Compose, expected worker-image Digest and deployment policy. This is domain evidence for the numeric DFL inference, not evidence of the agent's LLM prompt or Ollama reasoning.

DEVIATIONS / LIMITS: The per-call verifier explicitly returns dcap_collateral_verified=False. It does not authenticate the individual quote's Intel signature/certificate chain, TCB or revocation collateral; it relies on separately admitted Worker 0 and its registry-bound Sello receiver identity. It also does not reconcile all signed measurement-map values with the quote registers. The REPORTDATA layout is VITA-FL-specific; it differs from revision 02's illustrated TDX key-binding construction. The nonce is generated by the receiver, not the verifier, and there is no seen-cti store; the 300-second issue-time policy (30-second skew) gives recentness, not full replay protection. Describe the profile as AIR-style. Signature and attestation do not establish model accuracy or prove the forward-pass arithmetic independently of trusted execution.

CODE: vita-fl/tee_inference/air/v1.py; tee_inference/service/attestation.py; tee_inference/service/jobs.py; tee_inference/protocol/v1.py; agent/tee_inference_client.py; agent/scitt_client.py.
'''


def text(s, value, x, y, w, h=.36, size=17, color=INK, bold=False, center=False):
    shape = b.add_text(s, value, x, y, w, h, size, color, bold,
                      align=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT,
                      valign=MSO_ANCHOR.MIDDLE, margin=0)
    shape.name = 'Sello AIR: ' + value.replace('\n', ' / ')
    return shape


def panel(s, x, y, w, h, color=INK, fill=None, width=1.1):
    return b.add_box(s, x, y, w, h, fill=fill or TINT[color], line=color, line_width=width)


def route(s, points, color=INK, width=1.8, arrow=True):
    for i, (start, end) in enumerate(zip(points, points[1:])):
        b.add_line_segment(s, *start, *end, color, width, arrow=arrow and i == len(points)-2)


def icon(s, name, x, y, size=.4, color=BLUE):
    add_icon(s, name, x, y, size, color, TINT[color], prefix='Sello AIR icon: ')


def label(s, value, x, y, w, color=BLUE):
    return text(s, value, x, y, w, .25, 12, color, True)


def card(s, title, detail, x, y, w, h=1.2, color=BLUE, size=17):
    panel(s, x, y, w, h, color)
    text(s, title, x+.16, y+.13, w-.32, .37, size+1, color, True, True)
    text(s, detail, x+.16, y+.56, w-.32, h-.64, size, INK, False, True)


def page(prs, code, title, kind):
    s = b.new_content_slide(prs, len(prs.slides)+1, title, '')
    b.add_domain_navigation(s, 28)
    for shape in s.shapes:
        if shape.has_text_frame and shape.text.startswith('Page '):
            shape.text_frame.paragraphs[0].runs[0].text = code
            shape.width = b.Inches(1.10)
    notes = SELLO_NOTES if kind == 'sello' else AIR_NOTES
    b.add_note(s, 'Visual proposal ' + code + '\n\n' + notes)
    return s


def source(s, kind):
    value, url = (
        ('Figuera · Notarized Agents (2026) · arXiv:2606.04193', SELLO_URL)
        if kind == 'sello' else
        ('Tsyrulnikov · AIR Internet-Draft 02 (2026) · work in progress', AIR_URL)
    )
    shape = text(s, value, .66, 6.04, 12.0, .20, 10, MUTED)
    shape.text_frame.paragraphs[0].runs[0].hyperlink.address = url


def caveat(s, title, body, y=5.06):
    panel(s, .66, y, 12.0, .78, AMBER)
    text(s, title, .85, y+.07, 2.10, .64, 16, AMBER, True)
    text(s, body, 3.08, y+.08, 9.33, .61, 15, INK)
