"""Editable slide helpers and verified source map for evaluation alternatives."""
from pathlib import Path
import sys

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
REPO = ROOT.parent / "vita-fl"
sys.path.insert(0, str(ROOT))
import build_presentation as b
import render_preview as renderer

INK, MUTED, WHITE = "263440", "65737D", "FFFFFF"
GREEN, BLUE, PURPLE, RED, AMBER = "207548", "1764A1", "7446A6", "C40D1E", "9A6100"
LINE = "DCE3E8"
TINT = {GREEN: "EDF6F0", BLUE: "EDF4FA", PURPLE: "F3EEF8", RED: "FBECEF", AMBER: "FFF5E5", INK: "F2F5F7"}

SOURCES = [
    ("dfl/neural_network/tests/test_dicom_provenance.py", [
        "test_checked_in_shard_verifies_against_onchain_style_snapshot",
        "test_modified_label_is_rejected", "test_modified_image_is_rejected"]),
    ("smart_contracts/test/DeviceRegistryAttestationBinding.t.sol", [
        "testRegistersOnlyWithBoundAppComposeEvidence",
        "testRejectsWrongImageDigestDerivedFromAppCompose"]),
    ("smart_contracts/test/GMStorageContribution.t.sol", [
        "testIdenticalRetryIsIdempotent", "testRetryWithChangedPackageHashReverts",
        "testFirstSubmissionRejectsWrongActionKeySignature"]),
    ("smart_contracts/test/AggregatorSelectionAccess.t.sol", [
        "testAuthorizedNonAggregatorCanRecoverSelectionGapAfterCompletedRound",
        "testTimeoutReportCannotAbortAnActiveSubmissionWindow",
        "testTimeoutQuorumDoesNotShrinkAfterRegistryContraction"]),
    ("smart_contracts/test/GMStorageAbort.t.sol", [
        "testAbortBeforePublicationLeavesActiveArtifactsUnchanged",
        "testAbortAfterRejectedPublicationKeepsLastCompletedArtifacts"]),
    ("dfl/node_server/test/gm_crypto.test.js", [
        "gm crypto rejects decryption for participants without a wrapped key"]),
    ("agent/tests/test_tee_inference_client.py", [
        "test_valid_bundle_verifies", "test_request_substitution_is_rejected",
        "test_modified_event_log_is_rejected"]),
    ("tee_inference/tests/test_authorization.py", [
        "test_wrong_token_binding_scope_and_expiry_rejected_before_model_load",
        "test_another_authenticated_agent_cannot_read_or_run_job",
        "test_publication_failure_cannot_return_success_without_receipt"]),
    ("agent/tests/test_sello_v1.py", [
        "test_receiver_receipt_round_trip", "test_tampered_output_is_rejected",
        "test_wrong_receiver_is_rejected"]),
    ("agent/tests/test_receiver_transport.py", [
        "test_tool_response_without_receipt_fails_closed"]),
]

NOTES = """PURPOSE
One proposed evaluation slide, to be selected before insertion in the main deck.
All visible examples are implemented tests. This is an inventory and explanation
of the test design, not a new execution report. No test counts or new pass rates
are asserted. The existing September 2026 cloud run is separate empirical evidence.

CLASSIFICATION
Component tests exercise real production functions with selected dependencies
controlled: training provenance, native AES/RSA model exchange, AIR/Sello checks.
Solidity/Foundry examples are local contract component tests; they execute actual
contract logic, with registry/selection/TDX verifier stubs depending on the suite.
Local integration tests exercise the FastAPI receiver through an ASGI test client,
including authorization and receipt-publication error handling. Model loading,
transport context and publication dependencies are controlled. They do not run
the complete deployed service network.
The one Phala execution is the complete deployed workflow: six worker CVMs,
24 successful training rounds and one final inference using the round-25 model.
It is an existing run record, not an automated fault-injection test suite.

EXACT ASSERTIONS AND BOUNDARIES
Provenance: changing a signed label or image fails signature verification. Real
cryptography uses synthetic medical signer fixtures and an on-chain-style local
snapshot; clinical provenance and label correctness are outside this assertion.
Admission: matching image/Compose evidence is accepted; wrong image is rejected.
MockTdxV4Attestation supplies the verifier result; real DCAP is not evaluated here.
Contributions: changed package hash, wrong action key and nonce are rejected;
identical retry is idempotent. This does not evaluate protocol-valid poisoning.
Recovery: timeout reports cannot abort an active submission window; the quorum
does not shrink after registry contraction. Abort preserves the last finalized
artifacts. Selection-gap recovery after completed publication is a distinct path.
There is no positive full timeout-to-quorum-to-abort-to-replacement experiment
established by these selected tests. The cloud run did not trigger recovery.
Inference: a valid AIR bundle is accepted; a substituted request or modified
event log is rejected. The legacy fixture has a synthetic unsigned TDX quote;
its AIR signatures/hashes are real but do not establish hardware authentication.
Tool receipt: output mutation and the wrong receiver are rejected by real local
Sello signature/encryption checks. A missing receipt propagates an error through
the client wrapper; these are not measurements of an actual LLM's truthfulness.
Service integration: unauthorized calls fail before work; foreign subjects may
not access a job. A failed log publication gives HTTP 503 and no receipt. The
operation may already have executed and is not rolled back by this failure.
The SCITT/log failure is injected through a mock, not a live log outage.

VERSION SCOPE
The main presentation describes the evaluated legacy mTLS/AIR-v1 profile. Newer
RA-TLS-specific suites exist but are deliberately not mixed into that run's
claims. This slide uses examples from the legacy-compatible suites. A fresh
recorded execution should pin the tested commit and profile before reporting
additional numerical results. Framework labels describe scope, not pass status.

PRESENTATION
English matches the main deck. Suggested insertion: immediately after current
Page 24 (cloud-run overview), before the learning curves. Source names are
abbreviated on the slide; exact paths and functions are preserved below.
"""


def notes():
    blocks = [NOTES, "SOURCE MAP (paths relative to vita-fl):"]
    for path, names in SOURCES:
        blocks.append(path + "\n  " + "\n  ".join(names))
    blocks.append("CLOUD RUN: data/evaluation/authoritative-phala-6w-24r-20260901/README.md")
    return "\n\n".join(blocks)


def page(prs, letter, title, subtitle="EXISTING TEST IMPLEMENTATIONS"):
    slide = b.new_content_slide(prs, len(prs.slides) + 1, title, subtitle)
    b.add_domain_navigation(slide, 24, active_domains={"DFL", "Agent"})
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text.startswith("Page "):
            shape.text_frame.paragraphs[0].runs[0].text = "Variant " + letter
            shape.width = b.Inches(1.1)
    b.add_note(slide, "VARIANT " + letter + "\n\n" + notes())
    return slide


def text(slide, value, x, y, w, h=.32, size=15, color=INK, bold=False, center=False):
    shape = b.add_text(slide, value, x, y, w, h, size, color, bold,
                       align=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT,
                       valign=MSO_ANCHOR.MIDDLE, margin=0)
    shape.name = "Test evaluation: " + value.replace("\n", " / ")
    return shape


def panel(slide, x, y, w, h, color=INK, fill=None, width=1, radius=True):
    return b.add_box(slide, x, y, w, h, fill=fill or TINT[color], line=color,
                     radius=radius, line_width=width)


def tag(slide, value, x, y, w, color=INK, size=10.5, h=.29):
    panel(slide, x, y, w, h, color, color, 0, False)
    text(slide, value, x+.06, y+.015, w-.12, h-.03, size, WHITE, True, True)


def route(slide, points, color=INK, width=1.6, arrow=True):
    for index, (start, end) in enumerate(zip(points, points[1:])):
        b.add_line_segment(slide, *start, *end, color, width,
                           arrow=arrow and index == len(points)-2)


def line(slide, x, y, w, color=LINE):
    b.add_divider(slide, x, y, w, color, .012)


def scope(slide, value="Scope: controlled test cases; mocks delimit hardware and cloud-availability claims."):
    text(slide, value, .68, 5.97, 11.98, .26, 10.5, MUTED)


def verify_sources():
    for path, names in SOURCES:
        source = (REPO / path).read_text()
        for name in names:
            if name not in source:
                raise AssertionError("Missing test: " + path + " :: " + name)
