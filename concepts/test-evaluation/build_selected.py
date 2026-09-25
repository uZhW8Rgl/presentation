#!/usr/bin/env python3
"""Build selected A as a native slide, using the audited static inventory."""
import json
import sys

from pptx import Presentation

import common as c

sys.path.insert(0, str(c.ROOT / "concepts" / "color-navigation"))
from build_concepts import check_layout
from build_designs import check_bounds


def build(prs):
    inventory = json.loads((c.HERE / "test-inventory.json").read_text())
    # The inventory is independently generated from source; changes require review.
    assert inventory["counts"]["definitions"] == 808
    assert inventory["counts"]["by_category"] == {"unit_component": 731, "local_integration": 77, "system_e2e": 0}
    assert len(inventory["tests"]) == 808
    s = c.b.new_content_slide(
        prs, 25, "Implemented tests by test level",
        "EVALUATION · TEST COUNT, BOUNDARY AND EXECUTION")
    c.b.add_domain_navigation(s, 25, active_domains={"DFL", "Agent"})
    for shape in s.shapes:
        if shape.has_text_frame and shape.text.startswith("Page "):
            shape.text_frame.paragraphs[0].runs[0].text = "Page 25"

    c.text(s, "808 implemented test definitions", .68, 1.51, 7.20, .40, 21, c.INK, True)
    c.text(s, "Audited CI scope · current code, all profiles", 7.93, 1.54, 4.72, .34, 12, c.MUTED)

    columns = [("TEST TYPE", .83, 2.27), ("COUNT", 3.15, 1.22),
               ("BOUNDARY / CHECKED SCOPE", 4.58, 4.15),
               ("CONFIGURED EXECUTION", 8.98, 3.51)]
    for value, x, w in columns:
        c.text(s, value, x, 2.05, w, .27, 10.5, c.MUTED, True)
    c.line(s, .68, 2.40, 11.98)

    rows = [
        ("Unit / component", "731", "definitions", c.GREEN,
         "Functions and contract logic", "Validation · signatures · state guards",
         "CI · GitHub Actions", "Python / Node.js / Foundry"),
        ("Integration", "77", "definitions", c.BLUE,
         "Real local interfaces", "HTTP/TLS · contracts · external tools",
         "CI · GitHub Actions", "Local services, binaries and EVM"),
        ("System / end-to-end", "0", "automated", c.PURPLE,
         "Complete deployed workflow", "Training → publication → inference",
         "1 separate Phala run", "6 workers · 24 rounds + final inference"),
        ("Production checks*", "0", "identified", c.MUTED,
         "Tests against the live deployment", "Operational behaviour under real load",
         "No suite identified", "Within the audited CI scope"),
    ]
    for index, (kind, count, count_label, color, boundary, checks, location, detail) in enumerate(rows):
        y = 2.47 + index * .64
        c.b.add_box(s, .68, y, .045, .555, fill=color, line=color, radius=False, line_width=0)
        c.text(s, kind, .83, y+.08, 2.27, .42, 14, color, True)
        c.text(s, count, 3.15, y, 1.18, .40, 23, color, True)
        c.text(s, count_label, 3.15, y+.395, 1.24, .21, 10.5, c.MUTED)
        c.text(s, boundary, 4.58, y+.025, 4.17, .27, 14, c.INK, True)
        c.text(s, checks, 4.58, y+.33, 4.17, .24, 11.5, c.MUTED)
        c.text(s, location, 8.98, y+.025, 3.52, .27, 14, color, True)
        c.text(s, detail, 8.98, y+.33, 3.52, .24, 10.5, c.MUTED)
        c.line(s, .83, y+.61, 11.80)

    c.panel(s, .68, 5.17, 11.98, .53, c.INK, c.TINT[c.INK], 0, False)
    c.text(s, "CI: push / pull request / manual · Ubuntu runners", .84, 5.205, 6.40, .43, 13, c.INK, True)
    c.text(s, "+ 2 service smoke checks: Anvil + IPFS", 7.42, 5.205, 5.04, .43, 12.5, c.INK)
    c.text(s, "Static definitions, including conditional skips; parameter variants counted once. No pass count.",
           .68, 5.81, 11.98, .23, 10.5, c.MUTED)
    c.text(s, "* Production denotes an execution context. Current-code inventory and recorded cloud run are separate evidence.",
           .68, 6.075, 11.98, .23, 9.5, c.MUTED)

    report = (c.HERE / "test-inventory.md").read_text()
    c.b.add_note(s, "SELECTED VARIANT A — COUNTS STRICTLY BY TEST TYPE\n\n"
        "731 unit/component + 77 local integration = 808 unique implemented definitions. "
        "This is a static inventory, not an execution result, pass count or coverage percentage. "
        "The assignment is based on the asserted boundary, not suite names. Component combines "
        "isolated function/module/contract tests; integration crosses a real interface.\n\n"
        "CURRENT CODE VERSUS HISTORICAL EVIDENCE\n"
        "Includes the current RA-TLS, legacy mTLS, ZK, infrastructure and deployment-helper suites. "
        "These counts do not describe the exact code version of the September 1, 2026 Phala run. "
        "That separate run used six worker CVMs, 24 successful federated rounds and final inference "
        "with the round-25 model. It did not exercise aggregator failover. Its exported evidence "
        "does not preserve original evidence.cbor/transparent-statement byte files.\n\n"
        "SYSTEM AND PRODUCTION\n"
        "Zero system tests means no implemented automated full training-to-inference test suite "
        "in the audited CI scope. The manual Phala run is system-level evidence. Two shell-based "
        "service smoke checks query Anvil eth_chainId and IPFS version; counted separately from "
        "the 808 framework definitions and not considered complete system tests. Production is "
        "an environment, not a fourth formal test level. No dedicated production-test suite was "
        "identified; deployment readiness and existing operational metrics are not counted as one.\n\n"
        "CI/CD\n"
        "The configured test suites run on GitHub-hosted Ubuntu runners for pushes/PRs on main "
        "and phala_app_key, or manual dispatch. Python uses pytest/unittest, Node uses node --test, "
        "contracts use Foundry's local EVM; integration dependencies run locally on the runner. "
        "Tool/environment-dependent cases can skip. The general docker-builds job has no dependency "
        "on all test jobs; do not infer a global test gate before publication. Dedicated component "
        "publish workflows run selected checks before pushing, and optional Phala deployment depends "
        "on publication. No complete deployed lifecycle is automatically tested by these workflows.\n\n"
        "AUDIT AND IMPLEMENTATION MAP\n" + report)
    return s


def main():
    prs = c.b.prepare_template()
    prs.core_properties.title = "VITA-FL — Implemented tests by test level"
    build(prs)
    check_layout(prs)
    check_bounds(prs)
    output = c.ROOT / "assets" / "test-evaluation-a.pptx"
    prs.save(output)
    reopened = Presentation(output)
    preview = c.renderer.render_slide(reopened, reopened.slides[0])
    preview.save(c.HERE / "a-test-categories.png")
    preview.save(c.HERE / "a-test-categories.pdf", resolution=120)
    (c.HERE / "selected-validation.json").write_text(json.dumps({
        "layout": "passed", "bounds": "passed", "native_editable_slide": True,
        "test_definitions": 808, "unit_component": 731, "local_integration": 77,
        "automated_full_system": 0, "dedicated_production": 0,
        "separate_service_smoke_checks": 2, "separate_recorded_cloud_runs": 1,
        "test_execution": "not performed", "asset": str(output.relative_to(c.ROOT)),
    }, indent=2) + "\n")
    print("Built selected A: 808 definitions, 731 unit/component, 77 local integration.")


if __name__ == "__main__":
    main()
