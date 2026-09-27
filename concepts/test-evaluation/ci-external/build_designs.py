#!/usr/bin/env python3
"""Three editable proposals for CI coverage and configured external system tests."""
from pathlib import Path
import hashlib
from collections import Counter
import html
import json
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[2]
sys.path.insert(0,str(ROOT))
import build_presentation as b
import render_preview as renderer
sys.path.insert(0,str(ROOT/'concepts/color-navigation'))
from build_concepts import check_layout

INK, MUTED, LINE, WHITE='263440','65737D','DCE3E8','FFFFFF'
BLUE, PURPLE='1764A1','7446A6'
BLUE_TINT, PURPLE_TINT, PALE='EDF4FA','F3EEF8','F2F5F7'
DATA=json.loads((HERE/'slide-data.json').read_text())
CI=DATA['ci_groups']
TOTAL=sum(row['count'] for row in CI)
EXTERNAL=[
    ('Settings · configuration','Save/read parameters; reject invalid deadlines'),
    ('Lifecycle · normal / inference','Complete training; use and verify the finalized model'),
    ('Collection · early / partial / zero','Close early, publish partial inputs, or recover from zero'),
    ('Recovery · aggregator outage','Stop CVM; 1/3 → 2/3 → 3/3 votes; elect a replacement'),
    ('Security · access and contract guards','16 checks: access, contracts + positive inference'),
    ('Evidence · published receipt audit','17 checks: SCITT, AIR/TDX, sessions and Sello'),
]
SCENARIOS=[
    ('configuration','Readback matches; invalid deadline rejected'),
    ('normal','5 completed rounds; final model verified'),
    ('early','Accept 5/5 → close before the deadline'),
    ('partial','Accept 2/5 → publish those 2 after timeout'),
    ('zero','Accept 0/5 → no publication; quorum recovery'),
    ('recovery','Stop aggregator → quorum → new aggregator'),
    ('inference','Verify inference using the finalized model'),
    ('security','16 checks: access, contracts + positive control'),
]
NOTES='''PURPOSE
Alternative replacement for the current test-category slide. Unit/component and
local integration tests are grouped together as CI checks and divided by the
subject under test. External scenario types are a separate counting unit.
Static code/configuration inspection on 27 September 2026; no tests or cloud
campaigns were executed to build these slides. Counts are implementation counts,
not passed executions or independent empirical samples. The current presentation
and its selected test slide remain unchanged pending selection.

CI BOUNDARY AND COUNTING
878 unique CI-selected declarations: 617 Python, 145 Solidity and 116 Node tests.
The seven groups are mutually exclusive assignments of test source areas. Each
declaration counts once; parametrization, loops, repeated workflow invocations
and subtests are not expanded. Conditional skips remain included. The scope is
vita-fl/.github/workflows/ci.yml and the explicit test entrypoints it selects.
The node job installs Foundry and requires the Anvil transaction tests. The CI
runs on GitHub-hosted Ubuntu for push/PR on main and phala_app_key or manual
workflow dispatch. Lint, formatting, build steps and the two runtime readiness
checks for Anvil/IPFS are additional pipeline checks, excluded from the 878.
The previous 808 inventory is historical. No new unit/integration split is
asserted: new tests require fresh assertion-level classification if that split
is ever needed again. Unit/component/integration are merged by explicit request.
Tests of the TD, SMEW and SMA tools are not silently added to the vita-fl CI count.

EXTERNAL SCOPE AND CONFIGURATION
The TD has nine CLI modes. Eight are considered here: configuration, normal,
early, partial, zero, recovery, inference and security. The optional manipulation
mode is excluded according to the user's requested scope. Tool-order attacks are
also excluded. Eight implemented modes do not mean eight passed live scenarios.
The selected local Phala profile has six workers, five successful training rounds,
client limit five, epoch one, a 45-second collection deadline, a 3600-second
scenario deadline, one-second polling, deployed-agent inference, participant
coverage and required transparency audit. Configuration is a setup-only check;
the five-round setting applies to training scenarios. train_before_inference is
enabled for the inference-only mode in this profile. Bootstrap and aborted
attempts are excluded from the five successful rounds. Other runs can vary the
parameters, so the profile is not attributed to every historical result.
SMEW schedules campaigns/repetitions and validates reports; SMA supervises the
local TD subprocess and retrieves Prometheus data for the run window. The TD
drives the remote VITA-FL deployment, checks assertions and produces the verdict.
No general energy measurement is claimed. Measurements currently include round
progress, aborted attempts and transaction gas.

EXTERNAL ORACLES
configuration: five training parameters round-trip; invalid timeouts are rejected
with HTTP 400 and the saved configuration remains unchanged.
normal: completed nonempty rounds, count/root/policy consistency, unchanged frozen
roster and all-six-worker participation; finalized-model inference is verified.
early/partial/zero: a cooperative admin-controlled upload gate at the deployed
receiver limits accepted clients in the first training round. Respectively 5,
2 or 0 accepted inputs establish early closure, exact partial publication after
the deadline, or no publication followed by quorum recovery. This is explicit
test instrumentation; the gate is not described as a network outage. The receipt
of a gate command alone does not establish the expected contribution set.
recovery: stop and confirm the selected aggregator CVM, observe three distinct
eligible reports out of five voters, keep the old state below quorum, then prove
abort/reselection on the third vote and successful resumed publication. Restore
the CVM; cleanup failure invalidates the run. The local aggregator timer is not
replaced or advanced by the driver. Contract vote conditions remain separate.
security: 16 required checks comprise five UI/control checks, four read-only EVM
guards, one positive inference control and six inference authorization negatives.
EVM guards use eth_call; they do not submit an attack transaction. Wrong reasons,
TLS/network errors and unavailable positive controls do not count as rejection.
Audit: 17 additional checks are 10 SCITT publisher/inclusion checks, one AIR/TDX,
three attestation sessions and three Sello receipts. Re-fetch published bytes,
verify bindings to this run and local trust policy, decrypt owner-authorized
Sello context and request fresh Phala quote appraisal. This is independent
execution using production verifier primitives, not an independent Intel DCAP
implementation or a reconstruction of the original TLS/PoP handshake.
The audit spans inference-enabled scenarios; it is not a ninth scenario mode.
16 security checks and 17 audit checks are not added to the scenario count.

INTERPRETATION
This slide inventories configured tests, not clinical quality, LLM reasoning,
load capacity or complete coverage of Chapter 4 threats. Configuration and code
availability are kept separate from observed live outcomes. Existing reports
contain a five-round Recovery PASS and a complete 17-check audit; early, partial,
zero and security must not be described as demonstrated successes from that run.
A cross-boundary mapping in variant B connects related questions, not equivalent
coverage or one-to-one test cases. Some code properties have no live fault probe.

SOURCE MAP
Full current CI declaration manifest and source hashes: ci-inventory.json.
External scenario groups and public configuration: external-scope.json.
Presentation-facing labels/mapping: slide-data.json.
vita-fl/.github/workflows/ci.yml
vita-fl-td/src/vita_fl_td/config.py
vita-fl-td/src/vita_fl_td/runner.py
vita-fl-td/src/vita_fl_td/quorum.py
vita-fl-td/src/vita_fl_td/security_bridge.py
vita-fl-td/src/vita_fl_td/transparency_audit.py
vita-fl-td/docs/security-scenarios.md
vita-fl-td/docs/transparency-audit.md
SMEW/README.md
sustainability-measurement-agent/sma/sma.py
'''


def txt(s,value,x,y,w,h=.32,size=15,color=INK,bold=False,center=False):
    q=b.add_text(s,value,x,y,w,h,size,color,bold,
                 align=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT,
                 valign=MSO_ANCHOR.MIDDLE,margin=0)
    q.name='CI/external scope: '+value.replace('\n',' / ')
    return q


def box(s,x,y,w,h,fill=PALE,border=LINE,radius=False):
    return b.add_box(s,x,y,w,h,fill=fill,line=border,radius=radius,line_width=.8)


def line(s,x,y,w,color=LINE):
    b.add_line_segment(s,x,y,x+w,y,color=color,width=.7)


def page(prs,letter,title):
    s=b.new_content_slide(prs,len(prs.slides)+1,title,
                          'EVALUATION · CI COVERAGE AND EXTERNAL TEST CONFIGURATION')
    b.add_domain_navigation(s,26,active_domains={'DFL','Agent'})
    for q in s.shapes:
        if q.has_text_frame and q.text.startswith('Page '):
            q.text_frame.paragraphs[0].runs[0].text='Variant '+letter
            q.width=b.Inches(1.1)
    b.add_note(s,'VARIANT '+letter+'\n\n'+NOTES+'\n\nGROUP COUNTS\n'+
               '\n'.join(f"{g['label']}: {g['count']} — {g['detail']}" for g in CI))
    return s


def profile(s,y=5.98):
    txt(s,'Phala training profile: 6 workers · 5 successful rounds · client limit 5 · 45 s deadline. Counts ≠ passed runs.',
        .72,y,11.91,.29,11.3,MUTED)


def design_a(prs):
    s=page(prs,'A','CI checks and external system tests')
    box(s,.70,1.54,5.81,4.25,WHITE)
    box(s,6.70,1.54,5.93,4.25,WHITE)
    box(s,.70,1.54,5.81,.70,BLUE_TINT,BLUE_TINT)
    box(s,6.70,1.54,5.93,.70,PURPLE_TINT,PURPLE_TINT)
    txt(s,f'CI · {TOTAL} definitions',.91,1.63,5.39,.36,21,BLUE,True)
    txt(s,'Unit / component / integration · GitHub Actions',.91,2.04,5.39,.24,11.3,BLUE)
    txt(s,'External · 8 modes + audit',6.91,1.63,5.50,.36,21,PURPLE,True)
    txt(s,'SMEW → SMA → vita-fl-td · deployed VITA-FL in Phala',6.91,2.04,5.50,.24,11.3,PURPLE)
    for i,g in enumerate(CI):
        y=2.35+i*.48
        txt(s,g['label'],.91,y,4.61,.27,14,INK,True)
        txt(s,g['detail'],.91,y+.255,4.88,.23,11.2,MUTED)
        txt(s,str(g['count']),5.72,y+.04,.55,.34,18,BLUE,True,True)
        if i<6: line(s,.91,y+.475,5.36)
    for i,(name,detail) in enumerate(EXTERNAL):
        y=2.35+i*.555
        txt(s,name,6.91,y,5.50,.29,14,INK,True)
        txt(s,detail,6.91,y+.295,5.50,.23,11.4,MUTED)
        if i<5: line(s,6.91,y+.535,5.49)
    profile(s)


def design_b(prs):
    s=page(prs,'B','One system, two complementary test boundaries')
    txt(s,f'{TOTAL} CI definitions  +  8 external scenario modes and a receipt audit',
        .72,1.52,11.91,.45,19,INK)
    cols=[('TESTED AREA',.86,2.37),('CI #',3.20,.58),
          ('CI · UNIT / COMPONENT / INTEGRATION',3.96,4.0),
          ('EXTERNAL · DEPLOYED VITA-FL',8.33,4.08)]
    for label,x,w in cols: txt(s,label,x,2.02,w,.28,10.9,MUTED,True)
    for i,g in enumerate(CI):
        y=2.43+i*.47
        box(s,.70,y,11.93,.46,PALE if i%2==0 else WHITE,WHITE)
        txt(s,g['matrix_label'],.86,y+.015,2.20,.43,12.8,INK,True)
        txt(s,str(g['count']),3.20,y+.02,.58,.39,18,BLUE,True,True)
        txt(s,g['ci_check'],3.96,y+.015,4.12,.43,12.1)
        txt(s,g['external'],8.33,y+.015,4.08,.43,12.1,PURPLE)
    txt(s,'CI: GitHub Actions / Ubuntu, push + PR + manual.  External: SMEW / SMA / TD → Phala.',
        .72,5.79,11.91,.27,11.2,MUTED)
    profile(s,6.04)


def design_c(prs):
    s=page(prs,'C','From CI coverage to configured cloud scenarios')
    box(s,.70,1.54,4.19,4.28,WHITE)
    box(s,.70,1.54,4.19,.70,BLUE_TINT,BLUE_TINT)
    txt(s,f'CI · {TOTAL} definitions',.89,1.65,3.79,.35,20,BLUE,True)
    txt(s,'Unit / component / integration',.89,2.045,3.79,.23,11.3,BLUE)
    for i,g in enumerate(CI):
        y=2.40+i*.445
        txt(s,g['short_label'],.90,y,3.17,.26,13.4,INK,True)
        txt(s,g['short_detail'],.90,y+.26,3.42,.20,10.3,MUTED)
        txt(s,str(g['count']),4.15,y+.04,.49,.29,16,BLUE,True,True)
    txt(s,'GitHub Actions · Ubuntu · push / PR / manual',.90,5.61,3.80,.20,9.9,BLUE)
    txt(s,'8 external scenario modes',5.13,1.63,7.49,.37,21,PURPLE,True)
    txt(s,'SMEW / SMA / vita-fl-td → Phala',5.13,2.055,7.49,.27,12,PURPLE)
    txt(s,'SCENARIO',5.25,2.46,1.67,.26,11,MUTED,True)
    txt(s,'CONDITION → EXPECTED RESPONSE',7.16,2.46,5.24,.26,11,MUTED,True)
    for i,(name,detail) in enumerate(SCENARIOS):
        y=2.81+i*.315
        box(s,5.12,y,7.51,.305,PALE if i%2==0 else WHITE,WHITE)
        txt(s,name,5.25,y+.015,1.71,.275,12.1,PURPLE,True)
        txt(s,detail,7.16,y+.015,5.25,.275,12.1)
    box(s,5.12,5.45,7.51,.37,PURPLE_TINT,PURPLE_TINT)
    txt(s,'Additional audit: 17 checks of published SCITT / AIR / TDX / Sello evidence',
        5.26,5.49,7.21,.28,11.3,PURPLE,True)
    profile(s)


DESIGNS=[
    ('a-two-catalogs','A · Zwei Prüfkataloge','CI nach Prüfgegenstand; externe Szenarien mit ihren wichtigsten Assertions.',design_a),
    ('b-boundary-matrix','B · Gemeinsame Vergleichsmatrix','Verwandte Prüfgegenstände über die CI- und Deployment-Grenze hinweg vergleichen.',design_b),
    ('c-scenario-configuration','C · Konkrete Testbedingungen','CI kompakt links; jeder externe Szenariotyp mit Bedingung und erwartetem Verhalten rechts.',design_c),
]


def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()


def gallery():
    cards=[]
    for slug,label,desc,_ in DESIGNS:
        cards.append(f'<article><h2>{html.escape(label)}</h2><p>{html.escape(desc)}</p>'
            f'<a href="{slug}.png"><img src="{slug}.png" alt="{html.escape(label)}"></a>'
            f'<p><a href="{slug}.pptx">Editierbare PowerPoint-Folie</a> · <a href="{slug}.png">Vollbild</a></p></article>')
    (HERE/'index.html').write_text('''<!doctype html><html lang="de"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>VITA-FL · CI und externe Tests</title><style>body{font:16px/1.5 system-ui,sans-serif;background:#edf1f4;color:#263440;margin:0}header,main{max-width:1500px;margin:auto;padding:24px}h1,h2{margin:0}a{color:#1764a1}main{display:grid;gap:28px}article{background:white;padding:24px;border-radius:12px}img{width:100%;display:block;border:1px solid #dce3e8}header p{max-width:1050px}nav{display:flex;gap:24px;flex-wrap:wrap}</style><header><h1>CI-Abdeckung und externe Tests</h1><p>878 CI-Testdefinitionen, fachlich gegliedert. Dazu acht berücksichtigte externe Szenariotypen und ein szenarioübergreifender Receipt-Audit. Die Zahlen sind Implementierungsumfang, keine bestandenen Live-Läufe.</p><nav><a href="VITA-FL_CI_External_Test_Alternatives.pptx">Alle Varianten als PowerPoint</a><a href="VITA-FL_CI_External_Test_Alternatives.pdf">PDF-Vorschau</a><a href="overview.png">Gesamtübersicht</a></nav></header><main>'''+''.join(cards)+'</main></html>',encoding='utf-8')


def main():
    assert TOTAL==878 and len(CI)==7
    inventory=json.loads((HERE/'ci-inventory.json').read_text())
    assert len(inventory['tests'])==TOTAL
    assert len({t['id'] for t in inventory['tests']})==TOTAL
    counted=Counter(t['functional_group'] for t in inventory['tests'])
    assert dict(counted)=={g['id']:g['count'] for g in CI}
    for path, expected in inventory['source_files'].items():
        assert digest(ROOT.parent/'vita-fl'/path)==expected, 'CI inventory needs refresh: '+path
    assert len({g['id'] for g in CI})==7
    external=json.loads((HERE/'external-scope.json').read_text())
    assert external['considered_scenarios_total']==8
    assert sum(g['count'] or 0 for g in external['groups'])==8
    protected=[ROOT/'VITA-FL_Thesis_Presentation_TU_Berlin.pptx',
               ROOT/'assets/test-evaluation-a.pptx',HERE.parent/'a-test-categories.png',
               HERE.parent/'test-inventory.json']
    before={str(p.relative_to(ROOT)):digest(p) for p in protected}
    p=b.prepare_template()
    p.core_properties.title='VITA-FL — CI coverage and external test configuration'
    for _,_,_,build in DESIGNS: build(p)
    check_layout(p)
    output=HERE/'VITA-FL_CI_External_Test_Alternatives.pptx'
    p.save(output)
    q=Presentation(output)
    check_layout(q)
    previews=[]
    for s,(slug,_,_,build) in zip(q.slides,DESIGNS):
        assert 'SOURCE MAP' in s.notes_slide.notes_text_frame.text
        visible='\n'.join(x.text for x in s.shapes if x.has_text_frame)
        assert str(TOTAL) in visible and 'CI' in visible
        im=renderer.render_slide(q,s)
        im.save(HERE/(slug+'.png'))
        previews.append(im)
        single=b.prepare_template()
        build(single)
        check_layout(single)
        single.save(HERE/(slug+'.pptx'))
        assert len(Presentation(HERE/(slug+'.pptx')).slides)==1
    previews[0].save(HERE/'VITA-FL_CI_External_Test_Alternatives.pdf',save_all=True,
                     append_images=previews[1:],resolution=120)
    overview=Image.new('RGB',(1660,1550),'#E8ECEF')
    draw=ImageDraw.Draw(overview)
    font=ImageFont.truetype(renderer.FONT_BOLD,26)
    for i,(im,(_,label,_,_)) in enumerate(zip(previews,DESIGNS)):
        # Wide, readable previews, stacked rather than a four-up with one empty tile.
        y=14+i*510
        draw.text((35,y),label,fill='#263440',font=font)
        overview.paste(im.resize((800,450),Image.Resampling.LANCZOS),(35,y+43))
        description=[
            ['Zwei klar getrennte Bereiche','Links: 878 CI-Definitionen in sieben Gruppen','Rechts: acht Szenariotypen + Receipt-Audit','Empfehlung für einen schnellen Überblick'],
            ['Prüfgegenstände im direkten Vergleich','Welche Eigenschaft prüft der CI-Code?','Was wird zusätzlich am Deployment geprüft?','Zeigt auch Unterschiede im Prüfumfang'],
            ['Konfiguration und erwartetes Verhalten','Jeder externe Szenariotyp als eigene Zeile','Early: 5/5 · Partial: 2/5 · Zero: 0/5','Empfehlung für die Erklärung des Testplans'],
        ][i]
        bodyfont=ImageFont.truetype(renderer.FONT_REGULAR,24)
        headfont=ImageFont.truetype(renderer.FONT_BOLD,26)
        for j,t in enumerate(description):
            draw.text((890,y+140+j*58),t,fill='#263440' if j==0 else '#65737D',font=headfont if j==0 else bodyfont)
    overview.save(HERE/'overview.png')
    gallery()
    assert before=={str(p.relative_to(ROOT)):digest(p) for p in protected}
    (HERE/'validation.json').write_text(json.dumps({
        'slides':3,'ci_definitions':TOTAL,'disjoint_ci_groups':7,
        'external_scenario_modes':8,'receipt_audit_checks':17,'security_scenario_checks':16,
        'layout':'passed','pptx_roundtrip':'passed','speaker_notes':'present',
        'native_editable_shapes':True,'main_and_selected_slide_preserved':before,
        'test_execution':'none; source/configuration inventory only',
        'preview_renderer':'Pillow; PowerPoint remains authoritative for exact fonts',
    },indent=2)+'\n')
    print('Built 3 editable proposals, individual slides, PNG/PDF previews and gallery.')
    print('Layout passed; main deck, approved test slide and historical inventory unchanged.')


if __name__=='__main__': main()
