#!/usr/bin/env python3
"""Insert/update approved external-results variant B after the test inventory.

Preserve every other slide, its notes, relationships and media. Only page-number
footers change after insertion; the candidate is validated before replacement.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import shutil
import tempfile
from zipfile import ZipFile

from pptx import Presentation
from pptx.util import Inches

import build_presentation as b
from update_slide25_test_evaluation import snapshot, visible, footer, check_structure as check_inventory

ROOT = Path(__file__).resolve().parent
DECK = b.DESTINATION
ASSET = b.EXTERNAL_RESULTS_ASSET
RECORD = ROOT / 'assets/external-test-results-integration.json'
CHAPTER = ROOT.parent / 'overleaf/chapters/chapter4.tex'
PROPOSAL = ROOT / 'concepts/external-results/b-evidence-matrix.pptx'


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def source_footer(shape):
    return (shape.has_text_frame and shape.top > Inches(6.7)
            and shape.text.startswith(('Variant ', 'Page ')))


def check_structure(prs):
    assert b.external_results_count(prs) == 1
    check_inventory(prs)
    target = b.external_results_target(prs)
    assert 'What changed over the 24 federated rounds?' in visible(prs.slides[target+1])
    slide = prs.slides[target]
    assert not slide.element.xpath('.//a:blip'), 'Results must remain native editable shapes'
    assert 'Variant B' not in visible(slide)
    assert all(term in visible(slide) for term in ['3 / 3', '16 / 16', '17 / 17', 'HTTP 401', 'eth_call', 'RTMR3', 'Sello'])
    assert 'SOURCE MAP' in slide.notes_slide.notes_text_frame.text
    assert footer(slide).text == f'Page {target+1}'
    for shape in slide.shapes:
        assert shape.left >= 0 and shape.top >= 0, shape.name
        assert shape.left + shape.width <= prs.slide_width, shape.name
        assert shape.top + shape.height <= prs.slide_height, shape.name
    return prs


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--render', action='store_true', help='Refresh deck previews and contact sheet')
    args = parser.parse_args()
    assert ASSET.is_file()
    prs = Presentation(DECK)
    old_count = len(prs.slides)
    inserting = b.external_results_count(prs) == 0
    target = b.external_results_target(prs)
    assert old_count == 54 + int(not inserting), 'Unexpected deck size; refusing to overwrite'
    if inserting:
        check_inventory(prs)
        assert 'What changed over the 24 federated rounds?' in visible(prs.slides[target])
    else:
        check_structure(prs)
    before = [snapshot(slide) for slide in prs.slides]
    before_hash, chapter_hash = digest(DECK), digest(CHAPTER)
    backup = Path(tempfile.mkdtemp(prefix='vita-fl-before-external-results-')) / DECK.name
    shutil.copy2(DECK, backup)
    if inserting:
        b.insert_external_results_slide(prs)
    else:
        b.apply_external_results_asset(prs.slides[target], target+1)
    check_structure(prs)
    handle, filename = tempfile.mkstemp(prefix='.external-results-', suffix='.pptx', dir=ROOT)
    os.close(handle)
    candidate = Path(filename)
    try:
        prs.save(candidate)
        checked = check_structure(Presentation(candidate))
        preserved = 0
        for index, old in enumerate(before):
            if not inserting and index == target:
                continue
            new_index = index + int(inserting and index >= target)
            assert snapshot(checked.slides[new_index]) == old, f'Existing Page {index+1} changed beyond its footer'
            preserved += 1
        source = Presentation(ASSET).slides[0]
        integrated = checked.slides[target]
        assert len(source.shapes) == len(integrated.shapes)
        assert integrated.notes_slide.notes_text_frame.text == source.notes_slide.notes_text_frame.text
        for approved, actual in zip(source.shapes, integrated.shapes):
            if not source_footer(approved):
                assert str(approved.element.xml) == str(actual.element.xml), approved.name
        with ZipFile(backup) as old_zip, ZipFile(candidate) as new_zip:
            assert new_zip.testzip() is None
            for name in old_zip.namelist():
                if name.startswith(('ppt/media/', 'ppt/slideMasters/', 'ppt/slideLayouts/', 'ppt/notesMasters/')):
                    assert old_zip.read(name) == new_zip.read(name), name
        once = [snapshot(slide) for slide in checked.slides]
        b.apply_external_results_asset(checked.slides[target], target+1)
        assert [snapshot(slide) for slide in checked.slides] == once, 'Repeated asset application changed the deck'
        assert digest(DECK) == before_hash, 'Deck changed during validation'
        assert digest(CHAPTER) == chapter_hash
        os.replace(candidate, DECK)
    finally:
        candidate.unlink(missing_ok=True)
    RECORD.write_text(json.dumps({
        'page': target+1, 'mode': 'insert' if inserting else 'replace',
        'slides_before': old_count, 'slides_after': len(checked.slides),
        'talk_pages': len(checked.slides)-17, 'backup_pages': 17,
        'preserved_other_slides_and_notes': preserved,
        'allowed_existing_change': 'Page-number footer only after insertion; none outside replacement on update',
        'main_deck_sha256_before': before_hash, 'main_deck_sha256_after': digest(DECK),
        'asset_sha256': digest(ASSET), 'approved_proposal_sha256': digest(PROPOSAL),
        'chapter4_sha256': chapter_hash, 'chapter4_unchanged': True,
        'backup': str(backup), 'structure_check': 'passed', 'idempotent_asset_update': 'passed',
        'approved_shape_xml_preserved_except_footer': True, 'source_notes_preserved': True,
        'media_masters_layouts_unchanged': True, 'native_editable_shapes': True,
        'visual_review': 'pending',
    }, indent=2)+'\n')
    if args.render:
        import render_preview
        render_preview.main()
    print(f'{"Inserted" if inserting else "Updated"} Page {target+1}; {len(checked.slides)} slides. Preserved {preserved} other slides and notes.')
    print(f'Backup: {backup}')


if __name__ == '__main__':
    main()
